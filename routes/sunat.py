# routes/sunat.py - Módulo "SUNAT e impuestos" (pedido de Erika, 21/10/2026)
#
# Qué hace (en simple):
#   1. Resumen de IGV y renta por empresa y por mes, "a hoy" y "si hoy fuera fin de mes" (proyección).
#   2. Plan de acción como auditor: qué comprar antes del cierre, qué falta facturar o registrar, alertas legales.
#   3. Preliminar con el MISMO formato de la hoja del contador (PDT 621 + PLAME) para compararlo con lo que él envía.
#   4. Pagos del mes: checklist de AFP, EsSalud, ONP, renta, IGV, sueldos... con su vencimiento SUNAT y check de pagado.
#   5. Completar datos: importar la propuesta del SIRE (RVIE/RCE) o el Excel de las chicas, porque todavía no todas
#      las facturas están en el ERP. Lo importado se guarda aparte (sunat_registros) y NUNCA cambia ventas ni compras.
#
# Reglas:
#   - Es un cálculo PRELIMINAR. Quien declara es el contador.
#   - Solo LEE ventas, compras, guías, pedidos y órdenes. Escribe únicamente en sus tablas sunat_*.
#   - Acceso: solo administrador (Gerencia total) o quien tenga "Ver" en el permiso SUNAT. Se valida en el servidor.
#   - Empresa: respeta el selector de arriba (KCF / AGD). Con "Todas" se elige la empresa en la pantalla.
import calendar
import csv
import io
import json
import re
import unicodedata
import uuid
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal

from flask import Blueprint, render_template, request, jsonify, session, Response
from psycopg2.extras import RealDictCursor, Json

from database import db_query, db_tx
from utils import a_decimal, redondear, es_admin, tasa_igv, id_empresa
from routes.multiempresa import empresas as lista_empresas, empresa_actual, empresa_por_id

sunat_bp = Blueprint('sunat', __name__)

CERO = Decimal('0')
ANULADOS = ('ANULADO', 'ANULADA', 'ANULADOS', 'BORRADOR', 'RECHAZADO', 'RECHAZADA', 'ELIMINADO', 'ELIMINADA')
UIT_2026 = Decimal('5500')            # D.S. 301-2025-EF
BANCARIZACION = Decimal('2000')       # pagos desde S/ 2,000 (o US$ 500) deben ir por banco para usar el crédito fiscal
MESES = ('Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio', 'Julio', 'Agosto', 'Setiembre', 'Octubre',
         'Noviembre', 'Diciembre')

# ------------------------------------------------------------
# CRONOGRAMA SUNAT 2026 (obligaciones mensuales: PDT 621 IGV-Renta y PLAME). R.S. 281-2022/SUNAT, Anexo I.
# periodo -> (mes de vencimiento, {grupo de último dígito del RUC: día})
# ------------------------------------------------------------
_G = ('0', '1', '23', '45', '67', '89')
CRONOGRAMA = {
    (2026, 1): ((2026, 2), (16, 17, 18, 19, 20, 23)),
    (2026, 2): ((2026, 3), (16, 17, 18, 19, 20, 23)),
    (2026, 3): ((2026, 4), (17, 20, 21, 22, 23, 24)),
    (2026, 4): ((2026, 5), (18, 19, 20, 21, 22, 25)),
    (2026, 5): ((2026, 6), (15, 16, 17, 18, 19, 22)),
    (2026, 6): ((2026, 7), (15, 16, 17, 20, 21, 22)),
    (2026, 7): ((2026, 8), (18, 19, 20, 21, 24, 25)),
    (2026, 8): ((2026, 9), (15, 16, 17, 18, 21, 22)),
    (2026, 9): ((2026, 10), (16, 19, 20, 21, 22, 23)),
    (2026, 10): ((2026, 11), (16, 17, 18, 19, 20, 23)),
    (2026, 11): ((2026, 12), (17, 18, 21, 22, 23, 24)),
    (2026, 12): ((2027, 1), (18, 19, 20, 21, 22, 25)),
}
FERIADOS = {date(2026, 1, 1), date(2026, 4, 2), date(2026, 4, 3), date(2026, 5, 1), date(2026, 6, 7), date(2026, 6, 29),
            date(2026, 7, 23), date(2026, 7, 28), date(2026, 7, 29), date(2026, 8, 6), date(2026, 8, 30),
            date(2026, 10, 8), date(2026, 11, 1), date(2026, 12, 8), date(2026, 12, 9), date(2026, 12, 25),
            date(2027, 1, 1), date(2027, 3, 25), date(2027, 3, 26)}

# Obligaciones del checklist (orden = orden en que Erika paga). cuando: 'cronograma' (por RUC), 'habil5' (5.º día
# hábil del mes siguiente), 'fin_mes' (último día hábil del mismo mes), 'dia15' (día 15 de ciertos meses).
OBLIGACIONES = (
    {'codigo': 'REMUNERACIONES', 'nombre': 'Sueldos (remuneraciones)', 'donde': 'Banco', 'cuando': 'fin_mes',
     'ayuda': 'Pagar a los trabajadores en planilla. Se paga el mismo mes.'},
    {'codigo': 'AFP', 'nombre': 'AFP', 'donde': 'AFPnet', 'cuando': 'habil5',
     'ayuda': 'Aporte descontado a los trabajadores con AFP. No se paga en SUNAT, se paga en AFPnet.'},
    {'codigo': 'ESSALUD', 'nombre': 'EsSalud (9%)', 'donde': 'SUNAT · PLAME', 'cuando': 'cronograma',
     'ayuda': '9% de las remuneraciones. Lo paga la empresa.'},
    {'codigo': 'ONP', 'nombre': 'ONP (13%)', 'donde': 'SUNAT · PLAME', 'cuando': 'cronograma',
     'ayuda': '13% descontado a los trabajadores que están en ONP. La empresa lo retiene y lo paga.'},
    {'codigo': 'RENTA_5TA', 'nombre': 'Renta de 5.ª categoría', 'donde': 'SUNAT · PLAME', 'cuando': 'cronograma',
     'ayuda': 'Retención a sueldos altos. Si nadie gana lo suficiente, no aplica.'},
    {'codigo': 'RENTA_4TA', 'nombre': 'Renta de 4.ª categoría (recibos por honorarios)', 'donde': 'SUNAT · PLAME',
     'cuando': 'cronograma', 'ayuda': '8% retenido a recibos por honorarios mayores a S/ 1,500.'},
    {'codigo': 'RENTA_3RA', 'nombre': 'Renta mensual (pago a cuenta)', 'donde': 'SUNAT · PDT 621', 'cuando': 'cronograma',
     'ayuda': 'Porcentaje de las ventas del mes, según el régimen de la empresa.'},
    {'codigo': 'IGV', 'nombre': 'IGV', 'donde': 'SUNAT · PDT 621', 'cuando': 'cronograma',
     'ayuda': 'IGV de ventas menos crédito fiscal, retenciones y percepciones. Si sale cero, no se paga.'},
    {'codigo': 'DECLARACION_621', 'nombre': 'Presentar la declaración mensual (PDT 621)', 'donde': 'SUNAT · PDT 621',
     'cuando': 'cronograma', 'ayuda': 'Se presenta SIEMPRE, aunque no haya nada que pagar.'},
    {'codigo': 'DETRACCIONES', 'nombre': 'Detracciones de compras por depositar', 'donde': 'Banco de la Nación',
     'cuando': 'habil5', 'ayuda': 'Si compraste servicios con detracción, deposítala o pierdes el crédito fiscal de esa factura.'},
    {'codigo': 'CTS', 'nombre': 'CTS', 'donde': 'Banco', 'cuando': 'dia15', 'meses': (5, 11),
     'ayuda': 'Depósito de CTS de los trabajadores (mayo y noviembre).'},
    {'codigo': 'GRATIFICACION', 'nombre': 'Gratificación', 'donde': 'Banco', 'cuando': 'dia15', 'meses': (7, 12),
     'ayuda': 'Gratificación de julio y diciembre.'},
)
CODIGOS_OBLIGACION = {o['codigo']: o for o in OBLIGACIONES}
# Las que por defecto salen "No aplica" si no hay monto (Erika puede cambiarlo)
NO_APLICA_SI_CERO = ('ONP', 'RENTA_5TA', 'RENTA_4TA', 'IGV', 'DETRACCIONES', 'RENTA_3RA', 'AFP', 'ESSALUD',
                     'REMUNERACIONES', 'CTS', 'GRATIFICACION')

REGIMENES = {
    'RMT': {'nombre': 'Régimen MYPE Tributario', 'tasa': Decimal('1.0')},
    'GENERAL': {'nombre': 'Régimen General', 'tasa': Decimal('1.5')},
    'RER': {'nombre': 'Régimen Especial (RER)', 'tasa': Decimal('1.5')},
}

TIPOS_CP = {'01': 'Factura', '03': 'Boleta', '07': 'Nota de crédito', '08': 'Nota de débito', '02': 'Recibo por honorarios'}


class ErrorSunat(Exception):
    pass


# ============================================================
# PERMISO (en el servidor)
# ============================================================
def puede_ver_sunat():
    """Solo administración (Gerencia total) o quien tenga "Ver" en SUNAT (Configuración ▸ Usuarios y permisos)."""
    if es_admin():
        return True
    try:
        from routes.acceso import acceso_usuario
        acc = acceso_usuario()
    except Exception:
        return False
    return bool(acc.get('configurado') and 'sunat' in acc.get('modulos', set()))


@sunat_bp.before_request
def _proteger():
    if 'usuario_id' not in session:
        return None   # main.control_de_acceso ya responde
    if not puede_ver_sunat():
        msg = ('No tienes permiso para entrar a SUNAT e impuestos. Lo ve solo la gerencia y quien tenga el permiso '
               '"SUNAT" en Configuración ▸ Usuarios y permisos.')
        if '/api/' in request.path:
            return jsonify({'success': False, 'error': msg, 'codigo': 'sin_permiso'}), 403
        from flask import flash, redirect, url_for
        flash(msg, 'error')
        return redirect(url_for('index'))
    asegurar_tablas()
    return None


# ============================================================
# TABLAS PROPIAS (equivalente: migrations/017_sunat_impuestos.sql)
# ============================================================
SQL_DDL_SUNAT = """
    CREATE TABLE IF NOT EXISTS sunat_config (
        empresa_id VARCHAR(64) PRIMARY KEY,
        regimen VARCHAR(20) NOT NULL DEFAULT 'RMT',
        tasa_renta NUMERIC(6,3) NOT NULL DEFAULT 1.0,
        es_agente_retencion BOOLEAN NOT NULL DEFAULT false,
        es_buen_contribuyente BOOLEAN NOT NULL DEFAULT false,
        notas TEXT,
        actualizado_por VARCHAR(150),
        actualizado_en TIMESTAMP DEFAULT NOW()
    );
    CREATE TABLE IF NOT EXISTS sunat_periodos (
        id SERIAL PRIMARY KEY,
        empresa_id VARCHAR(64) NOT NULL,
        periodo VARCHAR(7) NOT NULL,
        manual JSONB NOT NULL DEFAULT '{}'::jsonb,
        planilla JSONB NOT NULL DEFAULT '{}'::jsonb,
        contador JSONB NOT NULL DEFAULT '{}'::jsonb,
        contador_archivo VARCHAR(255),
        contador_cargado_en TIMESTAMP,
        actualizado_por VARCHAR(150),
        actualizado_en TIMESTAMP DEFAULT NOW(),
        creado_en TIMESTAMP DEFAULT NOW(),
        UNIQUE (empresa_id, periodo)
    );
    CREATE TABLE IF NOT EXISTS sunat_registros (
        id SERIAL PRIMARY KEY,
        empresa_id VARCHAR(64) NOT NULL,
        periodo VARCHAR(7) NOT NULL,
        libro VARCHAR(10) NOT NULL CHECK (libro IN ('VENTAS', 'COMPRAS')),
        lote VARCHAR(40) NOT NULL,
        origen VARCHAR(30) NOT NULL,
        archivo VARCHAR(255),
        tipo_cp VARCHAR(4),
        serie VARCHAR(20),
        numero VARCHAR(30),
        fecha DATE,
        ruc VARCHAR(20),
        nombre VARCHAR(255),
        base NUMERIC(14,2) NOT NULL DEFAULT 0,
        igv NUMERIC(14,2) NOT NULL DEFAULT 0,
        no_gravado NUMERIC(14,2) NOT NULL DEFAULT 0,
        total NUMERIC(14,2) NOT NULL DEFAULT 0,
        moneda VARCHAR(5) DEFAULT 'PEN',
        tipo_cambio NUMERIC(10,4),
        creado_por VARCHAR(150),
        creado_en TIMESTAMP DEFAULT NOW()
    );
    CREATE INDEX IF NOT EXISTS idx_sunat_registros_periodo ON sunat_registros (empresa_id, periodo, libro);
    CREATE TABLE IF NOT EXISTS sunat_obligaciones (
        id SERIAL PRIMARY KEY,
        empresa_id VARCHAR(64) NOT NULL,
        periodo VARCHAR(7) NOT NULL,
        codigo VARCHAR(30) NOT NULL,
        estado VARCHAR(15) NOT NULL DEFAULT 'Pendiente' CHECK (estado IN ('Pendiente', 'Pagado', 'No aplica')),
        monto NUMERIC(14,2),
        vencimiento DATE,
        fecha_pago DATE,
        numero_operacion VARCHAR(80),
        observacion TEXT,
        marcado_por VARCHAR(150),
        marcado_en TIMESTAMP,
        historial JSONB NOT NULL DEFAULT '[]'::jsonb,
        UNIQUE (empresa_id, periodo, codigo)
    );
    -- Datos tributarios: nadie entra por la API pública de Supabase (el ERP usa su propia conexión)
    ALTER TABLE sunat_config ENABLE ROW LEVEL SECURITY;
    ALTER TABLE sunat_periodos ENABLE ROW LEVEL SECURITY;
    ALTER TABLE sunat_registros ENABLE ROW LEVEL SECURITY;
    ALTER TABLE sunat_obligaciones ENABLE ROW LEVEL SECURITY;
"""
_tablas_listas = {'ok': False}


def asegurar_tablas():
    if _tablas_listas['ok']:
        return
    with db_tx() as conn:
        conn.cursor().execute(SQL_DDL_SUNAT)
    _tablas_listas['ok'] = True


# ============================================================
# AYUDAS
# ============================================================
def _n(v):
    return float(redondear(v or 0, 2))


def _usuario():
    return session.get('nombre') or session.get('usuario') or f"Usuario {session.get('usuario_id')}"


def _ahora_lima():
    return datetime.now()   # main.py fija TZ=America/Lima


def _periodo(txt=None):
    """'2026-10' -> (2026, 10). Por defecto, el mes actual."""
    m = re.match(r'^(\d{4})-(\d{2})$', (txt or '').strip())
    if m and 1 <= int(m.group(2)) <= 12:
        return int(m.group(1)), int(m.group(2))
    h = date.today()
    return h.year, h.month


def _txt_periodo(anio, mes):
    return f'{anio:04d}-{mes:02d}'


def _nombre_periodo(anio, mes):
    return f'{MESES[mes - 1]} {anio}'


def _mes_mas(anio, mes, n):
    t = anio * 12 + (mes - 1) + n
    return t // 12, t % 12 + 1


def _rango(anio, mes):
    return date(anio, mes, 1), date(anio, mes, calendar.monthrange(anio, mes)[1])


def _es_habil(d):
    return d.weekday() < 5 and d not in FERIADOS


def _habil_n(anio, mes, n):
    d, cuenta = date(anio, mes, 1), 0
    while True:
        if _es_habil(d):
            cuenta += 1
            if cuenta == n:
                return d
        d += timedelta(days=1)


def _ultimo_habil(anio, mes):
    d = _rango(anio, mes)[1]
    while not _es_habil(d):
        d -= timedelta(days=1)
    return d


def vencimiento_cronograma(ruc, anio, mes):
    """Fecha de vencimiento SUNAT del periodo (anio, mes) según el último dígito del RUC (None si no está cargado)."""
    fila = CRONOGRAMA.get((anio, mes))
    digito = (str(ruc or '').strip() or '0')[-1]
    if not fila or not digito.isdigit():
        return None
    (va, vm), dias = fila
    grupo = next(i for i, g in enumerate(_G) if digito in g)
    return date(va, vm, dias[grupo])


def _empresa(eid_txt=None):
    """Empresa con la que se trabaja: la del parámetro (si es válida) o la elegida arriba."""
    eid = id_empresa(eid_txt) if eid_txt else None
    e = empresa_por_id(eid) if eid else None
    e = e or empresa_actual()
    if not e:
        lista = lista_empresas()
        e = lista[0] if lista else None
    if not e:
        raise ErrorSunat('No hay empresas configuradas (Configuración ▸ Empresas).')
    return e


def _otras_rucs(e):
    return {x['ruc']: x['codigo'] for x in lista_empresas() if x['id'] != e['id'] and x.get('ruc')}


def _norm(txt):
    t = unicodedata.normalize('NFKD', str(txt or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9/%. ]', ' ', t)).strip()


def _num_doc(numero):
    """'00012345' -> '12345' (para comparar números de comprobante)."""
    d = re.sub(r'\D', '', str(numero or ''))
    return str(int(d)) if d else str(numero or '').strip().upper()


def _clave(libro, tipo, serie, numero, ruc):
    base = (libro, (tipo or '01'), (serie or '').strip().upper(), _num_doc(numero))
    return base + ((re.sub(r'\D', '', str(ruc or '')),) if libro == 'COMPRAS' else ())


def _tipo_cp(txt):
    t = _norm(txt)
    if t in ('1', '01') or t.startswith('factura') or t == 'fac':
        return '01'
    if t in ('3', '03') or t.startswith('boleta'):
        return '03'
    if t in ('7', '07') or 'credito' in t or t == 'nc':
        return '07'
    if t in ('8', '08') or 'debito' in t or t == 'nd':
        return '08'
    if t in ('2', '02') or 'honorario' in t:
        return '02'
    return t[:4].upper() or '01'


def _cols(tabla):
    from routes.reportes import _columnas
    return _columnas(tabla)


def _tiene(tabla, col):
    from routes.reportes import _tiene as t
    return t(tabla, col)


def _sql_emp(tabla, col, eid):
    """Filtro por empresa con el mismo criterio del ERP (filas antiguas sin empresa = de todas)."""
    if not _tiene(tabla, col.split('.')[-1]):
        return '', ()
    return f' AND ({col} IS NULL OR {col}::text = %s)', (str(eid),)


# ============================================================
# DATOS DEL ERP
# ============================================================
def _tc(moneda, manual, propio=None):
    """TC para pasar a soles. 24/10/2026: primero el TC guardado en el documento; si no tiene, el TC manual del mes."""
    if (moneda or 'PEN').upper() in ('PEN', 'SOLES', 'S/', ''):
        return Decimal('1'), False
    if propio and a_decimal(propio) > 0:
        return a_decimal(propio), False
    tc = a_decimal(manual.get('tipo_cambio'))
    return (tc if tc > 0 else Decimal('1')), tc <= 0


def ventas_erp(eid, desde, hasta, manual=None):
    """Facturas, boletas y notas del ERP (fecha de emisión dentro del rango)."""
    manual = manual or {}
    sql_e, par_e = _sql_emp('comprobantes', 'empresa_id', eid)
    filas = db_query(f"""
        SELECT id, tipo_comprobante, serie, numero::text AS numero, fecha_emision::date AS fecha, moneda,
               cliente_numero_doc AS ruc, cliente_nombre AS nombre, COALESCE(subtotal, 0) AS base,
               COALESCE(igv, 0) AS igv, COALESCE(total, 0) AS total, COALESCE(tiene_retencion, false) AS tiene_ret,
               COALESCE(monto_retenido, 0) AS retenido, guia_vinculada, pc_vinculado,
               to_jsonb(comprobantes) ->> 'tipo_cambio' AS tc_doc
        FROM comprobantes
        WHERE COALESCE(UPPER(TRIM(estado_sunat)), 'BORRADOR') <> ALL(%s)
          AND fecha_emision >= %s AND fecha_emision < %s::date + 1 {sql_e}
        ORDER BY fecha_emision, id""", (list(ANULADOS), desde, hasta) + par_e)
    docs, sin_tc = [], False
    for f in filas:
        tipo = _tipo_cp(f['tipo_comprobante'])
        tc, falta = _tc(f['moneda'], manual, f.get('tc_doc'))
        sin_tc = sin_tc or falta
        signo = Decimal('-1') if tipo == '07' else Decimal('1')
        base, igv = abs(a_decimal(f['base'])) * tc * signo, abs(a_decimal(f['igv'])) * tc * signo
        docs.append({'origen': 'ERP', 'id': f['id'], 'tipo': tipo, 'serie': f['serie'], 'numero': f['numero'],
                     'fecha': f['fecha'], 'ruc': f['ruc'], 'nombre': f['nombre'], 'base': base, 'igv': igv,
                     'no_gravado': CERO, 'total': base + igv, 'moneda': f['moneda'] or 'PEN',
                     'tipo_cambio': tc if tc != 1 else None,
                     'retenido': a_decimal(f['retenido']) * tc if f['tiene_ret'] else CERO,
                     'guia': f['guia_vinculada'], 'pc': f['pc_vinculado']})
    # Notas de crédito (tabla propia): el monto incluye IGV
    if _cols('notas_credito'):
        sql_n, par_n = _sql_emp('notas_credito', 'empresa_id', eid)
        for f in db_query(f"""
                SELECT id, serie, numero, fecha_emision::date AS fecha, cliente_numero_doc AS ruc, cliente_nombre AS nombre,
                       COALESCE(monto, 0) AS monto, comprobante_asociado
                FROM notas_credito
                WHERE COALESCE(UPPER(TRIM(estado)), 'BORRADOR') <> ALL(%s)
                  AND fecha_emision >= %s AND fecha_emision < %s::date + 1 {sql_n}
                ORDER BY fecha_emision, id""", (list(ANULADOS), desde, hasta) + par_n):
            total = abs(a_decimal(f['monto']))
            base = redondear(total / (1 + tasa_igv()))
            docs.append({'origen': 'ERP', 'id': f'nc{f["id"]}', 'tipo': '07', 'serie': f['serie'], 'numero': f['numero'],
                         'fecha': f['fecha'], 'ruc': f['ruc'], 'nombre': f['nombre'], 'base': -base,
                         'igv': -(total - base), 'no_gravado': CERO, 'total': -total, 'moneda': 'PEN',
                         'retenido': CERO, 'referencia': f['comprobante_asociado']})
    return docs, sin_tc


def compras_erp(eid, desde, hasta, manual=None):
    """Facturas de proveedores registradas en Compras. El ERP guarda solo el total: base = total / 1.18.
    24/10/2026: las facturas en dólares pasan a soles con su TC (el de la factura o el manual del mes)."""
    manual = manual or {}
    sql_e, par_e = _sql_emp('comprobantes_proveedor', 'empresa_id', eid)
    filas = db_query(f"""
        SELECT id, COALESCE(tipo, 'Factura') AS tipo, numero, fecha::date AS fecha, COALESCE(monto, 0) AS monto,
               ruc_proveedor AS ruc, proveedor_nombre AS nombre, orden_compra_id, orden_compra_numero, estado,
               to_jsonb(comprobantes_proveedor) ->> 'moneda' AS moneda, to_jsonb(comprobantes_proveedor) ->> 'tipo_cambio' AS tc_doc
        FROM comprobantes_proveedor
        WHERE COALESCE(UPPER(TRIM(estado)), '') <> ALL(%s)
          AND fecha >= %s AND fecha < %s::date + 1 {sql_e}
        ORDER BY fecha, id""", (list(ANULADOS), desde, hasta) + par_e)
    docs = []
    for f in filas:
        tipo = _tipo_cp(f['tipo'])
        tc, _falta = _tc(f.get('moneda'), manual, f.get('tc_doc'))
        total = redondear(a_decimal(f['monto']) * tc)
        serie, _, numero = (f['numero'] or '').rpartition('-')
        if tipo == '01' or tipo == '08' or tipo == '07':
            base = redondear(total / (1 + tasa_igv()))
            igv, ng = total - base, CERO
        else:   # boletas y recibos por honorarios: no dan crédito fiscal
            base, igv, ng = CERO, CERO, total
        signo = Decimal('-1') if tipo == '07' else Decimal('1')
        docs.append({'origen': 'ERP', 'id': f['id'], 'tipo': tipo, 'serie': serie or f['numero'], 'numero': numero,
                     'fecha': f['fecha'], 'ruc': f['ruc'], 'nombre': f['nombre'], 'base': base * signo,
                     'igv': igv * signo, 'no_gravado': ng * signo, 'total': total * signo,
                     'moneda': (f.get('moneda') or 'PEN').upper(), 'tipo_cambio': tc if tc != 1 else None,
                     'oc': f['orden_compra_numero'], 'oc_id': f['orden_compra_id'], 'estado': f['estado']})
    return docs


def importados(eid, periodo, libro=None):
    filas = db_query("""SELECT * FROM sunat_registros WHERE empresa_id = %s AND periodo = %s
                        AND (%s::text IS NULL OR libro = %s) ORDER BY fecha NULLS LAST, id""",
                     (str(eid), periodo, libro, libro))
    for f in filas:
        tc = a_decimal(f['tipo_cambio']) if f['tipo_cambio'] else Decimal('1')
        for k in ('base', 'igv', 'no_gravado', 'total'):
            f[k] = a_decimal(f[k]) * tc
    return filas


def combinar(libro, erp, imp):
    """Une lo del ERP con lo importado (SIRE/Excel) sin contar dos veces el mismo comprobante.
    Si está en los dos, manda el importado (es lo que SUNAT tiene) y se avisa si los montos no cuadran."""
    por_clave = {}
    for d in erp:
        d = dict(d)
        d['fuente'] = 'Solo ERP'
        por_clave.setdefault(_clave(libro, d['tipo'], d['serie'], d['numero'], d['ruc']), d)
    diferencias = []
    for r in imp:
        k = _clave(libro, r['tipo_cp'], r['serie'], r['numero'], r['ruc'])
        nuevo = {'origen': r['origen'], 'id': f"imp{r['id']}", 'tipo': r['tipo_cp'] or '01', 'serie': r['serie'],
                 'numero': r['numero'], 'fecha': r['fecha'], 'ruc': r['ruc'], 'nombre': r['nombre'],
                 'base': a_decimal(r['base']), 'igv': a_decimal(r['igv']), 'no_gravado': a_decimal(r['no_gravado']),
                 'total': a_decimal(r['total']), 'moneda': r['moneda'], 'retenido': CERO, 'lote': r['lote']}
        previo = por_clave.get(k)
        if previo is not None and previo['fuente'] in ('Solo ERP', 'ERP y archivo'):
            if abs(a_decimal(previo['igv']) - nuevo['igv']) > Decimal('1'):
                diferencias.append({'documento': f"{nuevo['serie']}-{nuevo['numero']}", 'nombre': nuevo['nombre'],
                                    'igv_erp': _n(previo['igv']), 'igv_archivo': _n(nuevo['igv'])})
            nuevo['retenido'] = previo.get('retenido', CERO)
            nuevo['fuente'] = 'ERP y archivo'
        else:
            nuevo['fuente'] = 'Solo archivo'
        por_clave[k] = nuevo
    return list(por_clave.values()), diferencias


def _sumar(docs):
    s = {'base': CERO, 'igv': CERO, 'no_gravado': CERO, 'n_facturas': 0, 'n_boletas': 0, 'n_nc': 0, 'n_nd': 0,
         'nc_base': CERO, 'nc_igv': CERO, 'nd_base': CERO, 'nd_igv': CERO, 'n_total': 0,
         'n_erp': 0, 'n_archivo': 0, 'n_ambos': 0, 'n_rxh': 0, 'rxh_total': CERO}
    for d in docs:
        s['n_total'] += 1
        s['base'] += d['base']
        s['igv'] += d['igv']
        s['no_gravado'] += d['no_gravado']
        t = d['tipo']
        if t == '01':
            s['n_facturas'] += 1
        elif t == '03':
            s['n_boletas'] += 1
        elif t == '07':
            s['n_nc'] += 1
            s['nc_base'] += d['base']
            s['nc_igv'] += d['igv']
        elif t == '08':
            s['n_nd'] += 1
            s['nd_base'] += d['base']
            s['nd_igv'] += d['igv']
        elif t == '02':
            s['n_rxh'] += 1
            s['rxh_total'] += d['total']
        f = d.get('fuente')
        s['n_erp' if f == 'Solo ERP' else 'n_archivo' if f == 'Solo archivo' else 'n_ambos'] += 1
    return s


# ============================================================
# PERIODO: datos que se escriben a mano (retenciones, saldos, planilla, cifras del contador)
# ============================================================
CAMPOS_MANUAL = ('saldo_favor_anterior', 'retenciones_mes', 'saldo_retenciones_anterior', 'percepciones_mes',
                 'saldo_percepciones_anterior', 'detracciones_saldo_bn', 'tipo_cambio', 'ventas_extra', 'compras_extra')
CAMPOS_PLANILLA = ('remuneraciones', 'base_essalud', 'base_onp', 'afp', 'renta_5ta', 'renta_4ta')


def leer_periodo(eid, periodo):
    f = db_query("SELECT * FROM sunat_periodos WHERE empresa_id = %s AND periodo = %s", (str(eid), periodo))
    if f:
        return f[0]
    return {'manual': {}, 'planilla': {}, 'contador': {}, 'contador_archivo': None, 'actualizado_por': None,
            'actualizado_en': None, 'contador_cargado_en': None}


def guardar_periodo(eid, periodo, campo, valores):
    with db_tx() as conn:
        cur = conn.cursor()
        cur.execute(f"""
            INSERT INTO sunat_periodos (empresa_id, periodo, {campo}, actualizado_por, actualizado_en)
            VALUES (%s, %s, %s, %s, NOW())
            ON CONFLICT (empresa_id, periodo) DO UPDATE
               SET {campo} = sunat_periodos.{campo} || EXCLUDED.{campo},
                   actualizado_por = EXCLUDED.actualizado_por, actualizado_en = NOW()
        """, (str(eid), periodo, Json(valores), _usuario()))


def leer_config(eid):
    f = db_query("SELECT * FROM sunat_config WHERE empresa_id = %s", (str(eid),))
    if f:
        c = dict(f[0])
        c['configurado'] = True
        return c
    return {'empresa_id': str(eid), 'regimen': 'RMT', 'tasa_renta': Decimal('1.0'), 'es_agente_retencion': False,
            'es_buen_contribuyente': False, 'notas': None, 'configurado': False}


def _num(manual, k):
    v = manual.get(k)
    return None if v in (None, '') else a_decimal(v)


# ============================================================
# CÁLCULO (mismo orden que la hoja del contador)
# ============================================================
def retenciones_sugeridas(eid, desde, hasta):
    """Retenciones de clientes según el ERP: facturas con retención cuyo PRIMER cobro cae en el mes.
    (El cliente retiene cuando paga. Lo exacto está en SUNAT ▸ Consulta de comprobantes de retención.)"""
    if not _cols('finanzas_pagos'):
        return CERO, 0
    # Si en Finanzas ▸ Cuentas por cobrar se anotó el comprobante de retención del cliente, manda ese monto (23/10/2026)
    from routes.finanzas import asegurar_tablas_finanzas
    asegurar_tablas_finanzas()
    sql_e, par_e = _sql_emp('comprobantes', 'c.empresa_id', eid)
    f = db_query(f"""
        SELECT COALESCE(SUM(COALESCE(sg.retencion_monto, c.monto_retenido, 0)), 0) AS t, COUNT(*) AS n FROM comprobantes c
        JOIN (SELECT documento_id, MIN(fecha) AS primero FROM finanzas_pagos
              WHERE tipo = 'COBRO' AND documento_tabla = 'comprobantes' AND COALESCE(estado, '') <> 'Anulado'
              GROUP BY documento_id) p ON p.documento_id = c.id
        LEFT JOIN finanzas_seguimiento_cobro sg ON sg.comprobante_id = c.id
        WHERE (COALESCE(c.tiene_retencion, false) OR COALESCE(sg.retencion_monto, 0) > 0)
          AND p.primero BETWEEN %s AND %s {sql_e}""", (desde, hasta) + par_e)
    return a_decimal(f[0]['t']), f[0]['n']


def liquidar(debito, credito, manual):
    """IGV a pagar con el mismo orden del PDT 621 y de la hoja del contador."""
    saldo_favor = abs(_num(manual, 'saldo_favor_anterior') or CERO)
    resultante = debito - credito
    tributo = resultante - saldo_favor
    perc = abs(_num(manual, 'percepciones_mes') or CERO) + abs(_num(manual, 'saldo_percepciones_anterior') or CERO)
    ret = abs(_num(manual, 'retenciones_mes') or CERO) + abs(_num(manual, 'saldo_retenciones_anterior') or CERO)
    if tributo <= 0:
        pagar, perc_usada, ret_usada = CERO, CERO, CERO
        favor_siguiente = -tributo
    else:
        favor_siguiente = CERO
        perc_usada = min(perc, tributo)
        ret_usada = min(ret, tributo - perc_usada)
        pagar = tributo - perc_usada - ret_usada
    return {'resultante': resultante, 'saldo_favor_anterior': saldo_favor, 'tributo': tributo,
            'percepciones': perc, 'retenciones': ret, 'percepciones_usadas': perc_usada, 'retenciones_usadas': ret_usada,
            'igv_pagar': pagar, 'saldo_favor_siguiente': favor_siguiente,
            'saldo_retenciones_siguiente': ret - ret_usada, 'saldo_percepciones_siguiente': perc - perc_usada,
            'importe_contador': tributo - perc - ret}


def calcular(e, anio, mes, corte=None):
    """Todo el resumen de un mes para una empresa. corte = fecha 'a hoy' (por defecto hoy o fin de mes si ya pasó)."""
    eid = e['id']
    periodo = _txt_periodo(anio, mes)
    desde, hasta = _rango(anio, mes)
    hoy = date.today()
    corte = min(corte or hoy, hasta)
    if corte < desde:
        corte = desde
    per = leer_periodo(eid, periodo)
    manual = dict(per['manual'] or {})
    planilla = dict(per['planilla'] or {})
    conf = leer_config(eid)

    v_erp, sin_tc = ventas_erp(eid, desde, corte, manual)
    c_erp = compras_erp(eid, desde, corte, manual)
    v_imp = [r for r in importados(eid, periodo, 'VENTAS') if not r['fecha'] or r['fecha'] <= corte]
    c_imp = importados(eid, periodo, 'COMPRAS') if corte >= hasta else \
        [r for r in importados(eid, periodo, 'COMPRAS') if not r['fecha'] or r['fecha'] <= corte or r['fecha'] < desde]
    ventas, dif_v = combinar('VENTAS', v_erp, v_imp)
    compras, dif_c = combinar('COMPRAS', c_erp, c_imp)
    sv, sc = _sumar(ventas), _sumar(compras)

    # Retenciones: lo escrito a mano manda; si no hay, la sugerencia del ERP
    ret_sug, ret_n = retenciones_sugeridas(eid, desde, hasta)
    man_calc = dict(manual)
    if _num(manual, 'retenciones_mes') is None:
        man_calc['retenciones_mes'] = ret_sug
    # Saldos del mes anterior: los que se escribieron; si no, lo que dejó el contador o el ERP el mes pasado
    ant = _saldos_mes_anterior(e, anio, mes)
    origen_saldos = {}
    for k_man, k_ant in (('saldo_favor_anterior', 'saldo_favor'), ('saldo_retenciones_anterior', 'saldo_retenciones'),
                         ('saldo_percepciones_anterior', 'saldo_percepciones')):
        if _num(manual, k_man) is None and ant.get(k_ant) is not None:
            man_calc[k_man] = ant[k_ant]
            origen_saldos[k_man] = ant['origen']
        else:
            origen_saldos[k_man] = 'escrito a mano' if _num(manual, k_man) is not None else 'sin dato'

    liq = liquidar(sv['igv'], sc['igv'], man_calc)

    # Renta mensual (pago a cuenta de 3.ª categoría)
    ingresos = sv['base']
    tasa = a_decimal(conf['tasa_renta'])
    renta = redondear(max(ingresos, CERO) * tasa / 100, 0)
    acumulado = _ingresos_anio(eid, anio, mes) + max(ingresos, CERO)

    # Planilla (PLAME): se escribe a mano (el ERP no tiene planilla). Si está vacía, se copia del mes anterior.
    plan_origen = 'escrito a mano'
    if not any(planilla.get(k) not in (None, '') for k in CAMPOS_PLANILLA):
        pa, pm = _mes_mas(anio, mes, -1)
        previo = leer_periodo(eid, _txt_periodo(pa, pm))
        planilla = dict(previo['planilla'] or {})
        plan_origen = 'copiado del mes anterior (revísalo)' if planilla else 'sin dato'
    essalud = redondear(a_decimal(planilla.get('base_essalud')) * Decimal('0.09'))
    onp = redondear(a_decimal(planilla.get('base_onp')) * Decimal('0.13'))
    r4 = a_decimal(planilla.get('renta_4ta'))
    if planilla.get('renta_4ta') in (None, '') and sc['rxh_total'] > 0:
        r4 = redondear(sum((d['total'] for d in compras if d['tipo'] == '02' and d['total'] > 1500), CERO) * Decimal('0.08'))

    det_bn = abs(_num(manual, 'detracciones_saldo_bn') or CERO)
    total_sunat = liq['igv_pagar'] + renta + essalud + onp + a_decimal(planilla.get('renta_5ta')) + r4
    vence = vencimiento_cronograma(e.get('ruc'), anio, mes)
    return {
        'empresa': {'id': eid, 'codigo': e['codigo'], 'razon_social': e.get('razon_social'), 'ruc': e.get('ruc')},
        'periodo': periodo, 'nombre_periodo': _nombre_periodo(anio, mes), 'desde': desde, 'hasta': hasta, 'corte': corte,
        'es_mes_actual': desde <= hoy <= hasta, 'vencimiento': vence,
        'config': conf, 'manual': manual, 'manual_calculo': man_calc, 'origen_saldos': origen_saldos,
        'planilla': planilla, 'planilla_origen': plan_origen,
        'ventas': ventas, 'compras': compras, 'sv': sv, 'sc': sc, 'sin_tipo_cambio': sin_tc,
        'diferencias': dif_v + dif_c, 'retenciones_sugeridas': ret_sug, 'retenciones_sugeridas_n': ret_n,
        'liq': liq, 'renta': {'regimen': conf['regimen'], 'tasa': tasa, 'base': max(ingresos, CERO), 'monto': renta,
                              'acumulado_anio': acumulado, 'tope_rmt': UIT_2026 * 300},
        'plame': {'essalud': essalud, 'onp': onp, 'renta_5ta': a_decimal(planilla.get('renta_5ta')), 'renta_4ta': r4,
                  'afp': a_decimal(planilla.get('afp')), 'remuneraciones': a_decimal(planilla.get('remuneraciones'))},
        'detracciones_bn': det_bn, 'total_sunat': total_sunat,
        'pagar_con_detracciones': min(det_bn, total_sunat), 'pagar_efectivo': max(total_sunat - det_bn, CERO),
        'contador': per['contador'] or {}, 'contador_archivo': per.get('contador_archivo'),
        'contador_cargado_en': per.get('contador_cargado_en'),
    }


def _ingresos_anio(eid, anio, mes):
    """Ventas (base) del año antes de este mes, para el tope de 300 UIT del Régimen MYPE."""
    if mes == 1:
        return CERO
    docs, _ = ventas_erp(eid, date(anio, 1, 1), _rango(*_mes_mas(anio, mes, -1))[1])
    total = sum((d['base'] for d in docs), CERO)
    f = db_query("""SELECT COALESCE(SUM(base), 0) AS t FROM sunat_registros
                    WHERE empresa_id = %s AND libro = 'VENTAS' AND periodo >= %s AND periodo < %s""",
                 (str(eid), f'{anio}-01', _txt_periodo(anio, mes)))
    return max(total, a_decimal(f[0]['t']) if f else CERO)


def _saldos_mes_anterior(e, anio, mes):
    """Saldos que pasan al mes: primero los del contador (si se cargó su hoja), si no los del cálculo del ERP."""
    pa, pm = _mes_mas(anio, mes, -1)
    per = leer_periodo(e['id'], _txt_periodo(pa, pm))
    c = per['contador'] or {}
    if c:
        imp = a_decimal(c.get('importe_igv'))
        trib = a_decimal(c.get('tributo'))
        return {'origen': 'hoja del contador del mes anterior',
                'saldo_favor': -trib if trib < 0 else CERO,
                'saldo_retenciones': -imp if (imp < 0 and trib >= 0) else
                (abs(a_decimal(c.get('retenciones_mes'))) + abs(a_decimal(c.get('saldo_retenciones_anterior'))) if trib < 0 else CERO),
                'saldo_percepciones': CERO}
    if per.get('actualizado_en') or db_query("SELECT 1 FROM sunat_registros WHERE empresa_id = %s AND periodo = %s LIMIT 1",
                                             (str(e['id']), _txt_periodo(pa, pm))):
        desde, hasta = _rango(pa, pm)
        manual = dict(per['manual'] or {})
        v, _ = ventas_erp(e['id'], desde, hasta, manual)
        c_ = compras_erp(e["id"], desde, hasta, manual)
        ventas, _ = combinar('VENTAS', v, importados(e['id'], _txt_periodo(pa, pm), 'VENTAS'))
        compras, _ = combinar('COMPRAS', c_, importados(e['id'], _txt_periodo(pa, pm), 'COMPRAS'))
        liq = liquidar(_sumar(ventas)['igv'], _sumar(compras)['igv'], manual)
        return {'origen': 'cálculo del ERP del mes anterior', 'saldo_favor': liq['saldo_favor_siguiente'],
                'saldo_retenciones': liq['saldo_retenciones_siguiente'],
                'saldo_percepciones': liq['saldo_percepciones_siguiente']}
    return {}


# ============================================================
# PROYECCIÓN "si hoy fuera fin de mes"
# ============================================================
def _pc_por_cotizacion_o_numero():
    pcs = {}
    for p in db_query("SELECT id, numero, cotizacion_numero, COALESCE(monto, 0) AS monto, factura, cliente FROM pedido_compra_pc"):
        for k in (p['numero'], p['cotizacion_numero']):
            if k:
                pcs.setdefault(k.strip().upper(), p)
    return pcs


def guias_sin_factura(eid):
    """Guías de venta ya emitidas (mercadería entregada) que no tienen factura. Por ley hay que facturarlas."""
    sql_e, par_e = _sql_emp('guias_remision', 'g.empresa_id', eid)
    filas = db_query(f"""
        SELECT g.id, g.serie, g.numero, g.fecha_emision::date AS fecha, g.destinatario_nombre AS cliente,
               g.ruc_destinatario AS ruc, g.documento_asociado
        FROM guias_remision g
        WHERE COALESCE(UPPER(g.motivo_traslado), 'VENTA') LIKE 'VENTA%%'
          AND COALESCE(UPPER(TRIM(g.estado_sunat)), 'BORRADOR') <> ALL(%s)
          AND COALESCE(TRIM(g.factura), '') = ''
          AND NOT EXISTS (SELECT 1 FROM comprobantes c
                          WHERE COALESCE(UPPER(TRIM(c.estado_sunat)), '') NOT IN ('ANULADO', 'ANULADA')
                            AND UPPER(TRIM(c.guia_vinculada)) IN (UPPER(g.serie || '-' || g.numero), UPPER(g.numero)))
          {sql_e}
        ORDER BY g.fecha_emision NULLS LAST, g.id""", (list(ANULADOS),) + par_e)
    pcs = _pc_por_cotizacion_o_numero()
    salida = []
    for g in filas:
        pc = pcs.get((g['documento_asociado'] or '').strip().upper())
        salida.append({'guia': f"{g['serie']}-{g['numero']}", 'fecha': g['fecha'], 'cliente': g['cliente'] or (pc or {}).get('cliente'),
                       'ruc': g['ruc'], 'pc': pc['numero'] if pc else '', 'base': a_decimal(pc['monto']) if pc and not pc['factura'] else None})
    return salida


def oc_sin_factura(eid):
    """Órdenes de compra aprobadas o recibidas que todavía no tienen factura del proveedor registrada."""
    sql_e, par_e = _sql_emp('ordenes_compra', 'oc.empresa_id', eid)
    filas = db_query(f"""
        SELECT oc.id, oc.numero_orden, oc.estado, COALESCE(oc.subtotal, 0) AS base, COALESCE(oc.igv, 0) AS igv,
               COALESCE(oc.total, 0) AS total, COALESCE(oc.fecha_creacion, oc.created_at)::date AS fecha,
               COALESCE(p.razon_social, '') AS proveedor, COALESCE(p.ruc, '') AS ruc, oc.tipo_compra
        FROM ordenes_compra oc LEFT JOIN proveedores p ON p.id = oc.proveedor_id
        WHERE oc.estado IN ('Aprobada', 'Recibida parcial', 'Recibida', 'Cerrada')
          AND NOT EXISTS (SELECT 1 FROM comprobantes_proveedor cp WHERE cp.orden_compra_id = oc.id
                          AND COALESCE(UPPER(cp.estado), '') NOT IN ('ANULADO', 'ANULADA'))
          {sql_e}
        ORDER BY oc.id""", par_e)
    return [{'oc': f['numero_orden'], 'id': f['id'], 'estado': f['estado'], 'fecha': f['fecha'], 'proveedor': f['proveedor'],
             'ruc': f['ruc'], 'base': a_decimal(f['base']), 'igv': a_decimal(f['igv']), 'total': a_decimal(f['total']),
             'recibida': f['estado'] in ('Recibida', 'Cerrada'), 'tipo_compra': f['tipo_compra']} for f in filas]


def _costo_productos(ids):
    costos = {}
    if not ids:
        return costos
    for f in db_query("SELECT id, COALESCE(costo_unitario, 0) AS c FROM productos WHERE id = ANY(%s)", (list(ids),)):
        costos[f['id']] = a_decimal(f['c'])
    try:
        for f in db_query("""SELECT DISTINCT ON (oci.producto_id) oci.producto_id, oci.precio_unitario AS c
                             FROM ordenes_compra_items oci JOIN ordenes_compra oc ON oc.id = oci.orden_compra_id
                             WHERE oci.producto_id = ANY(%s) AND COALESCE(oc.estado, '') NOT IN ('Anulada', 'Borrador')
                               AND COALESCE(oci.precio_unitario, 0) > 0
                             ORDER BY oci.producto_id, oc.id DESC""", (list(ids),)):
            costos[f['producto_id']] = a_decimal(f['c'])
    except Exception:
        pass
    return costos


def pc_por_comprar(eid):
    """Pedidos de clientes (PC) a los que todavía les falta comprar algo: ya están vendidos, se compran igual."""
    try:
        from routes.compras_flujo import _pedidos_por_comprar
        pedidos = _pedidos_por_comprar(sin_filtro=True, solo_por_comprar=True)
    except Exception as ex:
        print(f'⚠️ sunat: pedidos por comprar: {ex}')
        return []
    pedidos = [p for p in pedidos if p.get('empresa_id') in (None, str(eid))]
    costos = _costo_productos({i['producto_id'] for p in pedidos for i in p['items'] if i.get('producto_id')})
    salida = []
    for p in pedidos:
        costo, sin_costo, lineas = CERO, 0, []
        for i in p['items']:
            falta = a_decimal(i.get('comprar_pedido') if i.get('producto_id') else i.get('falta'))
            if falta <= 0:
                continue
            c = costos.get(i.get('producto_id'))
            if c and c > 0:
                costo += c * falta
            else:
                sin_costo += 1
            lineas.append(f"{i.get('descripcion') or i.get('producto') or i.get('codigo')}: {_n(falta):g}")
        salida.append({'pc': p['numero'], 'id': p['id'], 'cliente': p['cliente'], 'ruc': p['ruc'], 'fecha': p['fecha'],
                       'costo': costo, 'igv': redondear(costo * tasa_igv()), 'sin_costo': sin_costo,
                       'productos': lineas[:5]})
    salida.sort(key=lambda x: x['costo'], reverse=True)
    return salida


def despachos_pendientes(eid):
    """Despachos que aún no salen: si se entregan el 1.º del mes siguiente, su IGV pasa a ese mes (legal)."""
    sql_e, par_e = _sql_emp('despachos', 'd.empresa_id', eid)
    cita = 'd.fecha_cita' if _tiene('despachos', 'fecha_cita') else 'NULL::timestamp'
    filas = db_query(f"""
        SELECT d.id, d.numero, d.cliente, d.ruc, d.pc_numero, {cita}::date AS cita, COALESCE(pc.monto, 0) AS base,
               pc.factura
        FROM despachos d LEFT JOIN pedido_compra_pc pc ON pc.id = d.pc_id
        WHERE COALESCE(d.estado, '') NOT IN ('Despachado', 'Entregado', 'Anulado') {sql_e}
        ORDER BY d.id""", par_e)
    return [{'despacho': f['numero'], 'cliente': f['cliente'], 'ruc': f['ruc'], 'pc': f['pc_numero'], 'cita': f['cita'],
             'base': a_decimal(f['base']), 'facturado': bool(f['factura'])} for f in filas]


def proyeccion(e, r):
    """Piezas de "si hoy fuera fin de mes". La pantalla deja marcar cuáles sumar."""
    tasa = tasa_igv()
    partes = []
    if r['es_mes_actual']:
        guias = guias_sin_factura(e['id'])
        g_base = sum((g['base'] for g in guias if g['base']), CERO)
        partes.append({'clave': 'guias', 'tipo': 'venta', 'incluir': True, 'n': len(guias), 'base': g_base,
                       'igv': redondear(g_base * tasa),
                       'texto': 'Guías de venta ya entregadas sin factura (hay que facturarlas sí o sí)',
                       'sin_monto': sum(1 for g in guias if not g['base'])})
        desp = [d for d in despachos_pendientes(e['id']) if not d['facturado']]
        d_base = sum((d['base'] for d in desp), CERO)
        partes.append({'clave': 'despachos', 'tipo': 'venta', 'incluir': True, 'n': len(desp), 'base': d_base,
                       'igv': redondear(d_base * tasa),
                       'texto': 'Despachos por salir este mes (si salen antes del 31)', 'sin_monto': 0})
        ocs = [o for o in oc_sin_factura(e['id'])]
        o_base = sum((o['base'] for o in ocs), CERO)
        partes.append({'clave': 'oc', 'tipo': 'compra', 'incluir': True, 'n': len(ocs), 'base': o_base,
                       'igv': sum((o['igv'] for o in ocs), CERO),
                       'texto': 'Órdenes de compra aprobadas o recibidas sin factura del proveedor', 'sin_monto': 0})
        pcs = pc_por_comprar(e['id'])
        p_base = sum((p['costo'] for p in pcs), CERO)
        partes.append({'clave': 'pc', 'tipo': 'compra', 'incluir': False, 'n': len(pcs), 'base': p_base,
                       'igv': redondear(p_base * tasa),
                       'texto': 'Pedidos (PC) que faltan comprar, si se compran y facturan antes del 31',
                       'sin_monto': sum(p['sin_costo'] for p in pcs)})
    for clave, tipo, texto in (('ventas_extra', 'venta', 'Otras ventas que esperas facturar (escrito a mano)'),
                               ('compras_extra', 'compra', 'Compra para stock que estás pensando (escrito a mano)')):
        base = abs(_num(r['manual'], clave) or CERO)
        partes.append({'clave': clave, 'tipo': tipo, 'incluir': True, 'n': 0, 'base': base,
                       'igv': redondear(base * tasa), 'texto': texto, 'sin_monto': 0})
    return partes


def proyectar(r, partes):
    deb = r['sv']['igv'] + sum((p['igv'] for p in partes if p['incluir'] and p['tipo'] == 'venta'), CERO)
    cre = r['sc']['igv'] + sum((p['igv'] for p in partes if p['incluir'] and p['tipo'] == 'compra'), CERO)
    liq = liquidar(deb, cre, r['manual_calculo'])
    base_v = r['sv']['base'] + sum((p['base'] for p in partes if p['incluir'] and p['tipo'] == 'venta'), CERO)
    return {'debito': deb, 'credito': cre, 'liq': liq,
            'renta': redondear(max(base_v, CERO) * a_decimal(r['config']['tasa_renta']) / 100, 0)}


# ============================================================
# PLAN DE ACCIÓN (como auditor)
# ============================================================
def plan_accion(e, r, proy):
    """Lista ordenada de lo que conviene hacer, con montos. Solo sugerencias legales."""
    acc = []
    tasa = tasa_igv()
    igv_proy = proy['liq']['igv_pagar'] if proy else r['liq']['igv_pagar']

    # 1. ¿Están completos los datos?
    sv, sc = r['sv'], r['sc']
    if sv['n_archivo'] + sv['n_ambos'] == 0 or sc['n_archivo'] + sc['n_ambos'] == 0:
        falta = [n for n, s in (('ventas', sv), ('compras', sc)) if s['n_archivo'] + s['n_ambos'] == 0]
        acc.append({'nivel': 'rojo', 'grupo': 'Datos', 'titulo': 'Completa las facturas que están solo en Excel',
                    'detalle': (f"El ERP tiene {sv['n_total']} comprobantes de venta y {sc['n_total']} de compra en "
                                f"{r['nombre_periodo']}. Como las chicas todavía registran en Excel, puede faltar. "
                                f"Importa la propuesta del SIRE (RVIE y RCE) o el Excel en la pestaña "
                                f"\"Completar datos\". Falta importar: {' y '.join(falta)}."),
                    'ir': 'completar'})
    if r['diferencias']:
        acc.append({'nivel': 'rojo', 'grupo': 'Datos', 'titulo': f"{len(r['diferencias'])} comprobantes con IGV distinto entre el ERP y el archivo",
                    'detalle': 'Revisa cuál está bien: ' + '; '.join(f"{d['documento']} (ERP S/ {d['igv_erp']:,.2f}, archivo S/ {d['igv_archivo']:,.2f})" for d in r['diferencias'][:5]),
                    'ir': 'detalle'})
    if r['sin_tipo_cambio']:
        acc.append({'nivel': 'amarillo', 'grupo': 'Datos', 'titulo': 'Hay comprobantes en dólares sin tipo de cambio',
                    'detalle': 'Escribe el tipo de cambio del mes en "Datos del mes"; mientras tanto se calcularon como soles.',
                    'ir': 'resumen'})

    # 2. Lo obligatorio (riesgo de multa o de perder crédito fiscal)
    if r['es_mes_actual']:
        guias = guias_sin_factura(e['id'])
        if guias:
            base = sum((g['base'] for g in guias if g['base']), CERO)
            acc.append({'nivel': 'rojo', 'grupo': 'Obligatorio', 'titulo': f'{len(guias)} guías de venta sin factura',
                        'detalle': ('La mercadería ya se entregó: la factura se debe emitir igual (el IGV nace con la entrega). '
                                    'No conviene "dejar de facturar" algo ya entregado: es una infracción y el IGV es el mismo.'),
                        'monto': base, 'igv': redondear(base * tasa), 'ir': 'ventas',
                        'items': [f"{g['guia']} · {g['cliente'] or ''} · {g['fecha'].strftime('%d/%m/%Y') if g['fecha'] else ''}"
                                  + (f" · S/ {_n(g['base']):,.2f}" if g['base'] else '') for g in guias[:8]]})
        ocs = oc_sin_factura(e['id'])
        recibidas = [o for o in ocs if o['recibida']]
        if recibidas:
            igv = sum((o['igv'] for o in recibidas), CERO)
            acc.append({'nivel': 'rojo', 'grupo': 'Obligatorio', 'titulo': f'{len(recibidas)} compras recibidas sin factura registrada',
                        'detalle': ('Ya llegó la mercadería pero no está la factura del proveedor en Compras. Pídela y '
                                    'regístrala: sin factura no hay crédito fiscal.'),
                        'monto': sum((o['base'] for o in recibidas), CERO), 'igv': igv, 'ir': 'compras',
                        'items': [f"{o['oc']} · {o['proveedor']} · IGV S/ {_n(o['igv']):,.2f}" for o in recibidas[:8]]})
    bancar = _compras_sin_banco(e['id'], r['desde'], r['hasta'])
    if bancar:
        acc.append({'nivel': 'rojo', 'grupo': 'Obligatorio', 'titulo': f'{len(bancar)} compras de S/ 2,000 o más pagadas en efectivo',
                    'detalle': ('Desde S/ 2,000 (o US$ 500) el pago debe ir por banco (bancarización). Pagado en efectivo, '
                                'SUNAT no acepta el crédito fiscal ni el gasto.'),
                    'items': [f"{b['numero']} · {b['proveedor']} · S/ {_n(b['total']):,.2f}" for b in bancar[:8]], 'ir': 'compras'})

    # 3. Si hay IGV por pagar: lo que se puede hacer legalmente
    if igv_proy > 0 and r['es_mes_actual']:
        pcs = pc_por_comprar(e['id'])
        if pcs:
            igv_pc = sum((p['igv'] for p in pcs), CERO)
            acc.append({'nivel': 'amarillo', 'grupo': 'Para pagar menos IGV',
                        'titulo': f'Compra antes del 31 lo que ya vendiste: {len(pcs)} pedidos por comprar',
                        'detalle': (f"Estos pedidos ya están vendidos y hay que comprarlos igual. Si el proveedor te factura "
                                    f"en {MESES[r['hasta'].month - 1].lower()}, su IGV (aprox. S/ {_n(igv_pc):,.2f}) baja tu pago "
                                    f"de este mes. Genera las órdenes en Compras ▸ Pedidos por comprar y pide que la factura "
                                    f"salga con fecha de este mes."),
                        'monto': sum((p['costo'] for p in pcs), CERO), 'igv': igv_pc, 'ir': 'oc',
                        'items': [f"{p['pc']} · {p['cliente']} · costo aprox. S/ {_n(p['costo']):,.2f}"
                                  + (f" ({p['sin_costo']} producto(s) sin costo conocido)" if p['sin_costo'] else '') for p in pcs[:8]]})
        aprob = [o for o in oc_sin_factura(e['id']) if not o['recibida']]
        if aprob:
            igv = sum((o['igv'] for o in aprob), CERO)
            acc.append({'nivel': 'amarillo', 'grupo': 'Para pagar menos IGV',
                        'titulo': f'{len(aprob)} órdenes de compra aprobadas: pide la factura este mes',
                        'detalle': (f"Si el proveedor entrega y factura antes del 31, sumas S/ {_n(igv):,.2f} de crédito fiscal. "
                                    f"Recuerda registrarla en Compras ▸ Facturas de proveedor."),
                        'monto': sum((o['base'] for o in aprob), CERO), 'igv': igv, 'ir': 'compras',
                        'items': [f"{o['oc']} · {o['proveedor']} · IGV S/ {_n(o['igv']):,.2f}" for o in aprob[:8]]})
        desp = [d for d in despachos_pendientes(e['id']) if not d['facturado'] and d['base'] > 0]
        if desp:
            acc.append({'nivel': 'info', 'grupo': 'Para pagar menos IGV',
                        'titulo': 'Entregas que podrías programar para el 1 del próximo mes',
                        'detalle': ('Solo si el cliente está de acuerdo y la mercadería aún NO sale: el IGV nace con la entrega '
                                    'o con la factura (lo que pase primero). Si se entrega y factura el 1, el IGV pasa al mes '
                                    'siguiente. Si el cliente exige factura antes, no se puede.'),
                        'monto': sum((d['base'] for d in desp), CERO),
                        'igv': redondear(sum((d['base'] for d in desp), CERO) * tasa), 'ir': 'ventas',
                        'items': [f"{d['despacho']} · {d['cliente']} · S/ {_n(d['base']):,.2f}" for d in desp[:8]]})
        stock_base = redondear(igv_proy / tasa)
        acc.append({'nivel': 'info', 'grupo': 'Para pagar menos IGV',
                    'titulo': f'Compra para stock: solo si la vas a vender pronto',
                    'detalle': (f"Para no pagar los S/ {_n(igv_proy):,.2f} de IGV tendrías que comprar S/ {_n(stock_base * (1 + tasa)):,.2f} "
                                f"(con IGV). Es decir, sacas S/ {_n(stock_base * (1 + tasa)):,.2f} de caja para ahorrar "
                                f"S/ {_n(igv_proy):,.2f} este mes; el IGV no se pierde, se recupera cuando vendes. Conviene solo con "
                                f"productos de alta rotación y si tienes liquidez. Usa Compras ▸ Compra directa a stock con "
                                f"motivo \"Crédito fiscal / IGV\"."),
                    'monto': stock_base, 'igv': igv_proy, 'ir': 'oc'})
    elif r['liq']['igv_pagar'] <= 0 and (r['liq']['saldo_retenciones_siguiente'] > 0 or r['liq']['saldo_favor_siguiente'] > 0):
        acc.append({'nivel': 'verde', 'grupo': 'Para pagar menos IGV', 'titulo': 'Este mes no pagas IGV',
                    'detalle': 'No compres para stock "por el IGV": ya tienes saldo a favor. Compra solo lo que necesitas.'})

    # 4. Retenciones acumuladas: se pueden pedir de vuelta
    ret_saldo = r['liq']['saldo_retenciones_siguiente']
    if ret_saldo >= 500:
        acc.append({'nivel': 'amarillo', 'grupo': 'Dinero a favor',
                    'titulo': f"Tienes S/ {_n(ret_saldo):,.2f} de retenciones sin usar",
                    'detalle': ('Si este saldo se mantiene 3 meses seguidos, puedes pedir la devolución a SUNAT '
                                '(Formulario 1649 o SUNAT Virtual). SUNAT responde en hasta 45 días hábiles. '
                                'Pregúntale al contador si ya cumple el plazo: es plata de la empresa.')})

    # 5. Intercompany (ventas entre KCF y AGD)
    otras = _otras_rucs(e)
    if otras:
        v_ic = [d for d in r['ventas'] if (d['ruc'] or '').strip() in otras]
        c_ic = [d for d in r['compras'] if (d['ruc'] or '').strip() in otras]
        if v_ic or c_ic:
            acc.append({'nivel': 'info', 'grupo': 'Entre empresas del grupo',
                        'titulo': 'Ventas entre empresas del grupo (intercompany)',
                        'detalle': (f"Este mes: ventas a {', '.join(sorted({otras[d['ruc'].strip()] for d in v_ic})) or '—'} "
                                    f"S/ {_n(sum((d['base'] for d in v_ic), CERO)):,.2f}; compras a "
                                    f"{', '.join(sorted({otras[d['ruc'].strip()] for d in c_ic})) or '—'} "
                                    f"S/ {_n(sum((d['base'] for d in c_ic), CERO)):,.2f} (sin IGV). El IGV del grupo es el mismo: "
                                    f"solo pasa de una empresa a la otra. El precio debe ser de mercado (como a un tercero) y con "
                                    f"factura, guía y pago por banco. Ojo: si una empresa acumula retenciones sin usar y la otra "
                                    f"paga IGV en efectivo, conversa con el contador cómo ordenar las ventas para usar esas "
                                    f"retenciones (siempre a valor de mercado).")})

    # 6. Reglas de oro (siempre)
    acc.append({'nivel': 'info', 'grupo': 'Reglas de oro', 'titulo': 'Para que el crédito fiscal valga',
                'items': ['Factura a nombre de la empresa, con IGV separado y proveedor habido.',
                          'Registrada en el Registro de Compras (SIRE) a más tardar dentro de los 12 meses.',
                          'Pagos desde S/ 2,000 por banco (transferencia o depósito), nunca en efectivo.',
                          'Si la compra tiene detracción, depositarla antes de usar el crédito.',
                          'Las boletas y los recibos por honorarios no dan crédito fiscal.',
                          'Si cobras un adelanto, el IGV nace ese día (aunque entregues después).']})
    return acc


def _compras_sin_banco(eid, desde, hasta):
    if not _cols('finanzas_pagos'):
        return []
    sql_e, par_e = _sql_emp('comprobantes_proveedor', 'cp.empresa_id', eid)
    return [dict(f) for f in db_query(f"""
        SELECT cp.numero, cp.proveedor_nombre AS proveedor, cp.monto AS total FROM comprobantes_proveedor cp
        WHERE COALESCE(cp.monto, 0) >= %s AND cp.fecha >= %s AND cp.fecha < %s::date + 1 {sql_e}
          AND EXISTS (SELECT 1 FROM finanzas_pagos p WHERE p.tipo = 'PAGO' AND p.documento_tabla = 'comprobantes_proveedor'
                      AND p.documento_id = cp.id AND COALESCE(p.estado, '') <> 'Anulado'
                      AND LOWER(COALESCE(p.metodo, '')) = 'efectivo')""", (BANCARIZACION, desde, hasta) + par_e)]


# ============================================================
# HOJA DEL CONTADOR (mismo formato) e importación de su Excel
# ============================================================
FILAS_PRELIMINAR = (
    # clave, texto, sección
    ('ventas_base', 'VENTAS NETAS LOCALES (base)', 'igv'),
    ('ventas_igv', 'IGV DE VENTAS', 'igv'),
    ('nc_base', '   (dentro de ventas: notas de crédito, base)', 'igv'),
    ('compras_base', 'COMPRAS LOCALES DESTINADAS A VENTAS GRAVADAS (base)', 'igv'),
    ('compras_igv', 'IGV DE COMPRAS (crédito fiscal)', 'igv'),
    ('compras_ng', 'COMPRAS NO GRAVADAS INTERNAS', 'igv'),
    ('resultante', 'IMPUESTO RESULTANTE O SALDO A FAVOR', 'igv'),
    ('saldo_favor_anterior', 'SALDO A FAVOR DEL PERIODO ANTERIOR', 'igv'),
    ('tributo', 'TRIBUTO A PAGAR O SALDO A FAVOR', 'igv'),
    ('percepciones', 'PERCEPCIONES (DEL PERIODO Y ANTERIORES)', 'igv'),
    ('retenciones_mes', 'RETENCIONES DECLARADAS EN EL PERIODO', 'igv'),
    ('saldo_retenciones_anterior', 'SALDO DE RETENCIONES DE PERIODOS ANTERIORES', 'igv'),
    ('importe_igv', 'IMPORTE A PAGAR O SALDO A FAVOR', 'igv'),
    ('renta', 'RENTA 3.ª CATEGORÍA: IMPUESTO DEL MES', 'renta'),
    ('essalud', 'PLAME: ESSALUD', 'plame'),
    ('onp', 'PLAME: SNP / ONP', 'plame'),
    ('total', 'TOTAL TRIBUTOS Y CONTRIBUCIONES POR PAGAR', 'total'),
)


def valores_preliminar(r):
    liq, sv, sc = r['liq'], r['sv'], r['sc']
    m = r['manual_calculo']
    v = {
        'ventas_base': sv['base'], 'ventas_igv': sv['igv'], 'nc_base': sv['nc_base'], 'nc_igv': sv['nc_igv'],
        'compras_base': sc['base'], 'compras_igv': -sc['igv'], 'compras_ng': sc['no_gravado'],
        'resultante': liq['resultante'], 'saldo_favor_anterior': -liq['saldo_favor_anterior'], 'tributo': liq['tributo'],
        'percepciones': -liq['percepciones'], 'retenciones_mes': -abs(a_decimal(m.get('retenciones_mes'))),
        'saldo_retenciones_anterior': -abs(a_decimal(m.get('saldo_retenciones_anterior'))),
        'importe_igv': liq['importe_contador'], 'renta': r['renta']['monto'], 'essalud': r['plame']['essalud'],
        'onp': r['plame']['onp'],
    }
    v['total'] = max(liq['igv_pagar'], CERO) + v['renta'] + v['essalud'] + v['onp']
    return v


def comparar_contador(r):
    pre = valores_preliminar(r)
    c = r['contador'] or {}
    filas = []
    for clave, texto, seccion in FILAS_PRELIMINAR:
        erp = pre.get(clave)
        cont = c.get(clave)
        dif = None if cont in (None, '') else a_decimal(cont) - a_decimal(erp)
        filas.append({'clave': clave, 'texto': texto, 'seccion': seccion, 'erp': _n(erp),
                      'contador': None if cont in (None, '') else _n(cont), 'diferencia': None if dif is None else _n(dif),
                      'alerta': dif is not None and abs(dif) > 1})
    obs = _observaciones_contador(c) if c else []
    return filas, obs


def _observaciones_contador(c, previo=None):
    """Cosas que vale la pena preguntarle al contador sobre su hoja."""
    obs = []
    if previo:
        trib_p, imp_p = a_decimal(previo.get('tributo')), a_decimal(previo.get('importe_igv'))
        if trib_p >= 0:
            esperado = abs(imp_p) if imp_p < 0 else CERO
        else:
            esperado = abs(a_decimal(previo.get('retenciones_mes'))) + abs(a_decimal(previo.get('saldo_retenciones_anterior')))
        actual = abs(a_decimal(c.get('saldo_retenciones_anterior')))
        if abs(esperado - actual) > 1:
            obs.append(f"El saldo de retenciones que viene del mes anterior debería ser S/ {_n(esperado):,.2f} y la hoja "
                       f"pone S/ {_n(actual):,.2f} (diferencia S/ {_n(actual - esperado):,.2f}). Pídele que lo revise.")
    onp = a_decimal(c.get('onp'))
    total = a_decimal(c.get('total'))
    suma = a_decimal(c.get('essalud')) + max(a_decimal(c.get('importe_igv')), CERO) + a_decimal(c.get('renta')) + onp
    if onp > 0 and total and abs(total - suma) > 1 and abs(total - (suma - onp)) <= 1:
        obs.append(f"El total de la hoja (S/ {_n(total):,.2f}) no suma la ONP (S/ {_n(onp):,.2f}). La ONP retenida a los "
                   f"trabajadores también se paga a SUNAT en el PLAME: pregúntale si la está pagando aparte.")
    if a_decimal(c.get('saldo_retenciones_siguiente')) >= 500 or (a_decimal(c.get('importe_igv')) < -500):
        obs.append(f"Hay S/ {_n(abs(a_decimal(c.get('importe_igv')))):,.2f} de retenciones sin usar. Pregúntale si ya se puede "
                   f"pedir la devolución (Formulario 1649: saldo mantenido 3 meses seguidos).")
    if c.get('vencimiento_hoja') and c.get('vencimiento_sunat') and c['vencimiento_hoja'] != c['vencimiento_sunat']:
        obs.append(f"La hoja dice que vence el {c['vencimiento_hoja']}, pero el cronograma SUNAT 2026 dice "
                   f"{c['vencimiento_sunat']}. Su tabla de vencimientos es del 2025: pídele que la actualice.")
    return obs


_HOJAS = {'ENE': 1, 'FEB': 2, 'MAR': 3, 'ABR': 4, 'MAY': 5, 'JUN': 6, 'JUL': 7, 'AGO': 8, 'SET': 9, 'SEP': 9,
          'OCT': 10, 'NOV': 11, 'DIC': 12}


def leer_hoja_contador(contenido, anio, mes):
    """Lee la hoja del mes del Excel del contador (formato 'LIQUIDACION DE IMPUESTOS Y CONTRIBUCIONES')."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
        wf = openpyxl.load_workbook(io.BytesIO(contenido), data_only=False)
    except Exception:
        raise ErrorSunat('No se pudo abrir el Excel. Súbelo tal como lo manda el contador (.xlsx).')
    hoja = None
    for ws in wb.worksheets:
        m = re.match(r'^\s*([A-Za-z]{3})\w*\s+(\d{4})\s*$', ws.title or '')
        if m and _HOJAS.get(m.group(1).upper()) == mes and int(m.group(2)) == anio:
            hoja = ws
            break
        titulo = ' '.join(str(c.value) for row in ws.iter_rows(max_row=6) for c in row if c.value).upper()
        nombres = ('SETIEMBRE', 'SEPTIEMBRE') if mes == 9 else (MESES[mes - 1].upper(),)
        if any(re.search(rf'LIQUIDACION.*{n}\s+{anio}', titulo) for n in nombres):
            hoja = ws
            break
    if hoja is None:
        raise ErrorSunat(f'No encontré la hoja de {_nombre_periodo(anio, mes)} en el Excel (por ejemplo "SET 2026").')
    formulas = wf[hoja.title]

    def valor(col, fila):
        v = hoja[f'{col}{fila}'].value
        return a_decimal(v) if isinstance(v, (int, float, Decimal)) else None

    def buscar(*textos, desde=1):
        for row in hoja.iter_rows(min_row=desde):
            for c in row:
                if c.column <= 5 and isinstance(c.value, str) and all(t in _norm(c.value) for t in textos):
                    return c.row
        return None

    d = {}
    f = buscar('ventas netas locales')
    if f:
        d['ventas_base'], d['ventas_igv'] = valor('F', f), valor('H', f)
    f = buscar('locales destinadas a ventas gravadas')
    if f:
        d['compras_base'], d['compras_igv'] = valor('F', f), valor('H', f)
    f = buscar('no gravadas internas')
    if f:
        d['compras_ng'] = valor('F', f)
    for clave, textos in (('resultante', ('impuesto resultante',)), ('saldo_favor_anterior', ('saldo favor del periodo anterior',)),
                          ('tributo', ('tributo a pagar',)), ('retenciones_mes', ('retenciones declaradas',)),
                          ('saldo_retenciones_anterior', ('saldo de retenciones de periodos',))):
        f = buscar(*textos)
        if f:
            d[clave] = valor('H', f)
    f = buscar('percepciones declaradas')
    if f:
        d['percepciones'] = (valor('H', f) or CERO) + (valor('H', f + 1) or CERO)
    f = buscar('importe a pagar o saldo')
    if f:
        d['importe_igv'] = valor('H', f)
    f = buscar('impuesto del mes')
    if f:
        d['renta'] = valor('H', f)
    f = buscar('essalud')
    if f:
        d['essalud'] = valor('H', f)
        d['base_essalud'] = valor('F', f)
    f = buscar('snp')
    if f:
        d['onp'] = valor('H', f)
        d['base_onp'] = valor('F', f)
    f = buscar('total tributos')
    if f:
        d['total'] = valor('H', f)
    # Vencimiento que escribió la hoja (para avisar si su tabla está desactualizada)
    f = buscar('vencimiento')
    if f:
        for col in 'EFGH':
            v = hoja[f'{col}{f}'].value
            if v:
                d['vencimiento_hoja'] = str(v).strip()
                break
    if not any(v is not None for k, v in d.items() if k != 'vencimiento_hoja'):
        raise ErrorSunat('La hoja no tiene el formato del contador (no encontré "VENTAS NETAS LOCALES").')
    salida = {k: (str(redondear(v)) if isinstance(v, Decimal) else v) for k, v in d.items() if v is not None}
    return hoja.title, salida


def excel_preliminar(r):
    """Excel con la misma forma de la hoja del contador, para comparar lado a lado."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{MESES[int(r['periodo'][5:]) - 1][:3].upper()} {r['periodo'][:4]}"
    neg = Font(bold=True)
    fino = Side(style='thin', color='CBD5E1')
    pre = valores_preliminar(r)
    ws['A1'] = 'PRELIMINAR DEL ERP (no es la declaración: la declaración la hace el contador)'
    ws['A1'].font = Font(bold=True, color='B91C1C')
    ws['A2'] = r['empresa']['razon_social'] or r['empresa']['codigo']
    ws['A2'].font = neg
    ws['A3'] = 'RUC  :'
    ws['B3'] = r['empresa']['ruc']
    ws['A5'] = f"LIQUIDACION DE IMPUESTOS Y CONTRIBUCIONES {r['nombre_periodo'].upper()}"
    ws['A5'].font = Font(bold=True, size=12)
    ws.merge_cells('A5:I5')
    ws['A5'].alignment = Alignment(horizontal='center')
    ws['D7'] = 'VENCIMIENTO'
    ws['E7'] = r['vencimiento'].strftime('%d/%m/%Y') if r['vencimiento'] else 'Revisar cronograma'
    ws['A9'], ws['F9'], ws['H9'] = 'PDT 0621 IGV- RENTA', 'BASE IMPONIBLE', 'IMPUESTO'
    for c in ('A9', 'F9', 'H9'):
        ws[c].font = neg
    filas = [
        (12, 'B', 'VENTAS NETAS LOCALES', pre['ventas_base'], 0.18, pre['ventas_igv']),
        (13, 'C', '(incluye notas de crédito)', pre['nc_base'], None, pre['nc_igv']),
        (19, 'B', 'COMPRAS  LOCALES DESTINADAS A VENTAS GRAVADAS', pre['compras_base'], 0.18, pre['compras_igv']),
        (23, 'C', 'NO GRAVADAS INTERNAS', pre['compras_ng'], None, None),
        (27, 'C', 'IMPUESTO RESULTANTE O SALDO A FAVOR', None, None, pre['resultante']),
        (28, 'C', 'SALDO FAVOR DEL PERIODO ANTERIOR', None, None, pre['saldo_favor_anterior']),
        (29, 'C', 'TRIBUTO A PAGAR O  SALDO A FAVOR', None, None, pre['tributo']),
        (31, 'C', 'PERCEPCIONES (DEL PERIODO Y ANTERIORES)', None, None, pre['percepciones']),
        (34, 'C', 'RETENCIONES DECLARADAS EN EL PERIODO', None, None, pre['retenciones_mes']),
        (35, 'C', 'SALDO DE RETENCIONES DE PERIODOS ANTERIORES', None, None, pre['saldo_retenciones_anterior']),
        (38, 'C', 'IMPORTE A PAGAR O SALDO  A FAVOR', None, None, pre['importe_igv']),
        (44, 'C', f"RENTA 3ERA. CATEG. ({REGIMENES.get(r['renta']['regimen'], {}).get('nombre', r['renta']['regimen'])})",
         r['renta']['base'], float(r['renta']['tasa']) / 100, pre['renta']),
        (70, 'B', 'ESSALUD', a_decimal(r['planilla'].get('base_essalud')), 0.09, pre['essalud']),
        (71, 'B', 'SNP', a_decimal(r['planilla'].get('base_onp')), 0.13, pre['onp']),
        (89, 'B', 'TOTAL TRIBUTOS Y CONTRIBUCIONES POR PAGAR', None, None, pre['total']),
    ]
    ws['A10'] = 'I. G. V.'
    ws['A40'] = 'RENTA 3ERA. CATEG.'
    ws['A68'] = 'PDT 0601 PLAME'
    for c in ('A10', 'A40', 'A68'):
        ws[c].font = neg
    for fila, col, texto, base, tasa, imp in filas:
        ws[f'{col}{fila}'] = texto
        if base is not None:
            ws[f'F{fila}'] = float(redondear(base))
            ws[f'F{fila}'].number_format = '#,##0.00'
        if tasa is not None:
            ws[f'G{fila}'] = tasa
            ws[f'G{fila}'].number_format = '0.0%'
        if imp is not None:
            ws[f'H{fila}'] = float(redondear(imp))
            ws[f'H{fila}'].number_format = '#,##0.00'
            ws[f'H{fila}'].border = Border(bottom=fino)
    ws['B89'].font = neg
    ws['H89'].font = neg
    ws['H89'].fill = PatternFill('solid', fgColor='FEF3C7')
    ws['B92'] = f"Generado por el ERP el {_ahora_lima().strftime('%d/%m/%Y %H:%M')} (hora de Lima) por {_usuario()}."
    ws['B93'] = (f"Ventas: {r['sv']['n_total']} comprobantes ({r['sv']['n_erp']} solo ERP, {r['sv']['n_archivo']} solo archivo, "
                 f"{r['sv']['n_ambos']} en ambos). Compras: {r['sc']['n_total']}.")
    for col, ancho in zip('ABCDEFGH', (16, 14, 44, 12, 14, 16, 8, 16)):
        ws.column_dimensions[col].width = ancho
    # Detalle de comprobantes
    for libro, docs in (('Ventas', r['ventas']), ('Compras', r['compras'])):
        d = wb.create_sheet(f'Detalle {libro.lower()}')
        d.append([f'{"Cliente" if libro == "Ventas" else "Proveedor"}', 'RUC / DNI', 'Tipo', 'N° comprobante', 'Fecha',
                  'Base (S/)', 'IGV (S/)', 'No gravado (S/)', 'Total (S/)', 'Fuente'])
        for c in d[1]:
            c.font = neg
        for x in sorted(docs, key=lambda x: (x['fecha'] or date.min, str(x['numero']))):
            d.append([x['nombre'], x['ruc'], TIPOS_CP.get(x['tipo'], x['tipo']), f"{x['serie']}-{x['numero']}",
                      x['fecha'].strftime('%d/%m/%Y') if x['fecha'] else '', float(redondear(x['base'])),
                      float(redondear(x['igv'])), float(redondear(x['no_gravado'])), float(redondear(x['total'])), x['fuente']])
        for col, ancho in zip('ABCDEFGHIJ', (40, 14, 16, 18, 12, 14, 12, 14, 14, 16)):
            d.column_dimensions[col].width = ancho
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ============================================================
# IMPORTAR SIRE (RVIE / RCE) o EXCEL
# ============================================================
SINONIMOS = (
    ('fecha', ('fecha de emision', 'fecha emision', 'fec. emision', 'fecha emi', 'f. emision', 'fecha de emisi', 'fecha')),
    ('tipo', ('tipo cp/doc', 'tipo cp', 'tipo de comprobante', 'tipo comprobante', 'tipo de cp', 'tipo doc.', 'tipo')),
    ('serie', ('serie del cdp', 'serie cdp', 'serie')),
    ('numero', ('nro cp o doc. nro inicial', 'nro cp', 'nro. cp', 'numero cp', 'numero de comprobante', 'nro comprobante',
                'numero', 'nro.', 'nro', 'n ')),
    ('comprobante', ('comprobante', 'documento', 'n factura', 'factura')),
    ('ruc', ('nro doc identidad', 'nro doc. identidad', 'numero de documento de identidad', 'doc identidad', 'ruc / dni',
             'ruc/dni', 'ruc dni', 'ruc', 'dni')),
    ('nombre', ('apellidos nombres/ razon social', 'apellidos y nombres', 'razon social', 'cliente', 'proveedor',
                'nombre')),
    ('base', ('bi gravada dg', 'bi gravada', 'base imponible', 'valor venta', 'valor de venta', 'subtotal', 'base')),
    ('igv', ('igv / ipm dg', 'igv / ipm', 'igv dg', 'igv')),
    ('no_gravado', ('valor adq. ng', 'mto exonerado', 'exonerado', 'inafecto', 'no gravad')),
    ('total', ('total cp', 'importe total', 'total')),
    ('moneda', ('moneda',)),
    ('tc', ('tipo de cambio', 'tipo cambio', 't.c.', 'tc')),
)


def _mapear_cabecera(cab):
    """{campo: índice de columna} buscando por nombre (no por posición)."""
    nombres = [_norm(c) for c in cab]
    usados, mapa = set(), {}
    for campo, opciones in SINONIMOS:
        for op in opciones:
            for i, n in enumerate(nombres):
                if i in usados or not n:
                    continue
                if campo == 'igv' and ('dscto' in n or 'desc' in n or 'ng' == n[-2:] or 'dng' in n):
                    continue
                if campo == 'base' and ('dscto' in n or 'desc' in n or 'ng' in n.split() or 'dgng' in n):
                    continue
                if campo == 'numero' and ('identidad' in n or 'final' == n[-5:] or 'doc' in n and 'cp' not in n):
                    continue
                if campo == 'total' and n not in ('total', 'total cp', 'importe total') and not n.startswith('total'):
                    continue
                if campo == 'tipo' and ('identidad' in n or 'cambio' in n):
                    continue
                if (n == op) or (len(op) > 3 and n.startswith(op)) or (len(op) > 5 and op in n):
                    mapa[campo] = i
                    usados.add(i)
                    break
            if campo in mapa:
                break
    return mapa


def _filas_archivo(nombre, contenido):
    """Lista de filas (listas de textos/números) de un .xlsx, .csv, .txt (SIRE, con |) o .zip con uno de ellos."""
    n = (nombre or '').lower()
    if n.endswith('.zip'):
        try:
            z = zipfile.ZipFile(io.BytesIO(contenido))
        except zipfile.BadZipFile:
            raise ErrorSunat('El .zip está dañado.')
        dentro = [x for x in z.namelist() if x.lower().endswith(('.txt', '.csv', '.xlsx'))]
        if not dentro:
            raise ErrorSunat('El .zip no tiene un archivo .txt, .csv o .xlsx.')
        return _filas_archivo(dentro[0], z.read(dentro[0]))
    if n.endswith(('.xlsx', '.xlsm')):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        except Exception:
            raise ErrorSunat('No se pudo abrir el Excel.')
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if n.endswith(('.txt', '.csv')):
        for cod in ('utf-8-sig', 'latin-1'):
            try:
                texto = contenido.decode(cod)
                break
            except UnicodeDecodeError:
                continue
        primera = texto.splitlines()[0] if texto.splitlines() else ''
        sep = '|' if primera.count('|') >= 3 else (';' if primera.count(';') > primera.count(',') else ',')
        return [r for r in csv.reader(io.StringIO(texto), delimiter=sep)]
    raise ErrorSunat('Sube un archivo .xlsx, .csv, .txt (SIRE) o .zip.')


def _fecha_txt(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = str(v or '').strip()[:10]
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y'):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            continue
    return None


def _dec(v):
    if isinstance(v, (int, float, Decimal)):
        return a_decimal(v)
    t = str(v or '').strip().replace('S/', '').replace(' ', '')
    if re.match(r'^-?\d{1,3}(\.\d{3})+,\d+$', t):     # 1.234,56
        t = t.replace('.', '').replace(',', '.')
    elif re.match(r'^-?\d+,\d{1,2}$', t):             # 1234,56
        t = t.replace(',', '.')
    return a_decimal(t)


def leer_archivo(nombre, contenido, libro, anio, mes):
    """Convierte el archivo en comprobantes. Devuelve (filas, avisos). No guarda nada."""
    filas = _filas_archivo(nombre, contenido)
    if not filas:
        raise ErrorSunat('El archivo está vacío.')
    idx_cab, mapa = None, {}
    for i, fila in enumerate(filas[:15]):
        m = _mapear_cabecera(['' if c is None else str(c) for c in fila])
        if len({'base', 'igv', 'total'} & set(m)) >= 2 and ({'numero', 'comprobante'} & set(m)):
            idx_cab, mapa = i, m
            break
    if idx_cab is None:
        raise ErrorSunat('No encontré los títulos de las columnas (Fecha, Serie, Número, RUC, Base, IGV, Total). '
                         'Usa la plantilla de esta pantalla o la propuesta del SIRE tal como la descarga SUNAT.')
    desde, hasta = _rango(anio, mes)
    salida, avisos, fuera = [], [], 0
    tasa = tasa_igv()

    def tomar(fila, campo):
        i = mapa.get(campo)
        return fila[i] if i is not None and i < len(fila) else None

    for fila in filas[idx_cab + 1:]:
        if not fila or all(c in (None, '') for c in fila):
            continue
        serie, numero = str(tomar(fila, 'serie') or '').strip().upper(), str(tomar(fila, 'numero') or '').strip()
        if (not serie or not numero) and tomar(fila, 'comprobante'):
            s, _, nn = str(tomar(fila, 'comprobante')).strip().upper().rpartition('-')
            serie, numero = (s or serie), (nn or numero)
        if not numero:
            continue
        if str(numero).endswith('.0'):
            numero = numero[:-2]
        tipo = _tipo_cp(tomar(fila, 'tipo') or ('07' if serie.startswith(('FC', 'BC')) else '08' if serie.startswith(('FD', 'BD')) else '01'))
        f = _fecha_txt(tomar(fila, 'fecha'))
        base, igv, total = _dec(tomar(fila, 'base')), _dec(tomar(fila, 'igv')), _dec(tomar(fila, 'total'))
        ng = _dec(tomar(fila, 'no_gravado'))
        if not total and (base or igv):
            total = base + igv + ng
        if not base and not igv and total and tipo != '02':
            base = redondear(total / (1 + tasa))
            igv = total - base
        if tipo == '07':
            base, igv, ng, total = -abs(base), -abs(igv), -abs(ng), -abs(total)
        if f and libro == 'VENTAS' and not (desde <= f <= hasta):
            fuera += 1
            continue
        if f and libro == 'COMPRAS' and f > hasta:
            fuera += 1
            continue
        moneda = str(tomar(fila, 'moneda') or 'PEN').strip().upper()[:3] or 'PEN'
        tc = _dec(tomar(fila, 'tc')) if moneda not in ('PEN', 'S/', 'SOL') else None
        ruc = re.sub(r'\.0$', '', str(tomar(fila, 'ruc') or '').strip())
        salida.append({'tipo_cp': tipo, 'serie': serie[:20], 'numero': str(numero)[:30], 'fecha': f,
                       'ruc': ruc[:20], 'nombre': str(tomar(fila, 'nombre') or '').strip()[:255],
                       'base': redondear(base), 'igv': redondear(igv), 'no_gravado': redondear(ng),
                       'total': redondear(total), 'moneda': 'PEN' if moneda in ('S/', 'SOL') else moneda,
                       'tipo_cambio': tc if tc and tc > 0 else None})
    if fuera:
        avisos.append(f'{fuera} filas son de otro mes y no se tomaron.')
    if not salida:
        raise ErrorSunat('No encontré comprobantes de este mes en el archivo.')
    return salida, avisos


def plantilla_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    for i, libro in enumerate(('Ventas', 'Compras')):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = libro
        cab = ['Fecha de emisión', 'Tipo', 'Serie', 'Número', 'Cliente' if libro == 'Ventas' else 'Proveedor',
               'RUC / DNI', 'Base imponible', 'IGV', 'No gravado', 'Total', 'Moneda', 'Tipo de cambio']
        ws.append(cab)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='0F766E')
        ws.append(['01/10/2026', 'Factura', 'F001', '123', 'EJEMPLO S.A.C. (borra esta fila)', '20100000001', 100, 18, 0,
                   118, 'PEN', ''])
        for col, ancho in zip('ABCDEFGHIJKL', (16, 16, 8, 12, 40, 14, 14, 10, 12, 12, 8, 14)):
            ws.column_dimensions[col].width = ancho
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ============================================================
# PAGOS DEL MES (checklist)
# ============================================================
def _vencimiento(ob, e, anio, mes):
    """Vencimiento de la obligación del periodo (anio, mes)."""
    c = ob['cuando']
    sa, sm = _mes_mas(anio, mes, 1)
    if c == 'cronograma':
        return vencimiento_cronograma(e.get('ruc'), anio, mes)
    if c == 'habil5':
        return _habil_n(sa, sm, 5)
    if c == 'fin_mes':
        return _ultimo_habil(anio, mes)
    if c == 'dia15':
        return date(anio, mes, 15)
    return None


def obligaciones_del_mes(e, pago_anio, pago_mes):
    """Lo que se paga en el mes (pago_anio, pago_mes): impuestos y AFP del mes anterior + sueldos de este mes
    (+ CTS o gratificación si toca)."""
    pa, pm = _mes_mas(pago_anio, pago_mes, -1)
    r = calcular(e, pa, pm)
    montos = {'IGV': r['liq']['igv_pagar'], 'RENTA_3RA': r['renta']['monto'], 'ESSALUD': r['plame']['essalud'],
              'ONP': r['plame']['onp'], 'RENTA_5TA': r['plame']['renta_5ta'], 'RENTA_4TA': r['plame']['renta_4ta'],
              'AFP': r['plame']['afp'] or None, 'DECLARACION_621': None, 'DETRACCIONES': None}
    # Si ya está la hoja del contador de ese mes, se usan sus montos (es lo que él declara)
    hoja = leer_periodo(e['id'], _txt_periodo(pa, pm))['contador'] or {}
    if hoja:
        imp = a_decimal(hoja.get('importe_igv'))
        if hoja.get('importe_igv') is not None:
            montos['IGV'] = imp if imp > 0 else Decimal('0')
        for cod, k in (('RENTA_3RA', 'renta'), ('ESSALUD', 'essalud'), ('ONP', 'onp')):
            if hoja.get(k) is not None:
                montos[cod] = a_decimal(hoja.get(k))
    else:
        for cod in ('IGV', 'RENTA_3RA', 'ONP', 'ESSALUD'):
            if not montos[cod]:
                montos[cod] = None
    plan_hoy = leer_periodo(e['id'], _txt_periodo(pago_anio, pago_mes))['planilla'] or {}
    montos['REMUNERACIONES'] = a_decimal(plan_hoy.get('remuneraciones')) or r['plame']['remuneraciones'] or None
    filas = []
    for ob in OBLIGACIONES:
        if ob['cuando'] in ('fin_mes', 'dia15'):
            if ob.get('meses') and pago_mes not in ob['meses']:
                continue
            anio, mes = pago_anio, pago_mes
        else:
            anio, mes = pa, pm
        filas.append((ob, anio, mes))
    guardados = {}
    periodos = list({_txt_periodo(a, m) for _, a, m in filas})
    for g in db_query("SELECT * FROM sunat_obligaciones WHERE empresa_id = %s AND periodo = ANY(%s)", (str(e['id']), periodos)):
        guardados[(g['periodo'], g['codigo'])] = g
    hoy = date.today()
    salida = []
    for ob, anio, mes in filas:
        per = _txt_periodo(anio, mes)
        g = guardados.get((per, ob['codigo']))
        monto_sug = montos.get(ob['codigo'])
        venc = (g and g['vencimiento']) or _vencimiento(ob, e, anio, mes)
        if g:
            estado = g['estado']
        elif ob['codigo'] in NO_APLICA_SI_CERO and monto_sug is not None and a_decimal(monto_sug) <= 0 \
                and ob['codigo'] not in ('REMUNERACIONES', 'AFP', 'ESSALUD') \
                and (hoja or ob['codigo'] not in ('IGV', 'RENTA_3RA', 'ONP')):
            # Sin la hoja del contador, un cero del ERP puede ser falta de datos: no se marca solo
            estado = 'No aplica'
        else:
            estado = 'Pendiente'
        monto = a_decimal(g['monto']) if g and g['monto'] is not None else monto_sug
        dias = (venc - hoy).days if venc else None
        if estado == 'Pagado':
            semaforo, texto = 'pagado', '✔ Pagado'
        elif estado == 'No aplica':
            semaforo, texto = 'no_aplica', 'No aplica'
        elif dias is not None and dias < 0:
            semaforo, texto = 'vencido', f'⚠ Vencido hace {-dias} día(s)'
        elif dias is not None and dias <= 3:
            semaforo, texto = 'urgente', 'Vence hoy' if dias == 0 else f'Pendiente · vence en {dias} día(s)'
        else:
            semaforo, texto = 'pendiente', 'Pendiente'
        if ob['codigo'] == 'IGV' and estado == 'No aplica' and not g:
            texto = 'No aplica (S/ 0.00)'
        salida.append({'codigo': ob['codigo'], 'nombre': ob['nombre'], 'donde': ob['donde'], 'ayuda': ob['ayuda'],
                       'periodo': per, 'nombre_periodo': _nombre_periodo(anio, mes), 'vencimiento': venc,
                       'monto': None if monto is None else _n(monto), 'monto_sugerido': None if monto_sug is None else _n(monto_sug),
                       'estado': estado, 'semaforo': semaforo, 'texto_estado': texto,
                       'fecha_pago': g['fecha_pago'] if g else None, 'numero_operacion': g['numero_operacion'] if g else None,
                       'observacion': g['observacion'] if g else None, 'marcado_por': g['marcado_por'] if g else None,
                       'marcado_en': g['marcado_en'] if g else None, 'guardado': bool(g)})
    salida.sort(key=lambda x: (x['vencimiento'] or date.max, [o['codigo'] for o in OBLIGACIONES].index(x['codigo'])))
    return salida


# ============================================================
# RUTAS
# ============================================================
def _err(ex, codigo=400):
    return jsonify({'success': False, 'error': str(ex)}), codigo


def _json(v):
    if isinstance(v, Decimal):
        return _n(v)
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d %H:%M')
    if isinstance(v, date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, dict):
        return {k: _json(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [_json(x) for x in v]
    return v


@sunat_bp.route('/sunat')
def pagina():
    emps = [{'id': x['id'], 'codigo': x['codigo'], 'ruc': x.get('ruc'), 'razon_social': x.get('razon_social')}
            for x in lista_empresas()]
    activa = empresa_actual()
    return render_template('sunat.html', empresas=emps, empresa_activa=activa['id'] if activa else None,
                           es_admin=es_admin(), regimenes={k: v['nombre'] for k, v in REGIMENES.items()})


@sunat_bp.route('/sunat/api/resumen')
def api_resumen():
    try:
        e = _empresa(request.args.get('empresa_id'))
        anio, mes = _periodo(request.args.get('periodo'))
        r = calcular(e, anio, mes)
        partes = proyeccion(e, r)
        proy = proyectar(r, partes)
        plan = plan_accion(e, r, proy)
        filas, obs = comparar_contador(r)
        if r['contador']:
            c = dict(r['contador'])
            c['vencimiento_sunat'] = r['vencimiento'].strftime('%d/%m/%Y') if r['vencimiento'] else None
            if c.get('vencimiento_hoja'):
                m = re.match(r'^(\d{1,2})\s+([A-ZÁÉÍÓÚ]+)\s+(\d{4})', c['vencimiento_hoja'].upper())
                if m and m.group(2)[:3] in _HOJAS:
                    c['vencimiento_hoja'] = f"{int(m.group(1)):02d}/{_HOJAS[m.group(2)[:3]]:02d}/{m.group(3)}"
            pa, pm = _mes_mas(anio, mes, -1)
            obs = _observaciones_contador(c, leer_periodo(e['id'], _txt_periodo(pa, pm))['contador'] or None)
        datos = {k: r[k] for k in ('empresa', 'periodo', 'nombre_periodo', 'desde', 'hasta', 'corte', 'es_mes_actual',
                                   'vencimiento', 'manual', 'origen_saldos', 'planilla', 'planilla_origen', 'sv', 'sc',
                                   'sin_tipo_cambio', 'diferencias', 'retenciones_sugeridas', 'retenciones_sugeridas_n',
                                   'liq', 'renta', 'plame', 'detracciones_bn', 'total_sunat', 'pagar_con_detracciones',
                                   'pagar_efectivo', 'contador_archivo', 'contador_cargado_en')}
        datos['config'] = {k: r['config'][k] for k in ('regimen', 'tasa_renta', 'es_agente_retencion', 'configurado')}
        datos['config']['regimen_nombre'] = REGIMENES.get(r['config']['regimen'], {}).get('nombre', r['config']['regimen'])
        datos['manual_calculo'] = {k: r['manual_calculo'].get(k) for k in CAMPOS_MANUAL}
        datos['proyeccion'] = {'partes': partes, **proy}
        datos['plan'] = plan
        datos['preliminar'] = filas
        datos['observaciones_contador'] = obs
        datos['tiene_contador'] = bool(r['contador'])
        # Lo pagado de verdad (voucher), según el check de "Pagos del mes"
        datos['voucher'] = {g['codigo']: {'estado': g['estado'], 'monto': g['monto'], 'fecha_pago': g['fecha_pago'],
                                          'numero_operacion': g['numero_operacion']}
                            for g in db_query("SELECT codigo, estado, monto, fecha_pago, numero_operacion FROM sunat_obligaciones "
                                              "WHERE empresa_id = %s AND periodo = %s AND codigo IN ('IGV', 'RENTA_3RA')",
                                              (str(e['id']), r['periodo']))}
        return jsonify({'success': True, 'data': _json(datos)})
    except ErrorSunat as ex:
        return _err(ex)


@sunat_bp.route('/sunat/api/documentos')
def api_documentos():
    try:
        e = _empresa(request.args.get('empresa_id'))
        anio, mes = _periodo(request.args.get('periodo'))
        r = calcular(e, anio, mes, corte=_rango(anio, mes)[1])
        libro = 'compras' if request.args.get('libro') == 'COMPRAS' else 'ventas'
        docs = sorted(r[libro], key=lambda x: (x['fecha'] or date.min, str(x['serie']), _num_doc(x['numero'])))
        return jsonify({'success': True, 'data': _json([{
            'tercero': d['nombre'], 'ruc': d['ruc'], 'tipo': TIPOS_CP.get(d['tipo'], d['tipo']),
            'documento': f"{d['serie']}-{d['numero']}", 'fecha': d['fecha'], 'base': d['base'], 'igv': d['igv'],
            'no_gravado': d['no_gravado'], 'total': d['total'], 'fuente': d['fuente']} for d in docs])})
    except ErrorSunat as ex:
        return _err(ex)


@sunat_bp.route('/sunat/api/periodo', methods=['POST'])
def api_guardar_periodo():
    d = request.get_json(silent=True) or {}
    try:
        e = _empresa(d.get('empresa_id'))
        anio, mes = _periodo(d.get('periodo'))
        valores = {}
        for k in CAMPOS_MANUAL:
            if k in (d.get('manual') or {}):
                v = d['manual'][k]
                if v in (None, ''):
                    valores[k] = None
                else:
                    x = _dec(v)
                    if x < 0:
                        raise ErrorSunat('Escribe los montos en positivo (sin signo menos).')
                    valores[k] = str(redondear(x, 4 if k == 'tipo_cambio' else 2))
        plan = {}
        for k in CAMPOS_PLANILLA:
            if k in (d.get('planilla') or {}):
                v = d['planilla'][k]
                plan[k] = None if v in (None, '') else str(redondear(abs(_dec(v))))
        if valores:
            guardar_periodo(e['id'], _txt_periodo(anio, mes), 'manual', valores)
        if plan:
            guardar_periodo(e['id'], _txt_periodo(anio, mes), 'planilla', plan)
        return jsonify({'success': True, 'mensaje': 'Datos del mes guardados.'})
    except ErrorSunat as ex:
        return _err(ex)


@sunat_bp.route('/sunat/api/config', methods=['GET', 'POST'])
def api_config():
    try:
        e = _empresa(request.values.get('empresa_id') or (request.get_json(silent=True) or {}).get('empresa_id'))
        if request.method == 'GET':
            return jsonify({'success': True, 'data': _json(leer_config(e['id']))})
        if not es_admin():
            return _err('Solo la gerencia cambia el régimen tributario.', 403)
        d = request.get_json(silent=True) or {}
        regimen = (d.get('regimen') or 'RMT').upper()
        if regimen not in REGIMENES:
            raise ErrorSunat('Régimen inválido.')
        tasa = _dec(d.get('tasa_renta')) if d.get('tasa_renta') not in (None, '') else REGIMENES[regimen]['tasa']
        if not (Decimal('0') <= tasa <= Decimal('10')):
            raise ErrorSunat('La tasa de renta mensual debe estar entre 0% y 10%.')
        with db_tx() as conn:
            conn.cursor().execute("""
                INSERT INTO sunat_config (empresa_id, regimen, tasa_renta, es_agente_retencion, notas, actualizado_por, actualizado_en)
                VALUES (%s, %s, %s, %s, %s, %s, NOW())
                ON CONFLICT (empresa_id) DO UPDATE SET regimen = EXCLUDED.regimen, tasa_renta = EXCLUDED.tasa_renta,
                    es_agente_retencion = EXCLUDED.es_agente_retencion, notas = EXCLUDED.notas,
                    actualizado_por = EXCLUDED.actualizado_por, actualizado_en = NOW()
            """, (e['id'], regimen, tasa, bool(d.get('es_agente_retencion')), (d.get('notas') or '').strip() or None, _usuario()))
        return jsonify({'success': True, 'mensaje': 'Régimen guardado.'})
    except ErrorSunat as ex:
        return _err(ex)


@sunat_bp.route('/sunat/api/importar/vista', methods=['POST'])
def api_importar_vista():
    try:
        e = _empresa(request.form.get('empresa_id'))
        anio, mes = _periodo(request.form.get('periodo'))
        libro = request.form.get('libro')
        if libro not in ('VENTAS', 'COMPRAS'):
            raise ErrorSunat('Elige si el archivo es de ventas o de compras.')
        arch = request.files.get('archivo')
        if not arch or not arch.filename:
            raise ErrorSunat('Elige el archivo.')
        contenido = arch.read()
        if len(contenido) > 8 * 1024 * 1024:
            raise ErrorSunat('El archivo pesa más de 8 MB.')
        filas, avisos = leer_archivo(arch.filename, contenido, libro, anio, mes)
        # ¿Cuáles ya están en el ERP?
        desde, hasta = _rango(anio, mes)
        erp = ventas_erp(e['id'], desde, hasta)[0] if libro == 'VENTAS' else compras_erp(e['id'], date(anio - 1, 1, 1), hasta)
        claves = {_clave(libro, d['tipo'], d['serie'], d['numero'], d['ruc']) for d in erp}
        for f in filas:
            f['en_erp'] = _clave(libro, f['tipo_cp'], f['serie'], f['numero'], f['ruc']) in claves
        ya = db_query("SELECT COUNT(*) AS n FROM sunat_registros WHERE empresa_id = %s AND periodo = %s AND libro = %s",
                      (e['id'], _txt_periodo(anio, mes), libro))[0]['n']
        if ya:
            avisos.append(f'Ya hay {ya} comprobantes importados de {libro.lower()} en este mes. Si es el mismo archivo '
                          f'actualizado, marca "Reemplazar lo importado antes".')
        resumen = {'n': len(filas), 'en_erp': sum(1 for f in filas if f['en_erp']),
                   'base': sum((f['base'] for f in filas), CERO), 'igv': sum((f['igv'] for f in filas), CERO),
                   'total': sum((f['total'] for f in filas), CERO), 'ya_importados': ya}
        return jsonify({'success': True, 'data': _json({'filas': filas, 'avisos': avisos, 'resumen': resumen,
                                                        'archivo': arch.filename})})
    except ErrorSunat as ex:
        return _err(ex)


@sunat_bp.route('/sunat/api/importar/guardar', methods=['POST'])
def api_importar_guardar():
    d = request.get_json(silent=True) or {}
    try:
        e = _empresa(d.get('empresa_id'))
        anio, mes = _periodo(d.get('periodo'))
        libro = d.get('libro')
        if libro not in ('VENTAS', 'COMPRAS'):
            raise ErrorSunat('Elige si el archivo es de ventas o de compras.')
        filas = d.get('filas') or []
        if not filas or len(filas) > 20000:
            raise ErrorSunat('No hay comprobantes para guardar.')
        origen = 'SIRE' if (d.get('origen') or '').upper() == 'SIRE' else 'EXCEL'
        lote = f"{_ahora_lima().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:6]}"
        periodo = _txt_periodo(anio, mes)
        with db_tx() as conn:
            cur = conn.cursor()
            if d.get('reemplazar'):
                cur.execute("DELETE FROM sunat_registros WHERE empresa_id = %s AND periodo = %s AND libro = %s",
                            (e['id'], periodo, libro))
            for f in filas:
                base, igv, ng, total = (redondear(_dec(f.get(k))) for k in ('base', 'igv', 'no_gravado', 'total'))
                if abs(total - (base + igv + ng)) > Decimal('1') and base:
                    raise ErrorSunat(f"El comprobante {f.get('serie')}-{f.get('numero')} no cuadra: base + IGV + no gravado "
                                     f"no da el total.")
                cur.execute("""
                    INSERT INTO sunat_registros (empresa_id, periodo, libro, lote, origen, archivo, tipo_cp, serie, numero,
                        fecha, ruc, nombre, base, igv, no_gravado, total, moneda, tipo_cambio, creado_por)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                            (e['id'], periodo, libro, lote, origen, (d.get('archivo') or '')[:255], _tipo_cp(f.get('tipo_cp')),
                             str(f.get('serie') or '')[:20], str(f.get('numero') or '')[:30], _fecha_txt(f.get('fecha')),
                             str(f.get('ruc') or '')[:20], str(f.get('nombre') or '')[:255], base, igv, ng, total,
                             str(f.get('moneda') or 'PEN')[:5], _dec(f.get('tipo_cambio')) if f.get('tipo_cambio') else None,
                             _usuario()))
        return jsonify({'success': True, 'mensaje': f'Se importaron {len(filas)} comprobantes de {libro.lower()}.', 'lote': lote})
    except ErrorSunat as ex:
        return _err(ex)


@sunat_bp.route('/sunat/api/importados')
def api_importados():
    try:
        e = _empresa(request.args.get('empresa_id'))
        anio, mes = _periodo(request.args.get('periodo'))
        filas = db_query("""
            SELECT lote, libro, origen, archivo, COUNT(*) AS n, SUM(base) AS base, SUM(igv) AS igv, SUM(total) AS total,
                   MIN(creado_por) AS creado_por, MIN(creado_en) AS creado_en
            FROM sunat_registros WHERE empresa_id = %s AND periodo = %s
            GROUP BY lote, libro, origen, archivo ORDER BY MIN(creado_en) DESC""", (e['id'], _txt_periodo(anio, mes)))
        return jsonify({'success': True, 'data': _json(filas)})
    except ErrorSunat as ex:
        return _err(ex)


@sunat_bp.route('/sunat/api/importados/quitar', methods=['POST'])
def api_importados_quitar():
    d = request.get_json(silent=True) or {}
    try:
        e = _empresa(d.get('empresa_id'))
        lote = str(d.get('lote') or '')
        if not lote:
            raise ErrorSunat('Falta el lote.')
        with db_tx() as conn:
            cur = conn.cursor()
            cur.execute("DELETE FROM sunat_registros WHERE empresa_id = %s AND lote = %s", (e['id'], lote))
            n = cur.rowcount
        return jsonify({'success': True, 'mensaje': f'Se quitó la importación ({n} comprobantes). Las ventas y compras del ERP no cambian.'})
    except ErrorSunat as ex:
        return _err(ex)


@sunat_bp.route('/sunat/api/plantilla.xlsx')
def api_plantilla():
    return Response(plantilla_excel(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': 'attachment; filename=plantilla_ventas_compras_SUNAT.xlsx'})


@sunat_bp.route('/sunat/api/preliminar.xlsx')
def api_preliminar_xlsx():
    try:
        e = _empresa(request.args.get('empresa_id'))
        anio, mes = _periodo(request.args.get('periodo'))
        r = calcular(e, anio, mes)
        nombre = f"PRELIMINAR_{e['codigo']}_IMPUESTOS_{MESES[mes - 1][:3].upper()}_{anio}.xlsx"
        return Response(excel_preliminar(r), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        headers={'Content-Disposition': f'attachment; filename={nombre}'})
    except ErrorSunat as ex:
        return _err(ex)


@sunat_bp.route('/sunat/api/contador', methods=['POST'])
def api_contador():
    """Sube la hoja del contador: se guardan sus cifras para compararlas con el preliminar."""
    try:
        e = _empresa(request.form.get('empresa_id'))
        anio, mes = _periodo(request.form.get('periodo'))
        arch = request.files.get('archivo')
        if not arch or not arch.filename.lower().endswith(('.xlsx', '.xlsm')):
            raise ErrorSunat('Sube el Excel del contador (.xlsx).')
        hoja, valores = leer_hoja_contador(arch.read(), anio, mes)
        periodo = _txt_periodo(anio, mes)
        with db_tx() as conn:
            conn.cursor().execute("""
                INSERT INTO sunat_periodos (empresa_id, periodo, contador, contador_archivo, contador_cargado_en,
                                            actualizado_por, actualizado_en)
                VALUES (%s, %s, %s, %s, NOW(), %s, NOW())
                ON CONFLICT (empresa_id, periodo) DO UPDATE SET contador = EXCLUDED.contador,
                    contador_archivo = EXCLUDED.contador_archivo, contador_cargado_en = NOW(),
                    actualizado_por = EXCLUDED.actualizado_por, actualizado_en = NOW()
            """, (e['id'], periodo, Json(valores), f'{arch.filename} · hoja {hoja}'[:255], _usuario()))
        # La planilla del contador sirve como dato del mes si no se escribió otra
        per = leer_periodo(e['id'], periodo)
        if not any((per['planilla'] or {}).get(k) not in (None, '') for k in CAMPOS_PLANILLA):
            plan = {k: valores[k] for k in ('base_essalud', 'base_onp') if valores.get(k)}
            if plan:
                guardar_periodo(e['id'], periodo, 'planilla', plan)
        return jsonify({'success': True, 'mensaje': f'Se leyó la hoja "{hoja}" del contador.', 'data': valores})
    except ErrorSunat as ex:
        return _err(ex)


@sunat_bp.route('/sunat/api/pagos')
def api_pagos():
    try:
        e = _empresa(request.args.get('empresa_id'))
        anio, mes = _periodo(request.args.get('mes'))
        filas = obligaciones_del_mes(e, anio, mes)
        resumen = {'pendientes': sum(1 for f in filas if f['estado'] == 'Pendiente'),
                   'vencidos': sum(1 for f in filas if f['semaforo'] == 'vencido'),
                   'pagados': sum(1 for f in filas if f['estado'] == 'Pagado'),
                   'no_aplica': sum(1 for f in filas if f['estado'] == 'No aplica'),
                   'monto_pendiente': sum((a_decimal(f['monto']) for f in filas if f['estado'] == 'Pendiente' and f['monto']), CERO),
                   'monto_pagado': sum((a_decimal(f['monto']) for f in filas if f['estado'] == 'Pagado' and f['monto']), CERO)}
        return jsonify({'success': True, 'data': _json({'filas': filas, 'resumen': resumen,
                                                        'mes': _txt_periodo(anio, mes), 'nombre_mes': _nombre_periodo(anio, mes),
                                                        'empresa': {'codigo': e['codigo'], 'ruc': e.get('ruc')}})})
    except ErrorSunat as ex:
        return _err(ex)


@sunat_bp.route('/sunat/api/pagos/marcar', methods=['POST'])
def api_pagos_marcar():
    d = request.get_json(silent=True) or {}
    try:
        e = _empresa(d.get('empresa_id'))
        codigo = d.get('codigo')
        if codigo not in CODIGOS_OBLIGACION:
            raise ErrorSunat('Obligación desconocida.')
        anio, mes = _periodo(d.get('periodo'))
        periodo = _txt_periodo(anio, mes)
        estado = d.get('estado')
        if estado not in ('Pendiente', 'Pagado', 'No aplica'):
            raise ErrorSunat('Estado inválido.')
        monto = None if d.get('monto') in (None, '') else redondear(abs(_dec(d.get('monto'))))
        fecha_pago = _fecha_txt(d.get('fecha_pago')) if estado == 'Pagado' else None
        if estado == 'Pagado':
            fecha_pago = fecha_pago or date.today()
            if fecha_pago > date.today():
                raise ErrorSunat('La fecha de pago no puede ser futura.')
            # Trazabilidad: cada pago queda con su monto y el N° de operación del voucher
            if monto is None:
                raise ErrorSunat('Escribe el monto pagado (el del voucher).')
            if codigo not in ('REMUNERACIONES', 'CTS', 'GRATIFICACION') and not (d.get('numero_operacion') or '').strip():
                raise ErrorSunat('Escribe el N° de operación o NPS del voucher de pago.')
        venc = _fecha_txt(d.get('vencimiento')) if d.get('vencimiento') else None
        obs = (d.get('observacion') or '').strip()[:500] or None
        with db_tx() as conn:
            cur = conn.cursor(cursor_factory=RealDictCursor)
            cur.execute("SELECT * FROM sunat_obligaciones WHERE empresa_id = %s AND periodo = %s AND codigo = %s FOR UPDATE",
                        (e['id'], periodo, codigo))
            previo = cur.fetchone()
            # Candado: lo pagado solo se desmarca con motivo y queda en el historial
            if previo and previo['estado'] == 'Pagado' and estado != 'Pagado':
                if not obs:
                    raise ErrorSunat('Para desmarcar un pago escribe el motivo (queda en el historial).')
                if not es_admin() and previo['marcado_por'] != _usuario():
                    raise ErrorSunat(f"Lo marcó {previo['marcado_por']}. Solo esa persona o la gerencia lo pueden desmarcar.")
            hist = list(previo['historial'] or []) if previo else []
            hist.append({'estado': estado, 'por': _usuario(), 'en': _ahora_lima().strftime('%Y-%m-%d %H:%M'),
                         'monto': str(monto) if monto is not None else None, 'obs': obs})
            cur.execute("""
                INSERT INTO sunat_obligaciones (empresa_id, periodo, codigo, estado, monto, vencimiento, fecha_pago,
                    numero_operacion, observacion, marcado_por, marcado_en, historial)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,NOW(),%s)
                ON CONFLICT (empresa_id, periodo, codigo) DO UPDATE SET estado = EXCLUDED.estado,
                    monto = COALESCE(EXCLUDED.monto, sunat_obligaciones.monto),
                    vencimiento = COALESCE(EXCLUDED.vencimiento, sunat_obligaciones.vencimiento),
                    fecha_pago = EXCLUDED.fecha_pago, numero_operacion = EXCLUDED.numero_operacion,
                    observacion = EXCLUDED.observacion, marcado_por = EXCLUDED.marcado_por, marcado_en = NOW(),
                    historial = EXCLUDED.historial
            """, (e['id'], periodo, codigo, estado, monto, venc, fecha_pago,
                  (d.get('numero_operacion') or '').strip()[:80] or None, obs, _usuario(), Json(hist)))
        textos = {'Pagado': 'Marcado como pagado.', 'No aplica': 'Marcado como "No aplica".', 'Pendiente': 'Vuelve a pendiente.'}
        return jsonify({'success': True, 'mensaje': textos[estado]})
    except ErrorSunat as ex:
        return _err(ex)
