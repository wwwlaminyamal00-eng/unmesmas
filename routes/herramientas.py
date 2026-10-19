# routes/herramientas.py - Módulo 9 "Herramientas" (solo administradores; pedido de Erika 16/10/2026)
#
#   1. Importar productos, clientes y proveedores desde Excel/CSV (plantilla, vista previa, validación).
#      Solo agrega nuevos o actualiza si se marca; NUNCA borra.
#   2. Exportar tablas de negocio a CSV/Excel (respaldo). Nunca exporta contraseñas ni claves.
#   3. Revisar numeraciones: avisa huecos y números repetidos. NUNCA renumera guías ni facturas.
#   4. Verificar integridad: registros huérfanos, duplicados y stock que no cuadra. Solo reporta.
#   5. Recalcular stock (productos.stock = suma de inv_stock).
#   6. Limpiar caché del sistema.  7. Ver / rotar el registro de errores.  8. Estado de la base.
# Cada uso queda en Auditoría (erp_auditoria).
import csv
import io
import logging
import os
import platform
import re
import time
import unicodedata
from collections import deque
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from logging.handlers import RotatingFileHandler

from flask import Blueprint, Response, jsonify, render_template, request, session, redirect, url_for, flash
from psycopg2.extras import RealDictCursor

from database import db_query, db_tx
from utils import es_admin
from routes.auditoria import registrar_auditoria, marcar_usuario

herramientas_bp = Blueprint('herramientas', __name__)
MODULO = 'Herramientas'


# ============================================================
# PROTECCIÓN: todo el módulo es solo para administradores
# ============================================================
@herramientas_bp.before_request
def _solo_admin():
    es_api = '/api/' in request.path
    if 'usuario_id' not in session:
        if es_api:
            return jsonify({'success': False, 'error': 'Sesión expirada. Inicia sesión.'}), 401
        return redirect(url_for('login'))
    if not es_admin():
        if es_api:
            return jsonify({'success': False, 'error': 'Solo un administrador puede usar Herramientas.'}), 403
        flash('Herramientas es solo para administradores.', 'error')
        return redirect(url_for('index'))
    return None


@herramientas_bp.route('/herramientas')
def herramientas():
    return render_template('herramientas.html', tablas_import=IMPORTABLES, tablas_export=EXPORTABLES)


# ============================================================
# AYUDAS DE BASE DE DATOS
# ============================================================
def columnas(tabla):
    return {f['column_name']: f['data_type'] for f in db_query(
        """SELECT column_name, data_type FROM information_schema.columns
           WHERE table_schema = current_schema() AND table_name = %s""", (tabla,))}


def existe_tabla(tabla):
    return bool(db_query("SELECT to_regclass(%s) IS NOT NULL AS e", (tabla,))[0]['e'])


def fecha_txt(v):
    if isinstance(v, datetime):
        return v.strftime('%d/%m/%Y %H:%M')
    if isinstance(v, date):
        return v.strftime('%d/%m/%Y')
    return v


def _plano(v):
    if isinstance(v, (datetime, date)):
        return fecha_txt(v)
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (dict, list)):
        import json
        return json.dumps(v, ensure_ascii=False, default=str)
    return v


def _err(msg, codigo=400):
    return jsonify({'success': False, 'error': msg}), codigo


def _normal(t):
    t = unicodedata.normalize('NFKD', str(t or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', '_', t).strip('_')


# ============================================================
# 1. IMPORTAR (productos, clientes, proveedores)
# ============================================================
# campo, etiqueta, obligatorio, tipo, valor por defecto (si es obligatorio y viene vacío)
IMPORTABLES = {
    'productos': {
        'titulo': 'Productos', 'clave': 'codigo',
        'campos': [
            ('codigo', 'Código', True, 'texto', None),
            ('descripcion', 'Descripción', True, 'texto', None),
            ('familia', 'Familia', True, 'texto', 'GENERAL'),
            ('unidad', 'Unidad', True, 'texto', 'UND'),
            ('marca', 'Marca', False, 'texto', None),
            ('modelo', 'Modelo', False, 'texto', None),
            ('costo_unitario', 'Costo unitario sin IGV', False, 'numero', None),
            ('precio_unitario', 'Precio de venta sin IGV', False, 'numero', None),
            ('stock_minimo', 'Stock mínimo', False, 'entero', None),
            ('codigo_barras', 'Código de barras', False, 'texto', None),
            ('descripcion_larga', 'Descripción larga', False, 'texto', None),
            ('observaciones', 'Observaciones', False, 'texto', None),
        ],
        'ejemplo': ['PRD-9001', 'GUANTE DE NITRILO TALLA M', 'EPP', 'PAR', 'ANSELL', '37-501', '10.50', '15.90', '20', '', '', ''],
        'nota': 'El stock NO se importa aquí: se carga en Inventario (entrada de stock inicial) para que quede en el kardex.',
    },
    'clientes': {
        'titulo': 'Clientes', 'clave': 'numero_documento',
        'campos': [
            ('numero_documento', 'RUC o DNI', True, 'documento', None),
            ('razon_social', 'Razón social', True, 'texto', None),
            ('direccion_fiscal', 'Dirección fiscal', True, 'texto', None),
            ('tipo_documento', 'Tipo de documento (RUC/DNI)', False, 'tipo_doc', None),
            ('nombre_comercial', 'Nombre comercial', False, 'texto', None),
            ('nombre_contacto', 'Contacto', False, 'texto', None),
            ('telefono_contacto', 'Teléfono', False, 'texto', None),
            ('email_contacto', 'Correo', False, 'correo', None),
            ('condicion_pago', 'Condición de pago', False, 'texto', None),
            ('dias_credito', 'Días de crédito', False, 'entero', None),
            ('ambito', 'Empresa (COMPARTIDO, KCF, AGD)', False, 'texto', None),
            ('observaciones', 'Observaciones', False, 'texto', None),
        ],
        'ejemplo': ['20100070970', 'EMPRESA EJEMPLO S.A.C.', 'Av. Ejemplo 123, Lima', 'RUC', '', 'Ana Pérez', '999888777',
                    'compras@ejemplo.pe', 'Crédito', '30', 'COMPARTIDO', ''],
        'nota': 'Si no escribes el tipo de documento: 11 dígitos = RUC, 8 dígitos = DNI.',
    },
    'proveedores': {
        'titulo': 'Proveedores', 'clave': 'ruc',
        'campos': [
            ('ruc', 'RUC', True, 'documento', None),
            ('razon_social', 'Razón social', True, 'texto', None),
            ('direccion', 'Dirección', True, 'texto', None),
            ('razon_comercial', 'Nombre comercial', False, 'texto', None),
            ('contacto', 'Contacto', False, 'texto', None),
            ('telefono', 'Teléfono', False, 'texto', None),
            ('email', 'Correo', False, 'correo', None),
            ('condicion_pago', 'Condición de pago', False, 'texto', None),
            ('tiempo_credito', 'Días de crédito', False, 'texto', None),
            ('banco', 'Banco', False, 'texto', None),
            ('numero_cuenta', 'Número de cuenta', False, 'texto', None),
            ('cci', 'CCI', False, 'texto', None),
            ('observaciones', 'Observaciones', False, 'texto', None),
        ],
        'ejemplo': ['20600011111', 'PROVEEDOR EJEMPLO S.A.C.', 'Av. Industrial 1, Lima', '', 'Luis Ramos', '988777666',
                    'ventas@proveedor.pe', 'Crédito 30 días', '30', 'BCP', '191-0000000-0-00', ''],
        'nota': 'El código de proveedor se asigna solo.',
    },
}
MAX_FILAS_IMPORTAR = 5000


@herramientas_bp.route('/herramientas/api/plantilla/<tipo>')
def api_plantilla(tipo):
    cfg = IMPORTABLES.get(tipo)
    if not cfg:
        return _err('Tipo de plantilla no válido.')
    encabezado = [c[1] + (' *' if c[2] else '') for c in cfg['campos']]
    formato = request.args.get('formato', 'xlsx')
    registrar_auditoria('herramienta', MODULO, detalle=f'Descargó la plantilla de {cfg["titulo"]} ({formato})')
    if formato == 'csv':
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=';')
        w.writerow(encabezado)
        w.writerow(cfg['ejemplo'])
        return Response('﻿' + buf.getvalue(), mimetype='text/csv; charset=utf-8',
                        headers={'Content-Disposition': f'attachment; filename=plantilla_{tipo}.csv'})
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = cfg['titulo']
    ws.append(encabezado)
    ws.append(cfg['ejemplo'])
    for i, c in enumerate(cfg['campos'], start=1):
        celda = ws.cell(row=1, column=i)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='DC2626' if c[2] else '334155')
        ws.column_dimensions[celda.column_letter].width = max(14, len(c[1]) + 4)
        for fila in (1, 2):  # que Excel no convierta RUC/códigos en números
            ws.cell(row=fila, column=i).number_format = '@'
    ins = wb.create_sheet('Instrucciones')
    for linea in ('Cómo llenar la plantilla',
                  '1. Una fila por registro. No cambies los títulos de la primera fila.',
                  '2. Las columnas en rojo (con *) son obligatorias.',
                  '3. La fila 2 es un EJEMPLO: bórrala o reemplázala con tus datos.',
                  '4. Si el código / RUC ya existe, el sistema NO lo duplica: solo lo actualiza si marcas "Actualizar existentes".',
                  '5. Nada se borra con la importación.',
                  cfg['nota']):
        ins.append([linea])
    ins['A1'].font = Font(bold=True, size=14)
    ins.column_dimensions['A'].width = 110
    salida = io.BytesIO()
    wb.save(salida)
    return Response(salida.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': f'attachment; filename=plantilla_{tipo}.xlsx'})


def leer_archivo(archivo):
    """Filas del archivo como lista de dicts {encabezado: valor}. Acepta .xlsx y .csv (coma o punto y coma)."""
    nombre = (archivo.filename or '').lower()
    datos = archivo.read()
    if nombre.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        filas = list(ws.iter_rows(values_only=True))
        wb.close()
    elif nombre.endswith(('.csv', '.txt')):
        texto = None
        for cod in ('utf-8-sig', 'latin-1'):
            try:
                texto = datos.decode(cod)
                break
            except UnicodeDecodeError:
                continue
        primera = texto.splitlines()[0] if texto.strip() else ''
        sep = ';' if primera.count(';') >= primera.count(',') else ','
        filas = list(csv.reader(io.StringIO(texto), delimiter=sep))
    else:
        raise ValueError('Sube un archivo Excel (.xlsx) o CSV (.csv).')
    filas = [f for f in filas if f and any(v not in (None, '') and str(v).strip() for v in f)]
    if not filas:
        raise ValueError('El archivo está vacío.')
    encabezado = [str(v or '').strip() for v in filas[0]]
    return encabezado, [list(f) + [None] * (len(encabezado) - len(f)) for f in filas[1:]]


def _mapa_columnas(cfg, encabezado):
    """Relaciona cada columna del archivo con un campo (acepta el título o el nombre del campo)."""
    alias = {}
    for campo, etiqueta, *_ in cfg['campos']:
        alias[_normal(campo)] = campo
        alias[_normal(etiqueta)] = campo
    if cfg['clave'] == 'numero_documento':
        alias.update({'ruc': 'numero_documento', 'dni': 'numero_documento', 'documento': 'numero_documento'})
    mapa, ignoradas = {}, []
    for i, h in enumerate(encabezado):
        campo = alias.get(_normal(h.replace('*', '')))
        if campo and campo not in mapa.values():
            mapa[i] = campo
        elif h:
            ignoradas.append(h)
    return mapa, ignoradas


def _valor(tipo, v):
    """Convierte y valida un valor. Devuelve (valor, error)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, None
    if isinstance(v, float) and v.is_integer() and tipo in ('documento', 'texto'):
        v = int(v)
    texto = str(v).strip()
    if tipo in ('numero', 'entero'):
        try:
            n = Decimal(texto.replace(' ', '').replace(',', '.') if texto.count(',') == 1 and '.' not in texto
                        else texto.replace(',', '').replace(' ', ''))
        except InvalidOperation:
            return None, f'"{texto}" no es un número'
        if n < 0:
            return None, f'{texto} no puede ser negativo'
        return (int(n) if tipo == 'entero' else n), None
    if tipo == 'documento':
        doc = re.sub(r'\D', '', texto)
        if len(doc) not in (8, 11):
            return None, f'"{texto}" no es un RUC (11 dígitos) ni un DNI (8 dígitos)'
        return doc, None
    if tipo == 'tipo_doc':
        t = texto.upper()
        if t not in ('RUC', 'DNI'):
            return None, f'Tipo de documento "{texto}" no válido (usa RUC o DNI)'
        return t, None
    if tipo == 'correo' and '@' not in texto:
        return None, f'"{texto}" no parece un correo'
    return texto, None


def validar_importacion(tipo, encabezado, filas, actualizar):
    cfg = IMPORTABLES[tipo]
    clave = cfg['clave']
    cols_bd = columnas(tipo)
    mapa, ignoradas = _mapa_columnas(cfg, encabezado)
    if clave not in mapa.values():
        raise ValueError(f'Falta la columna obligatoria "{dict((c[0], c[1]) for c in cfg["campos"])[clave]}". '
                         'Descarga la plantilla y usa sus títulos.')
    campos = {c[0]: c for c in cfg['campos'] if c[0] in cols_bd}
    if len(filas) > MAX_FILAS_IMPORTAR:
        raise ValueError(f'El archivo tiene {len(filas)} filas; el máximo por carga es {MAX_FILAS_IMPORTAR}.')
    claves_archivo = set()
    resultado = []
    for n, fila in enumerate(filas, start=2):
        datos, msgs, error = {}, [], False
        for i, campo in mapa.items():
            if campo not in campos:
                continue
            v, e = _valor(campos[campo][3], fila[i] if i < len(fila) else None)
            if e:
                msgs.append(e)
                error = True
            elif v is not None:
                datos[campo] = v
        if tipo == 'clientes' and datos.get('numero_documento') and 'tipo_documento' not in datos:
            datos['tipo_documento'] = 'RUC' if len(datos['numero_documento']) == 11 else 'DNI'
        if tipo == 'proveedores' and datos.get('ruc') and len(datos['ruc']) != 11:
            msgs.append('El RUC de un proveedor debe tener 11 dígitos')
            error = True
        # Obligatorios: solo se exigen si el registro es NUEVO (se decide más abajo)
        faltan = [(campo, etiqueta, defecto) for campo, etiqueta, oblig, _t, defecto in campos.values()
                  if oblig and campo not in datos and campo != clave]
        if clave not in datos and not error:
            msgs.append(f'Falta {campos[clave][1]}')
            error = True
        k = datos.get(clave)
        if k:
            k_norm = str(k).strip().upper()
            if k_norm in claves_archivo:
                msgs.append('Repetido dentro del mismo archivo (se usa solo la primera fila)')
                error = True
            claves_archivo.add(k_norm)
        resultado.append({'fila': n, 'clave': k, 'datos': datos, 'mensajes': msgs, 'estado': 'error' if error else 'nuevo',
                          '_faltan': faltan})
    existentes = {}
    if claves_archivo:
        for f in db_query(f"SELECT id, UPPER(TRIM({clave}::text)) AS k FROM {tipo} WHERE UPPER(TRIM({clave}::text)) = ANY(%s)",
                          (list(claves_archivo),)):
            existentes.setdefault(f['k'], []).append(f['id'])
    for r in resultado:
        if r['estado'] == 'error' or not r['clave']:
            continue
        ids = existentes.get(str(r['clave']).strip().upper())
        if ids:
            if len(ids) > 1:
                r['estado'] = 'error'
                r['mensajes'].append(f'Ya hay {len(ids)} registros con este código/RUC en el sistema: corrígelo primero en Maestros')
            elif actualizar:
                r['estado'] = 'actualizar'
                r['id'] = ids[0]
            else:
                r['estado'] = 'existe'
                r['mensajes'].append('Ya existe: no se toca (marca "Actualizar existentes" si quieres actualizarlo)')
        else:  # nuevo: aquí sí se exigen los obligatorios
            for campo, etiqueta, defecto in r['_faltan']:
                if defecto is not None:
                    r['datos'][campo] = defecto
                    r['mensajes'].append(f'{etiqueta} vacío: se usará "{defecto}"')
                else:
                    r['mensajes'].append(f'Falta {etiqueta}')
                    r['estado'] = 'error'
    for r in resultado:
        r.pop('_faltan', None)
    resumen = {e: sum(1 for r in resultado if r['estado'] == e) for e in ('nuevo', 'actualizar', 'existe', 'error')}
    return resultado, resumen, [campos[c][1] for c in mapa.values() if c in campos], ignoradas


def _hay_trigger(cur, tabla, texto):
    cur.execute("""SELECT 1 FROM pg_trigger WHERE tgrelid = to_regclass(%s) AND NOT tgisinternal AND tgname ILIKE %s""",
                (tabla, f'%{texto}%'))
    return cur.fetchone() is not None


def _codigo_siguiente(cur, tabla, columna, prefijo):
    cur.execute(f"""SELECT COALESCE(MAX(NULLIF(regexp_replace({columna}, '\\D', '', 'g'), '')::bigint), 0) + 1 AS n
                    FROM {tabla} WHERE {columna} LIKE %s""", (prefijo + '%',))
    return f"{prefijo}{str(cur.fetchone()['n']).zfill(6)}"


def ejecutar_importacion(tipo, resultado, actualizar):
    cols_bd = columnas(tipo)
    clave = IMPORTABLES[tipo]['clave']
    nuevos = actualizados = 0
    with db_tx() as conn:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        marcar_usuario(cur, origen='Importación desde Herramientas')
        trig_cli = tipo == 'clientes' and _hay_trigger(cur, 'clientes', 'codigo')
        trig_prov = tipo == 'proveedores' and _hay_trigger(cur, 'proveedores', 'codigo')
        for r in resultado:
            datos = dict(r['datos'])
            if r['estado'] == 'nuevo':
                if 'activo' in cols_bd:
                    datos['activo'] = True
                if tipo == 'productos':
                    datos.setdefault('costo_unitario', 0)
                    datos.setdefault('precio_unitario', 0)
                    if 'stock' in cols_bd:
                        datos['stock'] = 0
                if tipo == 'clientes' and 'codigo_cliente' in cols_bd and not trig_cli:
                    datos['codigo_cliente'] = _codigo_siguiente(cur, 'clientes', 'codigo_cliente', 'CLI-')
                if tipo == 'proveedores' and 'codigo_proveedor' in cols_bd and not trig_prov:
                    datos['codigo_proveedor'] = _codigo_siguiente(cur, 'proveedores', 'codigo_proveedor', 'PROV-')
                datos = {k: v for k, v in datos.items() if k in cols_bd}
                cols = list(datos)
                cur.execute(f"INSERT INTO {tipo} ({', '.join(cols)}) VALUES ({', '.join(['%s'] * len(cols))}) RETURNING id",
                            [datos[c] for c in cols])
                r['id'] = cur.fetchone()['id']
                nuevos += 1
            elif r['estado'] == 'actualizar' and actualizar:
                datos = {k: v for k, v in datos.items() if k in cols_bd and k != clave}
                if not datos:
                    continue
                sets = ', '.join(f'{c} = %s' for c in datos)
                if 'updated_at' in cols_bd:
                    sets += ', updated_at = NOW()'
                cur.execute(f"UPDATE {tipo} SET {sets} WHERE id = %s", list(datos.values()) + [r['id']])
                actualizados += cur.rowcount
    return nuevos, actualizados


@herramientas_bp.route('/herramientas/api/importar/<tipo>', methods=['POST'])
def api_importar(tipo):
    if tipo not in IMPORTABLES:
        return _err('Solo se puede importar productos, clientes o proveedores.')
    archivo = request.files.get('archivo')
    if not archivo or not archivo.filename:
        return _err('Elige el archivo Excel o CSV.')
    actualizar = str(request.form.get('actualizar', '')).lower() in ('1', 'true', 'on', 'si', 'sí')
    confirmar = str(request.form.get('confirmar', '')).lower() in ('1', 'true', 'on', 'si', 'sí')
    try:
        encabezado, filas = leer_archivo(archivo)
        resultado, resumen, reconocidas, ignoradas = validar_importacion(tipo, encabezado, filas, actualizar)
    except ValueError as e:
        return _err(str(e))
    except Exception as e:
        logging.getLogger(__name__).exception('Error leyendo importación')
        return _err(f'No se pudo leer el archivo: {e}')
    salida = {'success': True, 'resumen': resumen, 'columnas': reconocidas, 'ignoradas': ignoradas,
              'filas': [{**r, 'datos': {k: _plano(v) for k, v in r['datos'].items()}} for r in resultado[:1000]],
              'total_filas': len(resultado)}
    if not confirmar:
        return jsonify(salida)
    try:
        nuevos, actualizados = ejecutar_importacion(tipo, resultado, actualizar)
    except Exception as e:
        logging.getLogger(__name__).exception('Error en importación')
        return _err(f'No se importó nada (se deshizo todo): {e}', 500)
    detalle = (f'Importó {IMPORTABLES[tipo]["titulo"]} desde "{archivo.filename}": {nuevos} nuevos, '
               f'{actualizados} actualizados, {resumen["existe"]} ya existían, {resumen["error"]} con error')
    registrar_auditoria('importar', MODULO, tabla=tipo, detalle=detalle,
                        despues={'archivo': archivo.filename, 'nuevos': nuevos, 'actualizados': actualizados, **resumen})
    salida.update(message=detalle, nuevos=nuevos, actualizados=actualizados)
    return jsonify(salida)


# ============================================================
# 2. EXPORTAR (respaldo por tabla)
# ============================================================
EXPORTABLES = {
    'clientes': 'Clientes', 'clientes_contactos': 'Contactos de clientes', 'clientes_puntos_entrega': 'Puntos de entrega',
    'proveedores': 'Proveedores', 'productos': 'Productos', 'almacenes': 'Sedes / almacenes',
    'cotizaciones': 'Cotizaciones', 'cotizacion_detalle': 'Detalle de cotizaciones',
    'cotizaciones_eliminadas': 'Cotizaciones eliminadas (respaldo)',
    'pedido_compra_pc': 'Pedidos de compra del cliente (PC)', 'despachos': 'Despachos',
    'guias_remision': 'Guías de remisión', 'comprobantes': 'Facturas y boletas', 'notas_credito': 'Notas de crédito',
    'devoluciones': 'Devoluciones', 'ordenes_compra': 'Órdenes de compra', 'ordenes_compra_items': 'Ítems de órdenes de compra',
    'solicitudes_compra': 'Solicitudes de compra', 'recepciones_mercaderia': 'Recepciones de mercadería',
    'comprobantes_proveedor': 'Facturas de proveedores', 'finanzas_pagos': 'Cobros y pagos',
    'inv_stock': 'Stock por empresa y sede', 'inv_movimientos': 'Kardex (movimientos de inventario)',
    'inv_traslados': 'Traslados', 'transportistas': 'Transportistas', 'erp_empresas': 'Empresas',
    'usuarios': 'Usuarios (sin contraseñas)', 'erp_auditoria': 'Auditoría',
}
# Columnas que nunca se exportan
RE_SECRETO = re.compile(r'(password|contrasena|clave|secret|token|api_?key|hash|salt|credencial)', re.I)
MAX_FILAS_EXPORTAR = 200000


def filas_tabla(tabla):
    cols = [c for c in db_query("""SELECT column_name FROM information_schema.columns
                                   WHERE table_schema = current_schema() AND table_name = %s ORDER BY ordinal_position""",
                                (tabla,))]
    nombres = [c['column_name'] for c in cols if not RE_SECRETO.search(c['column_name'])]
    if not nombres:
        return [], []
    orden = 'id' if 'id' in nombres else nombres[0]
    filas = db_query(f"SELECT {', '.join(nombres)} FROM {tabla} ORDER BY {orden} LIMIT {MAX_FILAS_EXPORTAR}")
    return nombres, filas


def respuesta_tabla(nombre_archivo, columnas_, filas, formato, titulo='Datos'):
    if formato == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook(write_only=False)
        ws = wb.active
        ws.title = titulo[:31]
        ws.append(columnas_)
        for c in ws[1]:
            c.font = Font(bold=True)
        for f in filas:
            ws.append([_celda(f.get(c)) for c in columnas_])
        salida = io.BytesIO()
        wb.save(salida)
        return Response(salida.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        headers={'Content-Disposition': f'attachment; filename={nombre_archivo}.xlsx'})
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=';')
    w.writerow(columnas_)
    for f in filas:
        w.writerow([_celda(f.get(c)) for c in columnas_])
    return Response('﻿' + buf.getvalue(), mimetype='text/csv; charset=utf-8',
                    headers={'Content-Disposition': f'attachment; filename={nombre_archivo}.csv'})


def _celda(v):
    v = _plano(v)
    if isinstance(v, str) and v[:1] in ('=', '+', '-', '@') and not re.match(r'^-?\d', v):
        v = "'" + v  # evita fórmulas al abrir en Excel
    return v


@herramientas_bp.route('/herramientas/api/exportar')
def api_exportar():
    tabla = request.args.get('tabla', '')
    formato = 'xlsx' if request.args.get('formato') == 'xlsx' else 'csv'
    if tabla not in EXPORTABLES:
        return _err('Esa tabla no se puede exportar.')
    if not existe_tabla(tabla):
        return _err('Esa tabla todavía no existe en esta base.', 404)
    cols, filas = filas_tabla(tabla)
    registrar_auditoria('exportar', MODULO, tabla=tabla,
                        detalle=f'Exportó {EXPORTABLES[tabla]} ({len(filas)} filas, {formato.upper()})')
    return respuesta_tabla(f'respaldo_{tabla}_{datetime.now():%Y%m%d_%H%M}', cols, filas, formato, EXPORTABLES[tabla])


# ============================================================
# 3. NUMERACIONES (solo revisa; nunca renumera)
# ============================================================
NUMERACIONES = [
    # (código, título, tabla, columna número, columnas que separan la serie, es_fiscal)
    ('cotizaciones', 'Cotizaciones', 'cotizaciones', 'numero_cotizacion', [], False),
    ('pc', 'Pedidos de compra (PC)', 'pedido_compra_pc', 'numero', [], False),
    ('despachos', 'Despachos', 'despachos', 'numero', [], False),
    ('oc', 'Órdenes de compra', 'ordenes_compra', 'numero_orden', [], False),
    ('guias', 'Guías de remisión', 'guias_remision', 'numero', ['serie', 'empresa_id'], True),
    ('facturas', 'Facturas y boletas', 'comprobantes', 'numero', ['tipo_comprobante', 'serie', 'empresa_id'], True),
    ('nc', 'Notas de crédito', 'notas_credito', 'numero', ['serie', 'empresa_id'], True),
]
MAX_HUECOS = 30


def revisar_numeracion(codigo, titulo, tabla, col, series, fiscal):
    info = {'codigo': codigo, 'titulo': titulo, 'fiscal': fiscal, 'series': [], 'huecos': 0, 'duplicados': 0,
            'documentos': 0}
    if not existe_tabla(tabla):
        info['nota'] = 'La tabla no existe en esta base.'
        return info
    cols = columnas(tabla)
    series = [s for s in series if s in cols]
    if col not in cols:
        info['nota'] = 'Sin columna de número.'
        return info
    sel = ''.join(f"COALESCE({s}::text, '') AS s_{s}, " for s in series)
    grp = ''.join(f"COALESCE({s}::text, ''), " for s in series)
    pre = f"regexp_replace({col}::text, '\\d+$', '')"
    num = f"NULLIF(substring({col}::text from '(\\d+)$'), '')::bigint"
    # prefijo = lo que va antes de los últimos dígitos (ej. "PC-20261016-"), n = los últimos dígitos
    filas = db_query(f"""
        SELECT {sel}{pre} AS prefijo, {num} AS n, COUNT(*) AS veces, array_agg(id ORDER BY id) AS ids
        FROM {tabla}
        WHERE {col} IS NOT NULL AND {col}::text <> ''
        GROUP BY {grp}{pre}, {num}
        ORDER BY {grp}{pre}, {num}""")
    empresas_cod = {}
    try:
        empresas_cod = {str(e['id']): e['codigo'] for e in db_query("SELECT id, codigo FROM erp_empresas")}
    except Exception:
        pass
    grupos = {}
    for f in filas:
        partes = []
        for s in series:
            v = f.get(f's_{s}') or ''
            partes.append(empresas_cod.get(v, v or 'sin empresa') if s == 'empresa_id' else v)
        partes.append(f['prefijo'])
        clave = ' · '.join(p for p in partes if p) or '(sin serie)'
        grupos.setdefault(clave, []).append(f)
    for clave, lista in grupos.items():
        nums = sorted(f['n'] for f in lista if f['n'] is not None)
        dups = [{'numero': f"{lista[0]['prefijo']}{f['n']}", 'veces': f['veces'], 'ids': f['ids']} for f in lista if f['veces'] > 1]
        huecos, total_huecos = [], 0
        for a, b in zip(nums, nums[1:]):
            if b - a > 1:
                total_huecos += b - a - 1
                for x in range(a + 1, b):
                    if len(huecos) < MAX_HUECOS:
                        huecos.append(x)
        info['series'].append({'serie': clave, 'desde': nums[0] if nums else None, 'hasta': nums[-1] if nums else None,
                               'documentos': sum(f['veces'] for f in lista), 'huecos': huecos, 'total_huecos': total_huecos,
                               'duplicados': dups})
        info['huecos'] += total_huecos
        info['duplicados'] += len(dups)
        info['documentos'] += sum(f['veces'] for f in lista)
    # Series diarias (PC-AAAAMMDD-NNNN): solo se muestran las que tienen algo raro, para no llenar la pantalla
    if len(info['series']) > 15:
        raras = [s for s in info['series'] if s['total_huecos'] or s['duplicados']]
        info['series_ocultas'] = len(info['series']) - len(raras)
        info['series'] = raras
    return info


def correlativos_contador():
    """Contadores de Configuración ▸ Correlativos que quedaron por debajo del último número usado."""
    if not existe_tabla('erp_correlativos'):
        return []
    salida = []
    filas = db_query("""SELECT c.id, c.documento, c.prefijo, c.anio, c.ultimo_numero, e.codigo AS empresa
                        FROM erp_correlativos c LEFT JOIN erp_empresas e ON e.id = c.empresa_id
                        WHERE COALESCE(c.estado, 'activo') = 'activo' ORDER BY e.codigo, c.documento""")
    fuentes = [(t, c) for _, _, t, c, _, _ in NUMERACIONES if existe_tabla(t)]
    for f in filas:
        patron = f"{f['prefijo']}-{f['anio']}-%"
        maximo = 0
        for t, c in fuentes:
            r = db_query(f"""SELECT COALESCE(MAX(NULLIF(substring({c}::text from '(\\d+)$'), '')::bigint), 0) AS m
                             FROM {t} WHERE {c}::text LIKE %s""", (patron,))
            maximo = max(maximo, int(r[0]['m'] or 0))
        salida.append({'id': str(f['id']), 'documento': f['documento'], 'prefijo': f['prefijo'], 'anio': f['anio'],
                       'empresa': f['empresa'], 'ultimo_numero': f['ultimo_numero'], 'maximo_usado': maximo,
                       'atrasado': maximo > (f['ultimo_numero'] or 0)})
    return salida


@herramientas_bp.route('/herramientas/api/correlativos')
def api_correlativos():
    reporte = [revisar_numeracion(*n) for n in NUMERACIONES]
    contadores = correlativos_contador()
    registrar_auditoria('herramienta', MODULO, detalle='Revisó las numeraciones (huecos y repetidos)')
    return jsonify({'success': True, 'data': reporte, 'contadores': contadores})


@herramientas_bp.route('/herramientas/api/correlativos/ajustar', methods=['POST'])
def api_ajustar_contador():
    """Sube el "último número" de un contador al máximo ya usado. Nunca lo baja ni toca documentos."""
    cid = str((request.get_json(silent=True) or {}).get('id') or '')
    fila = next((c for c in correlativos_contador() if c['id'] == cid), None)
    if not fila:
        return _err('Contador no encontrado.', 404)
    if not fila['atrasado']:
        return _err('Ese contador ya está bien: no hace falta ajustarlo.')
    with db_tx() as conn:
        cur = conn.cursor()
        marcar_usuario(cur, origen='Herramientas')
        cur.execute("""UPDATE erp_correlativos SET ultimo_numero = %s WHERE id::text = %s AND ultimo_numero < %s""",
                    (fila['maximo_usado'], cid, fila['maximo_usado']))
    registrar_auditoria('editar', MODULO, tabla='erp_correlativos', registro_id=cid,
                        antes={'ultimo_numero': fila['ultimo_numero']}, despues={'ultimo_numero': fila['maximo_usado']},
                        detalle=f"Ajustó el contador {fila['prefijo']}-{fila['anio']} de {fila['ultimo_numero']} a {fila['maximo_usado']}")
    return jsonify({'success': True, 'message': f"Listo: el próximo número será {fila['maximo_usado'] + 1}."})


# ============================================================
# 4. INTEGRIDAD (solo reporta)
# ============================================================
# (código, título, explicación, gravedad, tablas/columnas necesarias, SQL que devuelve id + texto)
CHEQUEOS = [
    ('det_sin_cot', 'Detalle de cotización sin cotización', 'Líneas de productos cuya cotización ya no existe.', 'alta',
     {'cotizacion_detalle': ['cotizacion_id'], 'cotizaciones': []},
     "SELECT d.id, 'Línea ' || d.id || ' → cotización ' || d.cotizacion_id AS texto FROM cotizacion_detalle d "
     "WHERE NOT EXISTS (SELECT 1 FROM cotizaciones c WHERE c.id = d.cotizacion_id)"),
    ('det_sin_prod', 'Cotización con producto que no existe', 'Líneas que apuntan a un producto borrado.', 'media',
     {'cotizacion_detalle': ['producto_id'], 'productos': []},
     "SELECT d.id, 'Línea ' || d.id || ' de la cotización ' || d.cotizacion_id || ' → producto ' || d.producto_id AS texto "
     "FROM cotizacion_detalle d WHERE d.producto_id IS NOT NULL AND NOT EXISTS (SELECT 1 FROM productos p WHERE p.id = d.producto_id)"),
    ('cot_sin_cliente', 'Cotización sin cliente', 'Cotizaciones cuyo cliente no existe.', 'media',
     {'cotizaciones': ['cliente_id'], 'clientes': []},
     "SELECT c.id, c.numero_cotizacion || ' → cliente ' || COALESCE(c.cliente_id::text, '(vacío)') AS texto FROM cotizaciones c "
     "WHERE c.cliente_id IS NOT NULL AND NOT EXISTS (SELECT 1 FROM clientes x WHERE x.id = c.cliente_id)"),
    ('pc_sin_cot', 'PC sin cotización existente', 'Pedidos (PC) que dicen venir de una cotización que no existe.', 'alta',
     {'pedido_compra_pc': ['cotizacion_id', 'cotizacion_numero'], 'cotizaciones': []},
     "SELECT p.id, p.numero || ' → ' || COALESCE(p.cotizacion_numero, 'cotización ' || p.cotizacion_id) AS texto FROM pedido_compra_pc p "
     "WHERE (p.cotizacion_id IS NOT NULL AND NOT EXISTS (SELECT 1 FROM cotizaciones c WHERE c.id = p.cotizacion_id)) "
     "OR (p.cotizacion_id IS NULL AND COALESCE(p.cotizacion_numero, '') <> '' AND NOT EXISTS "
     "(SELECT 1 FROM cotizaciones c WHERE c.numero_cotizacion = p.cotizacion_numero OR c.codigo_cotizacion = p.cotizacion_numero))"),
    ('desp_sin_pc', 'Despacho sin PC existente', 'Despachos que dicen venir de un PC que no existe.', 'alta',
     {'despachos': ['pc_id', 'pc_numero'], 'pedido_compra_pc': []},
     "SELECT d.id, d.numero || ' → ' || COALESCE(d.pc_numero, 'PC ' || d.pc_id) AS texto FROM despachos d "
     "WHERE (d.pc_id IS NOT NULL AND NOT EXISTS (SELECT 1 FROM pedido_compra_pc p WHERE p.id = d.pc_id)) "
     "OR (d.pc_id IS NULL AND COALESCE(d.pc_numero, '') <> '' AND NOT EXISTS (SELECT 1 FROM pedido_compra_pc p WHERE p.numero = d.pc_numero))"),
    ('despprod_sin_desp', 'Productos de despacho sin despacho', 'Líneas de despacho cuyo despacho no existe.', 'media',
     {'despacho_productos': ['despacho_id'], 'despachos': []},
     "SELECT x.id, 'Línea ' || x.id || ' → despacho ' || x.despacho_id AS texto FROM despacho_productos x "
     "WHERE NOT EXISTS (SELECT 1 FROM despachos d WHERE d.id = x.despacho_id)"),
    ('ocitem_sin_oc', 'Ítems de OC sin orden de compra', 'Líneas de órdenes de compra cuya orden no existe.', 'alta',
     {'ordenes_compra_items': ['orden_compra_id'], 'ordenes_compra': []},
     "SELECT i.id, 'Ítem ' || i.id || ' → OC ' || i.orden_compra_id AS texto FROM ordenes_compra_items i "
     "WHERE NOT EXISTS (SELECT 1 FROM ordenes_compra o WHERE o.id = i.orden_compra_id)"),
    ('ocdet_sin_oc', 'Detalle de OC (antiguo) sin orden', 'Líneas antiguas de OC cuya orden no existe.', 'media',
     {'orden_compra_detalle': ['orden_id'], 'ordenes_compra': []},
     "SELECT i.id, 'Línea ' || i.id || ' → OC ' || i.orden_id AS texto FROM orden_compra_detalle i "
     "WHERE NOT EXISTS (SELECT 1 FROM ordenes_compra o WHERE o.id = i.orden_id)"),
    ('factprov_sin_oc', 'Factura de proveedor con OC inexistente', 'Facturas de proveedor ligadas a una OC que no existe.', 'media',
     {'comprobantes_proveedor': ['orden_compra_id'], 'ordenes_compra': []},
     "SELECT f.id, f.numero || ' → OC ' || f.orden_compra_id AS texto FROM comprobantes_proveedor f "
     "WHERE f.orden_compra_id IS NOT NULL AND NOT EXISTS (SELECT 1 FROM ordenes_compra o WHERE o.id = f.orden_compra_id)"),
    ('rec_sin_oc', 'Recepción con OC inexistente', 'Recepciones de mercadería ligadas a una OC que no existe.', 'media',
     {'recepciones_mercaderia': ['orden_compra_id'], 'ordenes_compra': []},
     "SELECT r.id, r.numero_recepcion || ' → OC ' || r.orden_compra_id AS texto FROM recepciones_mercaderia r "
     "WHERE r.orden_compra_id IS NOT NULL AND NOT EXISTS (SELECT 1 FROM ordenes_compra o WHERE o.id = r.orden_compra_id)"),
    ('pago_sin_doc', 'Cobro o pago sin documento', 'Cobros/pagos de Finanzas cuyo documento (factura u OC) no existe.', 'alta',
     {'finanzas_pagos': ['documento_tabla', 'documento_id']},
     "SELECT f.id, f.tipo || ' ' || COALESCE(f.documento_numero, '') || ' (' || f.documento_tabla || ' ' || f.documento_id || ')' AS texto "
     "FROM finanzas_pagos f WHERE (f.documento_tabla = 'comprobantes' AND NOT EXISTS (SELECT 1 FROM comprobantes c WHERE c.id = f.documento_id)) "
     "OR (f.documento_tabla = 'comprobantes_proveedor' AND to_regclass('comprobantes_proveedor') IS NOT NULL AND NOT EXISTS "
     "(SELECT 1 FROM comprobantes_proveedor c WHERE c.id = f.documento_id)) "
     "OR (f.documento_tabla = 'ordenes_compra' AND NOT EXISTS (SELECT 1 FROM ordenes_compra c WHERE c.id = f.documento_id))"),
    ('invstock_sin_prod', 'Stock de un producto que no existe', 'Filas de stock (inv_stock) de productos borrados.', 'alta',
     {'inv_stock': ['producto_id'], 'productos': []},
     "SELECT s.producto_id AS id, 'Producto ' || s.producto_id || ': ' || s.cantidad || ' unid.' AS texto FROM inv_stock s "
     "WHERE NOT EXISTS (SELECT 1 FROM productos p WHERE p.id = s.producto_id)"),
    ('invmov_sin_prod', 'Kardex de un producto que no existe', 'Movimientos de inventario de productos borrados.', 'media',
     {'inv_movimientos': ['producto_id'], 'productos': []},
     "SELECT m.id, COALESCE(m.numero, 'Mov. ' || m.id) || ' → producto ' || m.producto_id AS texto FROM inv_movimientos m "
     "WHERE NOT EXISTS (SELECT 1 FROM productos p WHERE p.id = m.producto_id)"),
    ('contacto_sin_cliente', 'Contacto o punto de entrega sin cliente', 'Datos de clientes que ya no existen.', 'baja',
     {'clientes_puntos_entrega': ['cliente_id'], 'clientes': []},
     "SELECT x.id, 'Punto ' || x.id || ' → cliente ' || x.cliente_id AS texto FROM clientes_puntos_entrega x "
     "WHERE NOT EXISTS (SELECT 1 FROM clientes c WHERE c.id = x.cliente_id)"),
    ('stock_no_cuadra', 'Stock del producto distinto al de Inventario', 'productos.stock no es igual a la suma de inv_stock '
     '(se arregla con "Recalcular stock").', 'alta',
     {'productos': ['stock'], 'inv_stock': ['producto_id']},
     "SELECT p.id, p.codigo || ': dice ' || COALESCE(p.stock, 0) || ', Inventario suma ' || s.total AS texto FROM productos p "
     "JOIN (SELECT producto_id, SUM(cantidad) AS total FROM inv_stock GROUP BY producto_id) s ON s.producto_id = p.id "
     "WHERE COALESCE(p.stock, 0) <> s.total"),
    ('dup_productos', 'Productos con código repetido', 'Dos o más productos con el mismo código.', 'alta',
     {'productos': ['codigo']},
     "SELECT MIN(id) AS id, UPPER(TRIM(codigo)) || ' (' || COUNT(*) || ' veces, ids ' || string_agg(id::text, ', ') || ')' AS texto "
     "FROM productos WHERE COALESCE(TRIM(codigo), '') <> '' GROUP BY UPPER(TRIM(codigo)) HAVING COUNT(*) > 1"),
    ('dup_clientes', 'Clientes con RUC/DNI repetido', 'Dos o más clientes con el mismo documento.', 'media',
     {'clientes': ['numero_documento']},
     "SELECT MIN(id) AS id, TRIM(numero_documento) || ' (' || COUNT(*) || ' veces, ids ' || string_agg(id::text, ', ') || ')' AS texto "
     "FROM clientes WHERE COALESCE(TRIM(numero_documento), '') <> '' GROUP BY TRIM(numero_documento) HAVING COUNT(*) > 1"),
    ('dup_proveedores', 'Proveedores con RUC repetido', 'Dos o más proveedores con el mismo RUC.', 'media',
     {'proveedores': ['ruc']},
     "SELECT MIN(id) AS id, TRIM(ruc) || ' (' || COUNT(*) || ' veces, ids ' || string_agg(id::text, ', ') || ')' AS texto "
     "FROM proveedores WHERE COALESCE(TRIM(ruc), '') <> '' GROUP BY TRIM(ruc) HAVING COUNT(*) > 1"),
    ('dup_cotizaciones', 'Cotizaciones con número repetido', 'Dos o más cotizaciones con el mismo número.', 'alta',
     {'cotizaciones': ['numero_cotizacion']},
     "SELECT MIN(id) AS id, numero_cotizacion || ' (' || COUNT(*) || ' veces)' AS texto FROM cotizaciones "
     "WHERE COALESCE(numero_cotizacion, '') <> '' GROUP BY numero_cotizacion HAVING COUNT(*) > 1"),
    ('dup_facturas', 'Facturas con número repetido', 'Mismo tipo, serie y número en la misma empresa.', 'alta',
     {'comprobantes': ['tipo_comprobante', 'serie', 'numero', 'empresa_id']},
     "SELECT MIN(id) AS id, tipo_comprobante || ' ' || serie || '-' || numero || ' (' || COUNT(*) || ' veces)' AS texto FROM comprobantes "
     "GROUP BY tipo_comprobante, serie, numero, empresa_id HAVING COUNT(*) > 1"),
    ('dup_guias', 'Guías con número repetido', 'Misma serie y número en la misma empresa.', 'alta',
     {'guias_remision': ['serie', 'numero', 'empresa_id']},
     "SELECT MIN(id) AS id, serie || '-' || numero || ' (' || COUNT(*) || ' veces)' AS texto FROM guias_remision "
     "GROUP BY serie, numero, empresa_id HAVING COUNT(*) > 1"),
]
MAX_EJEMPLOS = 25


def ejecutar_chequeos():
    salida = []
    for codigo, titulo, ayuda, gravedad, requisitos, sql in CHEQUEOS:
        r = {'codigo': codigo, 'titulo': titulo, 'ayuda': ayuda, 'gravedad': gravedad, 'cantidad': 0, 'ejemplos': []}
        faltan = []
        for tabla, cols in requisitos.items():
            cb = columnas(tabla)
            if not cb:
                faltan.append(tabla)
            faltan += [f'{tabla}.{c}' for c in cols if cb and c not in cb]
        if faltan:
            r['omitido'] = 'No aplica en esta base (falta ' + ', '.join(faltan) + ')'
            salida.append(r)
            continue
        try:
            filas = db_query(f"SELECT * FROM ({sql}) q ORDER BY 1 LIMIT 1000")
            r['cantidad'] = len(filas)
            r['ejemplos'] = [{'id': f['id'], 'texto': f['texto']} for f in filas[:MAX_EJEMPLOS]]
        except Exception as e:
            r['omitido'] = f'No se pudo revisar: {e}'
        salida.append(r)
    return salida


@herramientas_bp.route('/herramientas/api/integridad')
def api_integridad():
    data = ejecutar_chequeos()
    problemas = sum(1 for c in data if c['cantidad'])
    registrar_auditoria('herramienta', MODULO, detalle=f'Verificó la integridad: {problemas} revisión(es) con problemas')
    return jsonify({'success': True, 'data': data, 'problemas': problemas})


# ============================================================
# 5. RECALCULAR STOCK
# ============================================================
@herramientas_bp.route('/herramientas/api/recalcular_stock', methods=['POST'])
def api_recalcular_stock():
    from routes.inventario import asegurar_tablas_inventario
    asegurar_tablas_inventario()  # crea fn_recalcular_stock_producto si falta
    d = request.get_json(silent=True) or {}
    pid = d.get('producto_id')
    codigo = (d.get('codigo') or '').strip()
    with db_tx() as conn:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        marcar_usuario(cur, origen='Herramientas · Recalcular stock')
        if codigo and not pid:
            cur.execute("SELECT id FROM productos WHERE UPPER(codigo) = UPPER(%s) ORDER BY id LIMIT 1", (codigo,))
            f = cur.fetchone()
            if not f:
                return _err(f'No existe el producto {codigo}.', 404)
            pid = f['id']
        filtro, params = ('AND p.id = %s', (int(pid),)) if pid else ('', ())
        # Solo productos que ya están en Inventario: los demás conservan su stock antiguo
        cur.execute(f"""SELECT p.id, p.codigo, COALESCE(p.stock, 0) AS antes FROM productos p
                        WHERE EXISTS (SELECT 1 FROM inv_stock s WHERE s.producto_id = p.id) {filtro} ORDER BY p.id""", params)
        lista = cur.fetchall()
        if pid and not lista:
            return jsonify({'success': True, 'cambiados': [], 'revisados': 0,
                            'message': 'Ese producto no tiene stock en Inventario: no se cambió nada.'})
        cambiados = []
        for p in lista:
            cur.execute("SELECT fn_recalcular_stock_producto(%s)", (p['id'],))
            cur.execute("SELECT COALESCE(stock, 0) AS despues FROM productos WHERE id = %s", (p['id'],))
            despues = cur.fetchone()['despues']
            if Decimal(str(despues)) != Decimal(str(p['antes'])):
                cambiados.append({'id': p['id'], 'codigo': p['codigo'], 'antes': float(p['antes']), 'despues': float(despues)})
    msg = f'Se revisaron {len(lista)} producto(s); {len(cambiados)} tenían el stock desactualizado y se corrigieron.'
    registrar_auditoria('herramienta', MODULO, tabla='productos', registro_id=pid, detalle='Recalculó stock: ' + msg,
                        despues={'cambiados': cambiados[:200]})
    return jsonify({'success': True, 'revisados': len(lista), 'cambiados': cambiados, 'message': msg})


# ============================================================
# 6. CACHÉ
# ============================================================
def limpiar_caches(app=None):
    hechos = []
    from routes import multiempresa
    multiempresa._cache.update(t=0, filas=[])
    multiempresa._col_empresa.clear()
    hechos.append('Empresas y colores (se vuelven a leer de la base)')
    try:
        from routes import configuracion_seguridad
        configuracion_seguridad._cache_upsert.clear()
        hechos.append('Permisos de usuario (forma de guardado)')
    except Exception:
        pass
    # "Estructura revisada": en la próxima pantalla se vuelve a comprobar que existan tablas y columnas
    banderas = [('routes.multiempresa', '_listo'), ('routes.inventario', '_tablas_listas'), ('routes.finanzas', '_tablas_listas'),
                ('routes.compras_flujo', '_tablas_listas'), ('routes.stock_automatico', '_listo'),
                ('routes.productos', '_columnas_listas'), ('routes.auditoria', '_listo')]
    import importlib
    n = 0
    for mod, var in banderas:
        try:
            m = importlib.import_module(mod)
            if hasattr(m, var):
                setattr(m, var, False)
                n += 1
        except Exception:
            pass
    hechos.append(f'Revisión de estructura de {n} módulos (se repite al abrir la próxima pantalla)')
    try:
        from routes import auditoria
        auditoria._tipo_registro.clear()
    except Exception:
        pass
    if app is not None:
        try:
            app.jinja_env.cache.clear()
            hechos.append('Plantillas de pantallas')
        except Exception:
            pass
    return hechos


@herramientas_bp.route('/herramientas/api/cache/limpiar', methods=['POST'])
def api_limpiar_cache():
    from flask import current_app
    hechos = limpiar_caches(current_app)
    registrar_auditoria('herramienta', MODULO, detalle='Limpió la caché: ' + '; '.join(hechos))
    return jsonify({'success': True, 'limpiado': hechos, 'message': 'Caché limpiada.'})


# ============================================================
# 7. REGISTRO DE ERRORES (logs)
# ============================================================
# main.py solo usa logging.basicConfig (sale por consola, sin archivo). Aquí se agrega, sin tocar main.py:
#   - un registro en memoria con los últimos 2000 mensajes (WARNING o más graves, y los INFO de la app),
#   - un archivo rotativo erp.log (2 MB x 5) si se define la variable ERP_LOG_CARPETA.
class _RegistroMemoria(logging.Handler):
    def __init__(self, capacidad=2000):
        super().__init__(level=logging.INFO)
        self.filas = deque(maxlen=capacidad)
        self.anterior = []

    def emit(self, record):
        try:
            if record.levelno < logging.WARNING and record.name.startswith('werkzeug'):
                return  # cada pedido a la página: demasiado ruido
            self.filas.append({'fecha': datetime.fromtimestamp(record.created).strftime('%d/%m/%Y %H:%M:%S'),
                               'nivel': record.levelname, 'origen': record.name,
                               'mensaje': self.format(record)[:4000]})
        except Exception:
            pass


REGISTRO = None
ARCHIVO_LOG = None


def instalar_registro():
    global REGISTRO, ARCHIVO_LOG
    raiz = logging.getLogger()
    for h in raiz.handlers:
        if isinstance(h, _RegistroMemoria):
            REGISTRO = h
    if REGISTRO is None:
        REGISTRO = _RegistroMemoria()
        REGISTRO.setFormatter(logging.Formatter('%(message)s'))
        raiz.addHandler(REGISTRO)
        if raiz.level > logging.INFO or raiz.level == logging.NOTSET:
            raiz.setLevel(logging.INFO)
    carpeta = os.environ.get('ERP_LOG_CARPETA')  # archivo solo si se configura la carpeta (en Render el disco se borra)
    if ARCHIVO_LOG is None and carpeta:
        try:
            os.makedirs(carpeta, exist_ok=True)
            ARCHIVO_LOG = RotatingFileHandler(os.path.join(carpeta, 'erp.log'), maxBytes=2_000_000, backupCount=5,
                                              encoding='utf-8', delay=True)
            ARCHIVO_LOG.setLevel(logging.WARNING)
            ARCHIVO_LOG.setFormatter(logging.Formatter('%(asctime)s %(levelname)s %(name)s: %(message)s'))
            raiz.addHandler(ARCHIVO_LOG)
        except Exception:
            ARCHIVO_LOG = None


instalar_registro()


@herramientas_bp.route('/herramientas/api/logs')
def api_logs():
    try:
        n = max(1, min(int(request.args.get('n', 200)), 2000))
    except ValueError:
        n = 200
    nivel = (request.args.get('nivel') or 'WARNING').upper()
    minimo = getattr(logging, nivel, logging.WARNING) if nivel != 'TODOS' else 0
    filas = [f for f in REGISTRO.filas if getattr(logging, f['nivel'], 0) >= minimo] if REGISTRO else []
    filas = list(reversed(filas))[:n]
    registrar_auditoria('herramienta', MODULO, detalle=f'Vio el registro de errores (últimos {n}, nivel {nivel})')
    return jsonify({'success': True, 'data': filas, 'total_memoria': len(REGISTRO.filas) if REGISTRO else 0,
                    'anterior': len(REGISTRO.anterior) if REGISTRO else 0,
                    'archivo': bool(ARCHIVO_LOG), 'nota': 'Algunos módulos antiguos escriben sus errores con print(): esos '
                    'solo se ven en la consola del servidor (Render ▸ Logs).'})


@herramientas_bp.route('/herramientas/api/logs/rotar', methods=['POST'])
def api_rotar_logs():
    n = 0
    if REGISTRO:
        n = len(REGISTRO.filas)
        REGISTRO.anterior = list(REGISTRO.filas)  # se guarda la última tanda por si acaso
        REGISTRO.filas.clear()
    archivo = False
    if ARCHIVO_LOG:
        try:
            ARCHIVO_LOG.doRollover()
            archivo = True
        except Exception:
            pass
    detalle = f'Rotó el registro de errores ({n} mensajes archivados' + (', archivo rotado)' if archivo else ')')
    registrar_auditoria('herramienta', MODULO, detalle=detalle)
    logging.getLogger(__name__).warning('Registro rotado por %s', session.get('usuario'))
    return jsonify({'success': True, 'message': detalle})


# ============================================================
# 8. ESTADO DE LA BASE (nunca muestra la cadena de conexión)
# ============================================================
TABLAS_CLAVE = ['usuarios', 'clientes', 'proveedores', 'productos', 'cotizaciones', 'pedido_compra_pc', 'despachos',
                'guias_remision', 'comprobantes', 'notas_credito', 'ordenes_compra', 'inv_movimientos', 'finanzas_pagos',
                'erp_auditoria']


@herramientas_bp.route('/herramientas/api/salud')
def api_salud():
    t0 = time.perf_counter()
    try:
        db_query("SELECT 1")
    except Exception:
        return jsonify({'success': False, 'ok': False, 'error': 'No hay conexión con la base de datos.'}), 503
    latencia = round((time.perf_counter() - t0) * 1000, 1)
    info = db_query("""SELECT version() AS version, current_setting('TimeZone') AS zona, now() AS hora,
                              pg_size_pretty(pg_database_size(current_database())) AS tamano,
                              (SELECT COUNT(*) FROM pg_stat_activity WHERE datname = current_database()) AS conexiones""")[0]
    conteos = []
    for t in TABLAS_CLAVE:
        if existe_tabla(t):
            conteos.append({'tabla': t, 'filas': db_query(f"SELECT COUNT(*) AS n FROM {t}")[0]['n']})
    version = (info['version'] or '').split(' on ')[0].split(',')[0]
    salida = {'success': True, 'ok': True, 'latencia_ms': latencia,
              'estado': 'Rápida' if latencia < 150 else ('Lenta' if latencia < 800 else 'Muy lenta'),
              'version': version, 'zona_horaria': info['zona'], 'hora_base': fecha_txt(info['hora']),
              'hora_servidor': datetime.now().strftime('%d/%m/%Y %H:%M'), 'tamano': info['tamano'],
              'conexiones': info['conexiones'], 'conteos': conteos,
              'python': platform.python_version()}
    registrar_auditoria('herramienta', MODULO, detalle=f'Revisó la conexión a la base ({latencia} ms)')
    return jsonify(salida)
