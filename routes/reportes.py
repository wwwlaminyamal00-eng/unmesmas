# routes/reportes.py - Módulo 8: Reportes y tablero de ventas (pedido de Erika, 16/10/2026)
#
# Reglas (en simple):
#   1. Solo LEE datos que ya existen; no crea ni cambia documentos.
#   2. Ventas = comprobantes de venta emitidos (no anulados, no borradores). Compras = órdenes de compra
#      confirmadas (no anuladas, no borradores, no "por aprobar"). Cobros/pagos = finanzas_pagos no anulados.
#      Cuentas por cobrar/pagar usan las MISMAS consultas que Finanzas (routes/finanzas.py).
#      Inventario usa inv_stock (stock y costo promedio) e inv_movimientos registrados (kardex).
#   3. El vendedor de una venta es el usuario que hizo la cotización (cotizaciones.usuario_id) a la que apunta el
#      comprobante (directo o a través de su PV). Si no tiene cotización, quien registró el comprobante.
#      Solo salen usuarios reales de la tabla usuarios.
#   4. Costo de un producto: costo promedio del inventario (inv_stock); si no hay, productos.costo_unitario.
#   5. Empresa: respeta el selector de arriba (KCF / AGD / Todas). Con "Todas" se puede elegir una en el filtro.
#      Como en el resto del ERP, documentos antiguos sin empresa se ven en todas.
#   6. Acceso: montos solo para administración (ROLES_ADMIN) y Finanzas (ROLES_FINANZAS).
#      El KPI por colaborador solo para ROLES_ADMIN. Se valida en el servidor.
import csv
import io
import json
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal

from flask import Blueprint, render_template, request, jsonify, Response, session

from database import db_query
from utils import a_decimal, es_admin, rol_actual, id_empresa, tasa_igv
from routes.multiempresa import empresa_filtro, empresa_actual, es_todas, empresas as lista_empresas

reportes_bp = Blueprint('reportes', __name__)

ANULADOS = ('ANULADO', 'ANULADA', 'ANULADOS')
VENTA_NO_CUENTA = ANULADOS + ('BORRADOR', 'RECHAZADO', 'RECHAZADA', 'ELIMINADO', 'ELIMINADA')
OC_NO_CUENTA = ANULADOS + ('BORRADOR', 'POR APROBAR', 'RECHAZADA', 'RECHAZADO', 'ELIMINADA')
COT_CONVERTIDA = ('ACEPTADA POR CLIENTE', 'ACEPTADA', 'FACTURADA')
MESES = ('Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Set', 'Oct', 'Nov', 'Dic')


class ErrorReporte(Exception):
    pass


# ============================================================
# PERMISOS (en el servidor, no solo ocultando botones)
# ============================================================
def _roles_finanzas():
    from routes.finanzas import ROLES_FINANZAS
    return ROLES_FINANZAS


def puede_ver_montos():
    return es_admin() or rol_actual() in _roles_finanzas()


def puede_ver_kpi():
    return es_admin()


def _prohibido(msg='No tienes permiso para ver reportes con montos. Pide acceso a la gerencia.'):
    return jsonify({'success': False, 'error': msg}), 403


# ============================================================
# ESTRUCTURA (columnas que pueden faltar en bases antiguas)
# ============================================================
_cols_cache = {}


def _columnas(tabla):
    if tabla not in _cols_cache:
        try:
            _cols_cache[tabla] = {f['column_name'] for f in db_query(
                "SELECT column_name FROM information_schema.columns WHERE table_schema = current_schema() AND table_name = %s",
                (tabla,))}
        except Exception:
            _cols_cache[tabla] = set()
    return _cols_cache[tabla]


def _tiene(tabla, col):
    cols = _columnas(tabla)
    if col not in cols:
        _cols_cache.pop(tabla, None)  # pudo crearse después (otro módulo la agrega al arrancar)
        cols = _columnas(tabla)
    return col in cols


def _existe(tabla):
    return bool(_columnas(tabla))


def _asegurar():
    """Las columnas empresa_id / tablas de inventario y finanzas que los otros módulos crean al arrancar."""
    for mod, fn in (('routes.multiempresa', 'asegurar_columnas_empresa'), ('routes.compras_flujo', 'asegurar_tablas_compras')):
        try:
            getattr(__import__(mod, fromlist=[fn]), fn)()
        except Exception as e:  # nunca bloquear los reportes por esto
            print(f'⚠️ reportes: {fn}: {e}')


# ============================================================
# FILTROS
# ============================================================
def _fecha(txt, por_defecto):
    try:
        return datetime.strptime((txt or '')[:10], '%Y-%m-%d').date()
    except ValueError:
        return por_defecto


def _entero(txt, por_defecto, minimo=1, maximo=100000):
    try:
        return max(minimo, min(int(txt), maximo))
    except (TypeError, ValueError):
        return por_defecto


def leer_filtros(args=None):
    a = args if args is not None else request.args
    hoy = date.today()
    desde = _fecha(a.get('desde'), hoy.replace(day=1))
    hasta = _fecha(a.get('hasta'), hoy)
    if hasta < desde:
        desde, hasta = hasta, desde
    eid = empresa_filtro()
    if eid is None:  # "Todas": se puede elegir una empresa en el filtro
        eid = id_empresa(a.get('empresa_id'))
    return {
        'desde': desde, 'hasta': hasta, 'empresa_id': str(eid) if eid is not None else None,
        'almacen_id': _entero(a.get('almacen_id'), None) if a.get('almacen_id') else None,
        'usuario_id': _entero(a.get('usuario_id'), None) if a.get('usuario_id') else None,
        'cliente': (a.get('cliente') or '').strip(), 'proveedor': (a.get('proveedor') or '').strip(),
        'producto': (a.get('producto') or '').strip(), 'agrupar': (a.get('agrupar') or '').strip(),
        'estado': (a.get('estado') or '').strip(), 'limite': _entero(a.get('limite'), 10, 1, 500),
        'dias': _entero(a.get('dias'), 90, 1, 3650), 'orden': (a.get('orden') or 'monto').strip(),
    }


def _sql_empresa(col, f):
    """Mismo criterio que multiempresa.filtro_sql: las filas sin empresa se ven en todas."""
    if f['empresa_id'] is None:
        return '', ()
    return f' AND ({col} IS NULL OR {col}::text = %s)', (f['empresa_id'],)


def _coincide(texto, *valores):
    t = texto.lower()
    return any(t in (v or '').lower() for v in valores)


def _nombre_empresa(eid):
    if eid is None:
        return 'Sin empresa'
    e = next((x for x in lista_empresas() if x['id'] == str(eid)), None)
    return e['codigo'] if e else str(eid)


# ============================================================
# DATOS DE APOYO
# ============================================================
def _usuarios():
    """{id: nombre} de la tabla usuarios (las únicas personas que pueden salir en los reportes)."""
    cols = _columnas('usuarios')
    partes = [c for c in ('nombres_apellidos', 'nombre_completo', 'usuario_sistema', 'usuario') if c in cols]
    nombre = 'COALESCE(' + ', '.join(f"NULLIF({c}, '')" for c in partes) + ", 'Usuario ' || id)" if partes else "'Usuario ' || id"
    rol = 'rol' if 'rol' in cols else "''"
    estado = "COALESCE(LOWER(estado), 'activo')" if 'estado' in cols else "'activo'"
    return {f['id']: {'nombre': f['nombre'], 'rol': f['rol'] or '', 'estado': f['estado']}
            for f in db_query(f'SELECT id, {nombre} AS nombre, {rol} AS rol, {estado} AS estado FROM usuarios')}


def _productos():
    filas = db_query('SELECT id, codigo, descripcion, unidad, COALESCE(costo_unitario, 0) AS costo FROM productos')
    costo_inv = {}
    if _existe('inv_stock'):
        costo_inv = {f['producto_id']: a_decimal(f['c']) for f in db_query("""
            SELECT producto_id, SUM(valor) / NULLIF(SUM(cantidad), 0) AS c FROM inv_stock
            WHERE cantidad > 0 GROUP BY producto_id""") if f['c']}
    por_id, por_codigo = {}, {}
    for f in filas:
        d = dict(f)
        d['costo'] = costo_inv.get(f['id']) or a_decimal(f['costo'])
        por_id[f['id']] = d
        if f['codigo']:
            por_codigo.setdefault(f['codigo'].strip().upper(), d)
    return por_id, por_codigo


def _items(crudo):
    """Líneas de un items_json (lista de objetos, lista de listas o dict con 'items')."""
    if isinstance(crudo, str):
        try:
            crudo = json.loads(crudo)
        except ValueError:
            return []
    if isinstance(crudo, dict):
        crudo = crudo.get('items') or list(crudo.values())
    salida = []
    for it in crudo or []:
        if isinstance(it, dict):
            cant = a_decimal(it.get('cantidad') if it.get('cantidad') not in (None, '') else it.get('cantidad_pc'))
            precio = Decimal('0')
            for k in ('valorVenta', 'valor_venta', 'precio_unitario', 'precio', 'precio_pc', 'valor_unitario', 'precio_venta'):
                if it.get(k) not in (None, ''):
                    precio = a_decimal(it.get(k))
                    break
            importe = cant * precio
            if not importe:
                importe = a_decimal(it.get('subtotal') or it.get('importe') or it.get('total') or 0)
            salida.append({'producto_id': it.get('producto_id') or it.get('id_producto'),
                           'codigo': str(it.get('codigo') or '').strip(),
                           'descripcion': str(it.get('producto') or it.get('descripcion') or it.get('nombre') or '').strip(),
                           'cantidad': cant, 'importe': importe})
        elif isinstance(it, (list, tuple)) and len(it) > 1:
            cant = a_decimal(it[5] if len(it) > 5 else 1)
            salida.append({'producto_id': None, 'codigo': str(it[0] or '').strip(), 'descripcion': str(it[1] or ''),
                           'cantidad': cant, 'importe': Decimal('0')})
    return salida


# ============================================================
# VENTAS (base de casi todos los reportes comerciales)
# ============================================================
def _ventas(f, desde=None, hasta=None, con_lineas=True, ignorar_cliente=False):
    """Comprobantes de venta del periodo con vendedor, y sus líneas prorrateadas al subtotal (sin IGV)."""
    desde = desde or f['desde']
    hasta = hasta or f['hasta']
    emp_sql, emp_par = _sql_empresa('c.empresa_id', f) if _tiene('comprobantes', 'empresa_id') else ('', ())
    emp_col = 'c.empresa_id' if _tiene('comprobantes', 'empresa_id') else 'NULL'
    filas = db_query(f"""
        SELECT c.id, c.tipo_comprobante, c.serie, c.numero::text AS numero, c.fecha_emision, c.cliente_nombre,
               c.cliente_numero_doc, COALESCE(c.subtotal, 0) AS subtotal, COALESCE(c.igv, 0) AS igv,
               COALESCE(c.total, 0) AS total, c.items_json, {emp_col} AS empresa_id, c.documento_asociado,
               c.pc_vinculado, c.creado_por::text AS creado_por, c.estado_sunat
        FROM comprobantes c
        WHERE COALESCE(UPPER(TRIM(c.estado_sunat)), 'BORRADOR') <> ALL(%s)
          AND c.fecha_emision >= %s AND c.fecha_emision < %s::date + 1 {emp_sql}
        ORDER BY c.fecha_emision, c.id
    """, (list(VENTA_NO_CUENTA), desde, hasta) + emp_par)
    if not filas:
        return []
    usuarios = _usuarios()
    cots = {}
    for c in db_query('SELECT id, numero_cotizacion, usuario_id FROM cotizaciones'):
        cots[('id', c['id'])] = c['usuario_id']
        if c['numero_cotizacion']:
            cots[('num', c['numero_cotizacion'].strip().upper())] = c['usuario_id']
    pcs = {p['numero'].strip().upper(): p for p in db_query(
        'SELECT numero, cotizacion_id, cotizacion_numero FROM pedido_compra_pc WHERE numero IS NOT NULL')}
    por_id, por_codigo = _productos() if con_lineas else ({}, {})
    docs = []
    for c in filas:
        vend = None
        asoc = (c['documento_asociado'] or '').strip().upper()
        if asoc:
            vend = cots.get(('num', asoc))
        pc = pcs.get((c['pc_vinculado'] or '').strip().upper())
        if vend is None and pc:
            vend = cots.get(('id', pc['cotizacion_id'])) or cots.get(('num', (pc['cotizacion_numero'] or '').strip().upper()))
        if vend is None and (c['creado_por'] or '').isdigit():
            vend = int(c['creado_por'])
        if vend not in usuarios:
            vend = None
        if f['usuario_id'] and vend != f['usuario_id']:
            continue
        if f['cliente'] and not ignorar_cliente and not _coincide(f['cliente'], c['cliente_nombre'], c['cliente_numero_doc']):
            continue
        sub, total = a_decimal(c['subtotal']), a_decimal(c['total'])
        if not sub and total:
            sub = total / (1 + tasa_igv())
        doc = {'id': c['id'], 'fecha': c['fecha_emision'],
               'documento': f"{c['serie'] or ''}-{str(c['numero'] or '').zfill(8)}",
               'tipo': c['tipo_comprobante'] or '', 'cliente': (c['cliente_nombre'] or '').strip() or 'Sin nombre',
               'ruc': c['cliente_numero_doc'] or '', 'empresa_id': c['empresa_id'],
               'empresa': _nombre_empresa(c['empresa_id']), 'vendedor_id': vend,
               'vendedor': usuarios[vend]['nombre'] if vend else 'Sin vendedor asignado',
               'subtotal': sub, 'total': total, 'factor_total': (total / sub) if sub else Decimal('1'),
               'cotizacion': c['documento_asociado'] or '', 'pc': c['pc_vinculado'] or ''}
        if con_lineas:
            lineas = _items(c['items_json'])
            suma = sum((l['importe'] for l in lineas), Decimal('0'))
            if not lineas:
                lineas = [{'producto_id': None, 'codigo': '', 'descripcion': '(comprobante sin detalle de productos)',
                           'cantidad': Decimal('0'), 'importe': sub}]
                suma = sub
            for l in lineas:
                # Prorrateo: las líneas suman exactamente el subtotal del comprobante (incluye descuentos)
                l['venta'] = (l['importe'] * sub / suma) if suma else (sub / len(lineas))
                p = None
                try:
                    p = por_id.get(int(l['producto_id'])) if l['producto_id'] else None
                except (TypeError, ValueError):
                    p = None
                p = p or por_codigo.get(l['codigo'].upper())
                l['producto_id'] = p['id'] if p else None
                l['codigo'] = l['codigo'] or (p['codigo'] if p else '')
                l['descripcion'] = l['descripcion'] or (p['descripcion'] if p else '') or '(sin descripción)'
                l['costo_unit'] = p['costo'] if p else None
                l['costo'] = (p['costo'] * l['cantidad']) if p else Decimal('0')
            if f['producto']:
                lineas = [l for l in lineas if _coincide(f['producto'], l['codigo'], l['descripcion'])]
                if not lineas:
                    continue
                doc['subtotal'] = sum((l['venta'] for l in lineas), Decimal('0'))
                doc['total'] = doc['subtotal'] * doc['factor_total']
            doc['lineas'] = lineas
        docs.append(doc)
    return docs


def _notas_credito(f, desde, hasta):
    if not _existe('notas_credito'):
        return Decimal('0')
    emp_sql, emp_par = _sql_empresa('empresa_id', f) if _tiene('notas_credito', 'empresa_id') else ('', ())
    cli = ''
    par = ()
    if f['cliente']:
        cli = ' AND (cliente_nombre ILIKE %s OR cliente_numero_doc ILIKE %s)'
        par = (f"%{f['cliente']}%",) * 2
    r = db_query(f"""SELECT COALESCE(SUM(monto), 0) AS t FROM notas_credito
                     WHERE COALESCE(UPPER(estado), '') NOT IN ('ANULADA', 'ANULADO', 'BORRADOR', 'POR APROBAR', 'RECHAZADA')
                       AND fecha_emision >= %s AND fecha_emision < %s::date + 1 {emp_sql}{cli}""",
                 (desde, hasta) + emp_par + par)
    return a_decimal(r[0]['t'])


def _n(v, dec=2):
    return round(float(v or 0), dec)


def _col(k, t, tipo='texto'):
    return {'k': k, 't': t, 'tipo': tipo}


def _agrupar_docs(docs, clave, nombre):
    g = {}
    for d in docs:
        k = clave(d)
        x = g.setdefault(k, {'grupo': nombre(d), 'documentos': 0, 'subtotal': Decimal('0'), 'total': Decimal('0'),
                             'clientes': set(), 'ultima': None})
        x['documentos'] += 1
        x['subtotal'] += d['subtotal']
        x['total'] += d['total']
        x['clientes'].add(d['cliente'])
        if x['ultima'] is None or d['fecha'] > x['ultima']:
            x['ultima'] = d['fecha']
    return g


def rep_ventas(f):
    docs = _ventas(f)
    ag = f['agrupar'] or 'documento'
    tot_sub = sum((d['subtotal'] for d in docs), Decimal('0'))
    tot = sum((d['total'] for d in docs), Decimal('0'))
    nc = Decimal('0') if f['producto'] else _notas_credito(f, f['desde'], f['hasta'])
    resumen = [
        {'t': 'Ventas (con IGV)', 'v': _n(tot), 'tipo': 'dinero', 'color': 'ok'},
        {'t': 'Ventas sin IGV', 'v': _n(tot_sub), 'tipo': 'dinero'},
        {'t': 'Comprobantes', 'v': len(docs), 'tipo': 'num'},
        {'t': 'Ticket promedio', 'v': _n(tot / len(docs)) if docs else 0, 'tipo': 'dinero'},
    ]
    if nc:
        resumen.append({'t': 'Notas de crédito', 'v': _n(nc), 'tipo': 'dinero', 'color': 'warn'})
        resumen.append({'t': 'Venta neta', 'v': _n(tot - nc), 'tipo': 'dinero'})
    if ag == 'producto':
        g = {}
        for d in docs:
            for l in d['lineas']:
                k = l['producto_id'] or (l['codigo'] or l['descripcion']).upper()
                x = g.setdefault(k, {'codigo': l['codigo'], 'producto': l['descripcion'], 'cantidad': Decimal('0'),
                                     'subtotal': Decimal('0'), 'total': Decimal('0'), 'docs': set()})
                x['cantidad'] += l['cantidad']
                x['subtotal'] += l['venta']
                x['total'] += l['venta'] * d['factor_total']
                x['docs'].add(d['id'])
        filas = sorted(({'codigo': x['codigo'], 'producto': x['producto'], 'cantidad': _n(x['cantidad']),
                         'documentos': len(x['docs']), 'subtotal': _n(x['subtotal']), 'total': _n(x['total'])}
                        for x in g.values()), key=lambda r: -r['total'])
        cols = [_col('codigo', 'Código'), _col('producto', 'Producto'), _col('cantidad', 'Cantidad', 'num'),
                _col('documentos', 'Comprobantes', 'num'), _col('subtotal', 'Venta sin IGV', 'dinero'),
                _col('total', 'Venta con IGV', 'dinero')]
        return _salida('Ventas por producto', cols, filas, ['cantidad', 'documentos', 'subtotal', 'total'], resumen,
                       grafico=('producto', 'total'))
    if ag == 'documento':
        filas = [{'fecha': d['fecha'], 'empresa': d['empresa'], 'tipo': d['tipo'], 'documento': d['documento'],
                  'cliente': d['cliente'], 'ruc': d['ruc'], 'vendedor': d['vendedor'], 'subtotal': _n(d['subtotal']),
                  'igv': _n(d['total'] - d['subtotal']), 'total': _n(d['total'])} for d in docs]
        cols = [_col('fecha', 'Fecha', 'fechahora'), _col('empresa', 'Empresa'), _col('tipo', 'Tipo'),
                _col('documento', 'Comprobante'), _col('cliente', 'Cliente'), _col('ruc', 'RUC / DNI'),
                _col('vendedor', 'Vendedor'), _col('subtotal', 'Sin IGV', 'dinero'), _col('igv', 'IGV', 'dinero'),
                _col('total', 'Total', 'dinero')]
        return _salida('Ventas por comprobante', cols, filas, ['subtotal', 'igv', 'total'], resumen)
    claves = {
        'dia': (lambda d: d['fecha'].date(), lambda d: d['fecha'].strftime('%d/%m/%Y'), 'Día'),
        'mes': (lambda d: (d['fecha'].year, d['fecha'].month), lambda d: f"{MESES[d['fecha'].month - 1]} {d['fecha'].year}", 'Mes'),
        'empresa': (lambda d: str(d['empresa_id']), lambda d: d['empresa'], 'Empresa'),
        'vendedor': (lambda d: d['vendedor_id'] or 0, lambda d: d['vendedor'], 'Vendedor'),
        'cliente': (lambda d: (d['ruc'] or d['cliente']).upper(), lambda d: f"{d['cliente']} ({d['ruc']})" if d['ruc'] else d['cliente'], 'Cliente'),
    }
    if ag not in claves:
        raise ErrorReporte('Agrupación no válida.')
    clave, nombre, titulo = claves[ag]
    g = _agrupar_docs(docs, clave, nombre)
    orden = sorted(g.items(), key=(lambda kv: kv[0]) if ag in ('dia', 'mes') else (lambda kv: -kv[1]['total']))
    filas = [{'grupo': x['grupo'], 'documentos': x['documentos'], 'clientes': len(x['clientes']),
              'subtotal': _n(x['subtotal']), 'total': _n(x['total']),
              'ticket': _n(x['total'] / x['documentos']) if x['documentos'] else 0,
              'participacion': _n(x['total'] * 100 / tot, 1) if tot else 0} for _, x in orden]
    cols = [_col('grupo', titulo), _col('documentos', 'Comprobantes', 'num'), _col('clientes', 'Clientes', 'num'),
            _col('subtotal', 'Sin IGV', 'dinero'), _col('total', 'Total', 'dinero'), _col('ticket', 'Ticket promedio', 'dinero'),
            _col('participacion', '% del total', 'pct')]
    return _salida(f'Ventas por {titulo.lower()}', cols, filas, ['documentos', 'subtotal', 'total'], resumen,
                   grafico=('grupo', 'total'))


def rep_utilidad(f):
    docs = _ventas(f)
    ag = f['agrupar'] or 'producto'
    g = {}
    for d in docs:
        for l in d['lineas']:
            if ag == 'cliente':
                k, nom = (d['ruc'] or d['cliente']).upper(), d['cliente']
            elif ag == 'vendedor':
                k, nom = d['vendedor_id'] or 0, d['vendedor']
            else:
                k = l['producto_id'] or (l['codigo'] or l['descripcion']).upper()
                nom = f"{l['codigo']} · {l['descripcion']}" if l['codigo'] else l['descripcion']
            x = g.setdefault(k, {'grupo': nom, 'cantidad': Decimal('0'), 'venta': Decimal('0'), 'costo': Decimal('0'),
                                 'sin_costo': False})
            x['cantidad'] += l['cantidad']
            x['venta'] += l['venta']
            x['costo'] += l['costo']
            if l['costo_unit'] is None or not l['costo_unit']:
                x['sin_costo'] = True
    filas = []
    for x in g.values():
        util = x['venta'] - x['costo']
        filas.append({'grupo': x['grupo'], 'cantidad': _n(x['cantidad']), 'venta': _n(x['venta']), 'costo': _n(x['costo']),
                      'utilidad': _n(util), 'margen': _n(util * 100 / x['venta'], 1) if x['venta'] else 0,
                      'aviso': '⚠ Falta costo' if x['sin_costo'] else ''})
    filas.sort(key=lambda r: -r['utilidad'])
    venta = sum((x['venta'] for x in g.values()), Decimal('0'))
    costo = sum((x['costo'] for x in g.values()), Decimal('0'))
    resumen = [{'t': 'Venta sin IGV', 'v': _n(venta), 'tipo': 'dinero'},
               {'t': 'Costo de lo vendido', 'v': _n(costo), 'tipo': 'dinero'},
               {'t': 'Utilidad bruta', 'v': _n(venta - costo), 'tipo': 'dinero', 'color': 'ok' if venta >= costo else 'danger'},
               {'t': 'Margen', 'v': _n((venta - costo) * 100 / venta, 1) if venta else 0, 'tipo': 'pct'}]
    titulo = {'cliente': 'Cliente', 'vendedor': 'Vendedor'}.get(ag, 'Producto')
    cols = [_col('grupo', titulo), _col('cantidad', 'Cantidad', 'num'), _col('venta', 'Venta sin IGV', 'dinero'),
            _col('costo', 'Costo', 'dinero'), _col('utilidad', 'Utilidad', 'dinero'), _col('margen', 'Margen', 'margen'),
            _col('aviso', 'Revisar')]
    t = _salida(f'Utilidad por {titulo.lower()}', cols, filas, ['cantidad', 'venta', 'costo', 'utilidad'], resumen,
                grafico=('grupo', 'utilidad'))
    t['totales']['margen'] = _n((venta - costo) * 100 / venta, 1) if venta else 0
    t['nota'] = ('Venta sin IGV menos costo. Costo = costo promedio del inventario o, si no hay, el costo del producto. '
                 '"Falta costo" = producto sin costo registrado: su utilidad sale inflada.')
    return t


def rep_top_productos(f):
    docs = _ventas(f)
    g = {}
    for d in docs:
        for l in d['lineas']:
            if l['producto_id'] is None and not l['codigo']:
                continue
            k = l['producto_id'] or l['codigo'].upper()
            x = g.setdefault(k, {'codigo': l['codigo'], 'producto': l['descripcion'], 'cantidad': Decimal('0'),
                                 'venta': Decimal('0'), 'costo': Decimal('0'), 'clientes': set()})
            x['cantidad'] += l['cantidad']
            x['venta'] += l['venta']
            x['costo'] += l['costo']
            x['clientes'].add(d['cliente'])
    clave = (lambda x: -x['cantidad']) if f['orden'] == 'cantidad' else (lambda x: -x['venta'])
    lista = sorted(g.values(), key=clave)[:f['limite']]
    total = sum((x['venta'] for x in g.values()), Decimal('0'))
    filas = [{'puesto': i + 1, 'codigo': x['codigo'], 'producto': x['producto'], 'cantidad': _n(x['cantidad']),
              'venta': _n(x['venta']), 'utilidad': _n(x['venta'] - x['costo']), 'clientes': len(x['clientes']),
              'participacion': _n(x['venta'] * 100 / total, 1) if total else 0} for i, x in enumerate(lista)]
    cols = [_col('puesto', '#', 'num'), _col('codigo', 'Código'), _col('producto', 'Producto'),
            _col('cantidad', 'Cantidad', 'num'), _col('venta', 'Venta sin IGV', 'dinero'), _col('utilidad', 'Utilidad', 'dinero'),
            _col('clientes', 'Clientes', 'num'), _col('participacion', '% de la venta', 'pct')]
    return _salida(f"Top {f['limite']} productos más vendidos", cols, filas, ['cantidad', 'venta', 'utilidad'],
                   grafico=('producto', 'cantidad' if f['orden'] == 'cantidad' else 'venta'))


def rep_top_clientes(f):
    docs = _ventas(f, con_lineas=bool(f['producto']))
    g = _agrupar_docs(docs, lambda d: (d['ruc'] or d['cliente']).upper(), lambda d: d['cliente'])
    rucs = {}
    for d in docs:
        rucs[(d['ruc'] or d['cliente']).upper()] = d['ruc']
    total = sum((x['total'] for x in g.values()), Decimal('0'))
    lista = sorted(g.items(), key=lambda kv: -kv[1]['total'])[:f['limite']]
    filas = [{'puesto': i + 1, 'cliente': x['grupo'], 'ruc': rucs.get(k, ''), 'documentos': x['documentos'],
              'subtotal': _n(x['subtotal']), 'total': _n(x['total']), 'ticket': _n(x['total'] / x['documentos']),
              'ultima': x['ultima'], 'participacion': _n(x['total'] * 100 / total, 1) if total else 0}
             for i, (k, x) in enumerate(lista)]
    cols = [_col('puesto', '#', 'num'), _col('cliente', 'Cliente'), _col('ruc', 'RUC / DNI'),
            _col('documentos', 'Compras', 'num'), _col('subtotal', 'Sin IGV', 'dinero'), _col('total', 'Total', 'dinero'),
            _col('ticket', 'Ticket promedio', 'dinero'), _col('ultima', 'Última compra', 'fecha'),
            _col('participacion', '% del total', 'pct')]
    return _salida(f"Top {f['limite']} clientes que más compran", cols, filas, ['documentos', 'subtotal', 'total'],
                   grafico=('cliente', 'total'))


def rep_clientes_perdidos(f):
    """Clientes que compraban y no compran hace N días (o más)."""
    hoy = date.today()
    f2 = dict(f, usuario_id=None, producto='')
    docs = _ventas(f2, desde=date(2000, 1, 1), hasta=hoy, con_lineas=False)
    g = _agrupar_docs(docs, lambda d: (d['ruc'] or d['cliente']).upper(), lambda d: d['cliente'])
    vend = {}
    rucs = {}
    for d in docs:
        k = (d['ruc'] or d['cliente']).upper()
        vend[k] = d['vendedor']
        rucs[k] = d['ruc']
    filas = []
    for k, x in g.items():
        dias = (hoy - x['ultima'].date()).days
        if dias >= f['dias']:
            filas.append({'cliente': x['grupo'], 'ruc': rucs.get(k, ''), 'ultima': x['ultima'], 'dias': dias,
                          'documentos': x['documentos'], 'total': _n(x['total']),
                          'ticket': _n(x['total'] / x['documentos']), 'vendedor': vend.get(k, '')})
    filas.sort(key=lambda r: -r['total'])
    cols = [_col('cliente', 'Cliente'), _col('ruc', 'RUC / DNI'), _col('ultima', 'Última compra', 'fecha'),
            _col('dias', 'Días sin comprar', 'dias'), _col('documentos', 'Compras (histórico)', 'num'),
            _col('total', 'Total comprado', 'dinero'), _col('ticket', 'Ticket promedio', 'dinero'),
            _col('vendedor', 'Último vendedor')]
    resumen = [{'t': 'Clientes por recuperar', 'v': len(filas), 'tipo': 'num', 'color': 'danger' if filas else 'ok'},
               {'t': 'Lo que compraban (histórico)', 'v': _n(sum(r['total'] for r in filas)), 'tipo': 'dinero'}]
    t = _salida(f"Clientes que dejaron de comprar (hace {f['dias']} días o más)", cols, filas, ['documentos', 'total'], resumen)
    t['nota'] = 'Ordenado por lo que compraban: llama primero a los de arriba.'
    return t


# ============================================================
# KPI POR COLABORADOR (solo administración)
# ============================================================
def rep_kpi(f):
    usuarios = _usuarios()
    emp_sql, emp_par = _sql_empresa('c.empresa_id', f) if _tiene('cotizaciones', 'empresa_id') else ('', ())
    cots = db_query(f"""
        SELECT c.id, c.numero_cotizacion, c.usuario_id, c.fecha_creacion, c.estado, COALESCE(c.total, 0) AS total
        FROM cotizaciones c
        WHERE COALESCE(UPPER(TRIM(c.estado)), '') NOT IN ('ANULADA', 'ANULADO', 'ELIMINADA', 'ELIMINADO')
          AND c.fecha_creacion >= %s AND c.fecha_creacion < %s::date + 1 {emp_sql}
    """, (f['desde'], f['hasta']) + emp_par)
    # Primer PV (no anulado) y primer comprobante (no anulado) de cada cotización
    pcs = {}
    for p in db_query("""SELECT cotizacion_id, cotizacion_numero, numero, fecha FROM pedido_compra_pc
                         WHERE COALESCE(UPPER(estado), '') NOT IN ('ANULADO', 'ANULADA')"""):
        for k in (('id', p['cotizacion_id']), ('num', (p['cotizacion_numero'] or '').strip().upper())):
            if k[1] and p['fecha'] and (k not in pcs or p['fecha'] < pcs[k]):
                pcs[k] = p['fecha']
    pc_cot = {p['numero'].strip().upper(): p['cotizacion_id'] for p in db_query(
        'SELECT numero, cotizacion_id FROM pedido_compra_pc WHERE numero IS NOT NULL AND cotizacion_id IS NOT NULL')}
    comps = {}
    for c in db_query("""SELECT documento_asociado, pc_vinculado, fecha_emision FROM comprobantes
                         WHERE COALESCE(UPPER(TRIM(estado_sunat)), 'BORRADOR') <> ALL(%s)""", (list(VENTA_NO_CUENTA),)):
        claves = []
        if c['documento_asociado']:
            claves.append(('num', c['documento_asociado'].strip().upper()))
        cid = pc_cot.get((c['pc_vinculado'] or '').strip().upper())
        if cid:
            claves.append(('id', cid))
        for k in claves:
            if c['fecha_emision'] and (k not in comps or c['fecha_emision'] < comps[k]):
                comps[k] = c['fecha_emision']
    stats = defaultdict(lambda: {'cot': 0, 'cot_monto': Decimal('0'), 'conv': 0, 'dias': [], 'ventas': 0,
                                 'venta': Decimal('0'), 'venta_sub': Decimal('0'), 'clientes': set()})
    for c in cots:
        uid = c['usuario_id']
        if uid not in usuarios or (f['usuario_id'] and uid != f['usuario_id']):
            continue
        s = stats[uid]
        s['cot'] += 1
        s['cot_monto'] += a_decimal(c['total'])
        claves = [('id', c['id']), ('num', (c['numero_cotizacion'] or '').strip().upper())]
        fechas = [x for x in [pcs.get(k) for k in claves] + [comps.get(k) for k in claves] if x]
        if fechas or (c['estado'] or '').strip().upper() in COT_CONVERTIDA:
            s['conv'] += 1
            if fechas and c['fecha_creacion']:
                s['dias'].append(max((min(fechas) - c['fecha_creacion']).total_seconds() / 86400, 0))
    for d in _ventas(dict(f, cliente='', producto=''), con_lineas=False):
        if d['vendedor_id'] is None:
            continue
        s = stats[d['vendedor_id']]
        s['ventas'] += 1
        s['venta'] += d['total']
        s['venta_sub'] += d['subtotal']
        s['clientes'].add(d['cliente'])
    total = sum((s['venta'] for s in stats.values()), Decimal('0'))
    filas = []
    for uid, s in stats.items():
        filas.append({'colaborador': usuarios[uid]['nombre'], 'rol': usuarios[uid]['rol'],
                      'cotizaciones': s['cot'], 'cotizado': _n(s['cot_monto']), 'convertidas': s['conv'],
                      'conversion': _n(s['conv'] * 100 / s['cot'], 1) if s['cot'] else 0,
                      'dias_cierre': _n(sum(s['dias']) / len(s['dias']), 1) if s['dias'] else None,
                      'ventas': s['ventas'], 'venta_sub': _n(s['venta_sub']), 'venta': _n(s['venta']),
                      'ticket': _n(s['venta'] / s['ventas']) if s['ventas'] else 0, 'clientes': len(s['clientes']),
                      'participacion': _n(s['venta'] * 100 / total, 1) if total else 0})
    filas.sort(key=lambda r: (-r['venta'], -r['cotizaciones']))
    for i, r in enumerate(filas):
        r['puesto'] = i + 1
    cols = [_col('puesto', '#', 'num'), _col('colaborador', 'Colaborador'), _col('rol', 'Rol'),
            _col('cotizaciones', 'Cotizaciones hechas', 'num'), _col('cotizado', 'Monto cotizado', 'dinero'),
            _col('convertidas', 'Convertidas en venta', 'num'), _col('conversion', 'Tasa de conversión', 'semaforo'),
            _col('dias_cierre', 'Días promedio cotización → PV/factura', 'dias1'), _col('ventas', 'Facturas', 'num'),
            _col('venta_sub', 'Vendido sin IGV', 'dinero'), _col('venta', 'Vendido con IGV', 'dinero'),
            _col('ticket', 'Ticket promedio', 'dinero'), _col('clientes', 'Clientes atendidos', 'num'),
            _col('participacion', '% de la venta', 'pct')]
    tc = sum(r['cotizaciones'] for r in filas)
    tv = sum(r['convertidas'] for r in filas)
    resumen = [{'t': 'Vendido (con IGV)', 'v': _n(total), 'tipo': 'dinero', 'color': 'ok'},
               {'t': 'Cotizaciones hechas', 'v': tc, 'tipo': 'num'},
               {'t': 'Convertidas en venta', 'v': tv, 'tipo': 'num'},
               {'t': 'Conversión del equipo', 'v': _n(tv * 100 / tc, 1) if tc else 0, 'tipo': 'pct'}]
    t = _salida('KPI por colaborador', cols, filas, ['cotizaciones', 'cotizado', 'convertidas', 'ventas', 'venta_sub', 'venta'],
                resumen, grafico=('colaborador', 'venta'))
    t['nota'] = ('Cotizaciones: las creadas en el periodo (sin anuladas). Convertida = tiene PV o factura, o está aceptada. '
                 'Vendido: facturas emitidas en el periodo cuya cotización es del colaborador '
                 '(si la factura no tiene cotización, cuenta para quien la registró). Solo usuarios del sistema.')
    return t


# ============================================================
# COMPRAS
# ============================================================
def _compras(f):
    emp_sql, emp_par = _sql_empresa('oc.empresa_id', f) if _tiene('ordenes_compra', 'empresa_id') else ('', ())
    emp_col = 'oc.empresa_id' if _tiene('ordenes_compra', 'empresa_id') else 'NULL'
    ocs = db_query(f"""
        SELECT oc.id, oc.numero_orden, COALESCE(oc.fecha_creacion, oc.created_at) AS fecha, oc.estado,
               {emp_col} AS empresa_id, COALESCE(p.razon_social, oc.contacto_proveedor, 'Sin proveedor') AS proveedor,
               COALESCE(p.ruc, '') AS ruc, COALESCE(oc.subtotal, 0) AS subtotal, COALESCE(oc.total, 0) AS total
        FROM ordenes_compra oc LEFT JOIN proveedores p ON p.id = oc.proveedor_id
        WHERE COALESCE(UPPER(TRIM(oc.estado)), '') <> ALL(%s)
          AND COALESCE(oc.fecha_creacion, oc.created_at) >= %s AND COALESCE(oc.fecha_creacion, oc.created_at) < %s::date + 1
          {emp_sql}
        ORDER BY 3, 1
    """, (list(OC_NO_CUENTA), f['desde'], f['hasta']) + emp_par)
    if f['proveedor']:
        ocs = [o for o in ocs if _coincide(f['proveedor'], o['proveedor'], o['ruc'])]
    if not ocs:
        return []
    ids = [o['id'] for o in ocs]
    por_id, _ = _productos()
    lineas = defaultdict(list)
    pid = 'producto_id' if _tiene('ordenes_compra_items', 'producto_id') else 'NULL::int AS producto_id'
    for it in db_query(f"""SELECT orden_compra_id, producto, {pid}, COALESCE(cantidad, 0) AS cantidad,
                                  COALESCE(precio_unitario, 0) AS precio, COALESCE(total, 0) AS total
                           FROM ordenes_compra_items WHERE orden_compra_id = ANY(%s) ORDER BY id""", (ids,)):
        p = por_id.get(it['producto_id'])
        imp = a_decimal(it['cantidad']) * a_decimal(it['precio']) or a_decimal(it['total'])
        lineas[it['orden_compra_id']].append({'codigo': p['codigo'] if p else '', 'descripcion': (p['descripcion'] if p else None) or it['producto'] or '',
                                              'producto_id': it['producto_id'], 'cantidad': a_decimal(it['cantidad']), 'importe': imp})
    faltan = [i for i in ids if i not in lineas]
    if faltan and _existe('orden_compra_detalle'):
        for it in db_query("""SELECT orden_id, producto_id, COALESCE(cantidad, 0) AS cantidad, COALESCE(costo_unitario, 0) AS precio
                              FROM orden_compra_detalle WHERE orden_id = ANY(%s) ORDER BY id""", (faltan,)):
            p = por_id.get(it['producto_id'])
            lineas[it['orden_id']].append({'codigo': p['codigo'] if p else '', 'descripcion': p['descripcion'] if p else '(producto)',
                                           'producto_id': it['producto_id'], 'cantidad': a_decimal(it['cantidad']),
                                           'importe': a_decimal(it['cantidad']) * a_decimal(it['precio'])})
    docs = []
    for o in ocs:
        sub, total = a_decimal(o['subtotal']), a_decimal(o['total'])
        ls = lineas.get(o['id']) or [{'codigo': '', 'descripcion': '(OC sin detalle)', 'producto_id': None,
                                      'cantidad': Decimal('0'), 'importe': sub}]
        suma = sum((l['importe'] for l in ls), Decimal('0'))
        if not sub:
            sub = suma
            total = total or sub * (1 + tasa_igv())
        for l in ls:
            l['compra'] = (l['importe'] * sub / suma) if suma else Decimal('0')
        factor = (total / sub) if sub else Decimal('1')
        if f['producto']:
            ls = [l for l in ls if _coincide(f['producto'], l['codigo'], l['descripcion'])]
            if not ls:
                continue
            sub = sum((l['compra'] for l in ls), Decimal('0'))
            total = sub * factor
        docs.append({'id': o['id'], 'fecha': o['fecha'], 'documento': o['numero_orden'] or f"OC #{o['id']}",
                     'estado': o['estado'] or '', 'empresa_id': o['empresa_id'], 'empresa': _nombre_empresa(o['empresa_id']),
                     'proveedor': o['proveedor'], 'ruc': o['ruc'], 'subtotal': sub, 'total': total, 'factor': factor,
                     'lineas': ls})
    return docs


def rep_compras(f):
    docs = _compras(f)
    ag = f['agrupar'] or 'documento'
    tot_sub = sum((d['subtotal'] for d in docs), Decimal('0'))
    tot = sum((d['total'] for d in docs), Decimal('0'))
    resumen = [{'t': 'Compras (con IGV)', 'v': _n(tot), 'tipo': 'dinero'},
               {'t': 'Compras sin IGV', 'v': _n(tot_sub), 'tipo': 'dinero'},
               {'t': 'Órdenes de compra (Proveedor)', 'v': len(docs), 'tipo': 'num'},
               {'t': 'Proveedores', 'v': len({d['proveedor'] for d in docs}), 'tipo': 'num'}]
    if ag == 'documento':
        filas = [{'fecha': d['fecha'], 'empresa': d['empresa'], 'documento': d['documento'], 'estado': d['estado'],
                  'proveedor': d['proveedor'], 'ruc': d['ruc'], 'subtotal': _n(d['subtotal']),
                  'igv': _n(d['total'] - d['subtotal']), 'total': _n(d['total'])} for d in docs]
        cols = [_col('fecha', 'Fecha', 'fechahora'), _col('empresa', 'Empresa'), _col('documento', 'Orden de compra'),
                _col('estado', 'Estado'), _col('proveedor', 'Proveedor'), _col('ruc', 'RUC'),
                _col('subtotal', 'Sin IGV', 'dinero'), _col('igv', 'IGV', 'dinero'), _col('total', 'Total', 'dinero')]
        return _salida('Compras por orden de compra', cols, filas, ['subtotal', 'igv', 'total'], resumen)
    if ag == 'producto':
        g = {}
        for d in docs:
            for l in d['lineas']:
                k = l['producto_id'] or (l['codigo'] or l['descripcion']).upper()
                x = g.setdefault(k, {'codigo': l['codigo'], 'producto': l['descripcion'], 'cantidad': Decimal('0'),
                                     'subtotal': Decimal('0'), 'docs': set(), 'prov': set()})
                x['cantidad'] += l['cantidad']
                x['subtotal'] += l['compra']
                x['docs'].add(d['id'])
                x['prov'].add(d['proveedor'])
        filas = sorted(({'codigo': x['codigo'], 'producto': x['producto'], 'cantidad': _n(x['cantidad']),
                         'precio': _n(x['subtotal'] / x['cantidad'], 4) if x['cantidad'] else 0,
                         'ordenes': len(x['docs']), 'proveedores': len(x['prov']), 'subtotal': _n(x['subtotal'])}
                        for x in g.values()), key=lambda r: -r['subtotal'])
        cols = [_col('codigo', 'Código'), _col('producto', 'Producto'), _col('cantidad', 'Cantidad', 'num'),
                _col('precio', 'Costo promedio', 'dinero'), _col('ordenes', 'OC', 'num'), _col('proveedores', 'Proveedores', 'num'),
                _col('subtotal', 'Comprado sin IGV', 'dinero')]
        return _salida('Compras por producto', cols, filas, ['cantidad', 'ordenes', 'subtotal'], resumen,
                       grafico=('producto', 'subtotal'))
    claves = {
        'mes': (lambda d: (d['fecha'].year, d['fecha'].month), lambda d: f"{MESES[d['fecha'].month - 1]} {d['fecha'].year}", 'Mes'),
        'empresa': (lambda d: str(d['empresa_id']), lambda d: d['empresa'], 'Empresa'),
        'proveedor': (lambda d: (d['ruc'] or d['proveedor']).upper(), lambda d: d['proveedor'], 'Proveedor'),
    }
    if ag not in claves:
        raise ErrorReporte('Agrupación no válida.')
    clave, nombre, titulo = claves[ag]
    g = {}
    for d in docs:
        x = g.setdefault(clave(d), {'grupo': nombre(d), 'ordenes': 0, 'subtotal': Decimal('0'), 'total': Decimal('0')})
        x['ordenes'] += 1
        x['subtotal'] += d['subtotal']
        x['total'] += d['total']
    orden = sorted(g.items(), key=(lambda kv: kv[0]) if ag == 'mes' else (lambda kv: -kv[1]['total']))
    filas = [{'grupo': x['grupo'], 'ordenes': x['ordenes'], 'subtotal': _n(x['subtotal']), 'total': _n(x['total']),
              'participacion': _n(x['total'] * 100 / tot, 1) if tot else 0} for _, x in orden]
    cols = [_col('grupo', titulo), _col('ordenes', 'Órdenes de compra (Proveedor)', 'num'), _col('subtotal', 'Sin IGV', 'dinero'),
            _col('total', 'Total', 'dinero'), _col('participacion', '% del total', 'pct')]
    return _salida(f'Compras por {titulo.lower()}', cols, filas, ['ordenes', 'subtotal', 'total'], resumen,
                   grafico=('grupo', 'total'))


# ============================================================
# CUENTAS POR COBRAR / POR PAGAR (mismas consultas que Finanzas)
# ============================================================
def _cuentas(tipo, f):
    from routes.finanzas import SQL_CXC, SQL_CXP, asegurar_tablas_finanzas, _semaforo
    asegurar_tablas_finanzas()
    if tipo == 'cxc':
        sql = SQL_CXC + " AND COALESCE(UPPER(c.estado_sunat), 'BORRADOR') <> 'BORRADOR'"
        emp_sql, emp_par = _sql_empresa('c.empresa_id', f)
        sql = sql.replace('SELECT c.id,', 'SELECT c.empresa_id AS rep_empresa_id, c.id,', 1)
    else:
        sql = SQL_CXP
        emp_sql, emp_par = _sql_empresa('cp.empresa_id', f) if _tiene('comprobantes_proveedor', 'empresa_id') else ('', ())
        if _tiene('comprobantes_proveedor', 'empresa_id'):
            sql = sql.replace('SELECT cp.id,', 'SELECT cp.empresa_id AS rep_empresa_id, cp.id,', 1)
    filas = db_query(sql + emp_sql, emp_par)
    hoy = date.today()
    texto = f['cliente'] if tipo == 'cxc' else f['proveedor']
    salida = []
    for r in filas:
        if texto and not _coincide(texto, r['tercero'], r['ruc'], r['documento']):
            continue
        importe = a_decimal(r['importe_cobrar'])
        abonado = a_decimal(r['cobrado']) + a_decimal(r['notas_credito'])
        saldo = max(importe - abonado, Decimal('0'))
        if saldo <= Decimal('0.004'):
            continue
        venc = r['fecha_vencimiento']
        if isinstance(venc, datetime):
            venc = venc.date()
        sem = _semaforo(saldo, venc, hoy)
        dias = (hoy - venc).days if venc else 0
        # 24/10/2026: documentos en dólares -> se muestra la moneda y los totales van en soles con su TC
        mon = 'USD' if str(r.get('moneda') or '').upper() == 'USD' else 'PEN'
        tc = a_decimal(r.get('tipo_cambio') or 0) if mon == 'USD' else Decimal('1')
        if mon == 'USD' and tc <= 0:
            from routes.moneda import tc_referencia
            tc = a_decimal(tc_referencia(hoy) or 0)
        salida.append({'moneda': 'US$' if mon == 'USD' else 'S/', 'saldo_soles': _n(saldo * tc),
                       'estado': sem, 'dias_atraso': dias if sem == 'Vencido' else 0,
                       'dias_vence': -dias if venc and sem != 'Vencido' else None,
                       'empresa': _nombre_empresa(r.get('rep_empresa_id')),
                       'documento': r['documento'], 'tercero': r['tercero'] or '', 'ruc': r['ruc'] or '',
                       'emision': r['fecha_emision'], 'vencimiento': venc, 'importe': _n(importe),
                       'abonado': _n(abonado), 'saldo': _n(saldo)})
    est = f['estado'].lower()
    if est in ('vencido', 'por vencer', 'al día'):
        salida = [x for x in salida if x['estado'].lower() == est]
    prio = {'Vencido': 0, 'Por vencer': 1, 'Al día': 2}
    salida.sort(key=lambda x: (prio.get(x['estado'], 3), -x['dias_atraso'],
                               x['vencimiento'] or date.max, x['tercero']))
    return salida


def _rep_cuentas(tipo, f):
    filas = _cuentas(tipo, f)
    quien = 'Cliente' if tipo == 'cxc' else 'Proveedor'

    def suma(e):
        return _n(sum(x['saldo_soles'] for x in filas if x['estado'] == e))
    resumen = [{'t': 'Vencido', 'v': suma('Vencido'), 'tipo': 'dinero', 'color': 'danger' if suma('Vencido') else '',
                'sub': f"{sum(1 for x in filas if x['estado'] == 'Vencido')} documentos"},
               {'t': 'Por vencer (7 días)', 'v': suma('Por vencer'), 'tipo': 'dinero', 'color': 'warn' if suma('Por vencer') else '',
                'sub': f"{sum(1 for x in filas if x['estado'] == 'Por vencer')} documentos"},
               {'t': 'Al día', 'v': suma('Al día'), 'tipo': 'dinero', 'color': 'ok',
                'sub': f"{sum(1 for x in filas if x['estado'] == 'Al día')} documentos"},
               # Colores de Finanzas (Erika 25/10/2026): total por cobrar verde azulado, por pagar morado
               {'t': 'Total pendiente', 'v': _n(sum(x['saldo_soles'] for x in filas)), 'tipo': 'dinero',
                'color': 'cobro' if tipo == 'cxc' else 'pago'}]
    cols = [_col('estado', 'Estado', 'estado'), _col('dias_atraso', 'Días de atraso', 'atraso'),
            _col('dias_vence', 'Vence en (días)', 'num'), _col('empresa', 'Empresa'), _col('documento', 'Documento'),
            _col('tercero', quien), _col('ruc', 'RUC'), _col('emision', 'Emisión', 'fecha'),
            _col('vencimiento', 'Vencimiento', 'fecha'), _col('moneda', 'Moneda'), _col('importe', 'Importe', 'num'),
            _col('abonado', 'Cobrado' if tipo == 'cxc' else 'Pagado', 'num'), _col('saldo', 'Saldo', 'num'),
            _col('saldo_soles', 'Saldo en S/', 'dinero')]
    t = _salida('Cuentas por cobrar' if tipo == 'cxc' else 'Cuentas por pagar', cols, filas,
                ['importe', 'abonado', 'saldo', 'saldo_soles'], resumen)
    t['nota'] = ('Todo lo pendiente a hoy (no depende de las fechas). Primero lo vencido, con más días de atraso arriba. '
                 '"Por vencer" = vence en los próximos 7 días. Mismo cálculo que Finanzas. '
                 'Los documentos en dólares se suman en soles con su tipo de cambio.')
    return t


def rep_cxc(f):
    return _rep_cuentas('cxc', f)


def rep_cxp(f):
    return _rep_cuentas('cxp', f)


# ============================================================
# INVENTARIO
# ============================================================
def _sql_inv(f, alias='s'):
    sql, par = '', ()
    if f['empresa_id'] is not None:
        sql += f' AND {alias}.empresa_id::text = %s'
        par += (f['empresa_id'],)
    if f['almacen_id']:
        sql += f' AND {alias}.almacen_id = %s'
        par += (f['almacen_id'],)
    if f['producto']:
        sql += ' AND (p.codigo ILIKE %s OR p.descripcion ILIKE %s)'
        par += (f"%{f['producto']}%",) * 2
    return sql, par


def rep_stock(f):
    from routes.inventario import asegurar_tablas_inventario
    asegurar_tablas_inventario()
    sql, par = _sql_inv(f)
    filas = db_query(f"""
        SELECT s.empresa_id, a.nombre AS sede, p.codigo, p.descripcion, p.unidad, s.cantidad, s.costo_promedio, s.valor,
               s.ultimo_movimiento
        FROM inv_stock s JOIN productos p ON p.id = s.producto_id LEFT JOIN almacenes a ON a.id = s.almacen_id
        WHERE s.cantidad > 0 {sql}
        ORDER BY p.descripcion, s.empresa_id, a.nombre
    """, par)
    ag = f['agrupar'] or 'detalle'
    datos = [{'empresa': _nombre_empresa(r['empresa_id']), 'sede': r['sede'] or '', 'codigo': r['codigo'],
              'producto': r['descripcion'], 'unidad': r['unidad'] or '', 'cantidad': _n(r['cantidad']),
              'costo': _n(r['costo_promedio'], 4), 'valor': _n(r['valor']), 'ultimo': r['ultimo_movimiento']} for r in filas]
    total = sum(d['valor'] for d in datos)
    resumen = [{'t': 'Valor del stock', 'v': _n(total), 'tipo': 'dinero'},
               {'t': 'Productos con stock', 'v': len({d['codigo'] for d in datos}), 'tipo': 'num'},
               {'t': 'Unidades', 'v': _n(sum(d['cantidad'] for d in datos)), 'tipo': 'num'}]
    if ag == 'detalle':
        cols = [_col('empresa', 'Empresa'), _col('sede', 'Sede'), _col('codigo', 'Código'), _col('producto', 'Producto'),
                _col('unidad', 'Unidad'), _col('cantidad', 'Cantidad', 'num'), _col('costo', 'Costo promedio', 'dinero'),
                _col('valor', 'Valor S/', 'dinero'), _col('ultimo', 'Último movimiento', 'fechahora')]
        return _salida('Stock por empresa y sede', cols, datos, ['cantidad', 'valor'], resumen)
    nombres = {'empresa': ('Empresa', lambda d: d['empresa']), 'sede': ('Sede', lambda d: f"{d['empresa']} · {d['sede']}"),
               'producto': ('Producto', lambda d: f"{d['codigo']} · {d['producto']}")}
    if ag not in nombres:
        raise ErrorReporte('Agrupación no válida.')
    titulo, fn = nombres[ag]
    g = {}
    for d in datos:
        x = g.setdefault(fn(d), {'grupo': fn(d), 'cantidad': 0.0, 'valor': 0.0, 'productos': set()})
        x['cantidad'] += d['cantidad']
        x['valor'] += d['valor']
        x['productos'].add(d['codigo'])
    filas2 = sorted(({'grupo': x['grupo'], 'productos': len(x['productos']), 'cantidad': _n(x['cantidad']),
                      'valor': _n(x['valor']), 'participacion': _n(x['valor'] * 100 / total, 1) if total else 0}
                     for x in g.values()), key=lambda r: -r['valor'])
    cols = [_col('grupo', titulo), _col('productos', 'Productos', 'num'), _col('cantidad', 'Unidades', 'num'),
            _col('valor', 'Valor S/', 'dinero'), _col('participacion', '% del valor', 'pct')]
    return _salida(f'Stock por {titulo.lower()}', cols, filas2, ['cantidad', 'valor'], resumen, grafico=('grupo', 'valor'))


def rep_kardex(f):
    from routes.inventario import asegurar_tablas_inventario, MOTIVOS
    asegurar_tablas_inventario()
    sql, par = _sql_inv(f, 'm')
    filas = db_query(f"""
        SELECT m.fecha, m.numero, m.empresa_id, a.nombre AS sede, p.codigo, p.descripcion, m.tipo, m.motivo,
               m.documento_numero, m.tercero, m.cantidad, m.costo_unitario, m.costo_total, m.saldo_cantidad, m.saldo_valor,
               m.usuario_nombre
        FROM inv_movimientos m JOIN productos p ON p.id = m.producto_id LEFT JOIN almacenes a ON a.id = m.almacen_id
        WHERE m.estado = 'Registrado' AND m.fecha >= %s AND m.fecha < %s::date + 1 {sql}
        ORDER BY m.fecha, m.id LIMIT 5000
    """, (f['desde'], f['hasta']) + par)
    datos = []
    for r in filas:
        ent = r['tipo'] == 'ENTRADA'
        datos.append({'fecha': r['fecha'], 'numero': r['numero'] or '', 'empresa': _nombre_empresa(r['empresa_id']),
                      'sede': r['sede'] or '', 'producto': f"{r['codigo']} · {r['descripcion']}",
                      'motivo': MOTIVOS.get(r['motivo'], ('', r['motivo']))[1],
                      'documento': ' '.join(x for x in (r['documento_numero'], r['tercero']) if x),
                      'entrada': _n(r['cantidad']) if ent else None, 'salida': None if ent else _n(r['cantidad']),
                      'costo_unit': _n(r['costo_unitario'], 4), 'valor_entrada': _n(r['costo_total']) if ent else None,
                      'valor_salida': None if ent else _n(r['costo_total']), 'saldo': _n(r['saldo_cantidad']),
                      'saldo_valor': _n(r['saldo_valor']), 'usuario': r['usuario_nombre'] or ''})
    cols = [_col('fecha', 'Fecha', 'fechahora'), _col('numero', 'N°'), _col('empresa', 'Empresa'), _col('sede', 'Sede'),
            _col('producto', 'Producto'), _col('motivo', 'Motivo'), _col('documento', 'Documento / tercero'),
            _col('entrada', 'Entrada', 'num'), _col('salida', 'Salida', 'num'), _col('costo_unit', 'Costo unit.', 'dinero'),
            _col('valor_entrada', 'Valor entrada', 'dinero'), _col('valor_salida', 'Valor salida', 'dinero'),
            _col('saldo', 'Saldo cant.', 'num'), _col('saldo_valor', 'Saldo valor', 'dinero'), _col('usuario', 'Usuario')]
    resumen = [{'t': 'Entradas (valor)', 'v': _n(sum(d['valor_entrada'] or 0 for d in datos)), 'tipo': 'dinero', 'color': 'ok'},
               {'t': 'Salidas (valor)', 'v': _n(sum(d['valor_salida'] or 0 for d in datos)), 'tipo': 'dinero', 'color': 'warn'},
               {'t': 'Movimientos', 'v': len(datos), 'tipo': 'num'}]
    t = _salida('Kardex valorizado', cols, datos, ['entrada', 'salida', 'valor_entrada', 'valor_salida'], resumen)
    t['nota'] = 'Solo movimientos registrados (aprobados). El saldo es el de cada empresa y sede después del movimiento.'
    return t


def rep_rotacion(f):
    from routes.inventario import asegurar_tablas_inventario, _clasificar_rotacion
    asegurar_tablas_inventario()
    sql, par = _sql_inv(f)
    stock = db_query(f"""
        SELECT s.producto_id, p.codigo, p.descripcion, SUM(s.cantidad) AS cantidad, SUM(s.valor) AS valor
        FROM inv_stock s JOIN productos p ON p.id = s.producto_id WHERE TRUE {sql}
        GROUP BY s.producto_id, p.codigo, p.descripcion
    """, par)
    sqlm, parm = _sql_inv(f, 'm')
    mov = {r['producto_id']: r for r in db_query(f"""
        SELECT m.producto_id, SUM(m.cantidad) AS cantidad, SUM(m.costo_total) AS costo, MAX(m.fecha) AS ultima
        FROM inv_movimientos m JOIN productos p ON p.id = m.producto_id
        WHERE m.estado = 'Registrado' AND m.motivo IN ('VENTA', 'INTERCOMPANY_SALIDA')
          AND m.fecha >= %s AND m.fecha < %s::date + 1 {sqlm}
        GROUP BY m.producto_id
    """, (f['desde'], f['hasta']) + parm)}
    ult = {r['producto_id']: r['ultima'] for r in db_query(f"""
        SELECT m.producto_id, MAX(m.fecha) AS ultima FROM inv_movimientos m JOIN productos p ON p.id = m.producto_id
        WHERE m.estado = 'Registrado' AND m.motivo IN ('VENTA', 'INTERCOMPANY_SALIDA') {sqlm} GROUP BY m.producto_id
    """, parm)}
    dias_periodo = (f['hasta'] - f['desde']).days + 1
    hoy = datetime.now()
    filas = []
    for s in stock:
        m = mov.get(s['producto_id']) or {}
        if not a_decimal(s['cantidad']) and not m:
            continue
        costo_v = a_decimal(m.get('costo'))
        valor = a_decimal(s['valor'])
        rot = (costo_v / valor) if valor else None
        dias_inv = (valor / (costo_v / dias_periodo)) if costo_v else None
        ultima = ult.get(s['producto_id'])
        filas.append({'producto': f"{s['codigo']} · {s['descripcion']}", 'stock': _n(s['cantidad']), 'valor': _n(valor),
                      'vendido': _n(m.get('cantidad')), 'costo_vendido': _n(costo_v),
                      'rotacion': _n(rot, 2) if rot is not None else None,
                      'dias_inventario': _n(dias_inv, 0) if dias_inv is not None else None,
                      'ultima_venta': ultima,
                      'clase': _clasificar_rotacion((hoy - ultima).days if ultima else None) if a_decimal(s['cantidad']) else 'Sin stock'})
    filas.sort(key=lambda r: (r['rotacion'] is None, -(r['rotacion'] or 0)))
    cols = [_col('producto', 'Producto'), _col('stock', 'Stock hoy', 'num'), _col('valor', 'Valor stock', 'dinero'),
            _col('vendido', 'Vendido en el periodo', 'num'), _col('costo_vendido', 'Costo de lo vendido', 'dinero'),
            _col('rotacion', 'Rotación (veces)', 'num'), _col('dias_inventario', 'Días de inventario', 'num'),
            _col('ultima_venta', 'Última venta', 'fecha'), _col('clase', 'Movimiento', 'rot')]
    t = _salida('Rotación de inventario', cols, filas, ['stock', 'valor', 'vendido', 'costo_vendido'])
    t['nota'] = ('Rotación = costo de lo vendido en el periodo ÷ valor del stock de hoy (más alto = se vende más rápido). '
                 'Días de inventario = cuántos días alcanzaría el stock al ritmo de venta del periodo.')
    return t


def rep_inmovilizado(f):
    from routes.inventario import asegurar_tablas_inventario
    asegurar_tablas_inventario()
    sql, par = _sql_inv(f)
    filas = db_query(f"""
        WITH ult AS (
            SELECT empresa_id, almacen_id, producto_id, MAX(fecha) AS ultima,
                   MAX(fecha) FILTER (WHERE tipo = 'SALIDA') AS ultima_salida
            FROM inv_movimientos WHERE estado = 'Registrado' GROUP BY 1, 2, 3
        )
        SELECT s.empresa_id, a.nombre AS sede, p.codigo, p.descripcion, s.cantidad, s.costo_promedio, s.valor,
               COALESCE(u.ultima, s.ultimo_movimiento, s.ultimo_ingreso) AS ultima,
               COALESCE(u.ultima_salida, s.ultima_salida) AS ultima_salida
        FROM inv_stock s JOIN productos p ON p.id = s.producto_id LEFT JOIN almacenes a ON a.id = s.almacen_id
        LEFT JOIN ult u ON u.empresa_id = s.empresa_id AND u.almacen_id = s.almacen_id AND u.producto_id = s.producto_id
        WHERE s.cantidad > 0 {sql}
    """, par)
    hoy = datetime.now()
    datos = []
    for r in filas:
        dias = (hoy - r['ultima']).days if r['ultima'] else None
        if dias is not None and dias < f['dias']:
            continue
        datos.append({'empresa': _nombre_empresa(r['empresa_id']), 'sede': r['sede'] or '',
                      'producto': f"{r['codigo']} · {r['descripcion']}", 'cantidad': _n(r['cantidad']),
                      'valor': _n(r['valor']), 'ultima': r['ultima'], 'ultima_salida': r['ultima_salida'],
                      'dias': dias if dias is not None else 'Sin movimientos'})
    datos.sort(key=lambda r: -r['valor'])
    cols = [_col('empresa', 'Empresa'), _col('sede', 'Sede'), _col('producto', 'Producto'), _col('cantidad', 'Cantidad', 'num'),
            _col('valor', 'Valor parado', 'dinero'), _col('ultima', 'Último movimiento', 'fecha'),
            _col('ultima_salida', 'Última salida', 'fecha'), _col('dias', 'Días sin moverse', 'atraso')]
    resumen = [{'t': 'Dinero parado', 'v': _n(sum(d['valor'] for d in datos)), 'tipo': 'dinero', 'color': 'danger' if datos else 'ok'},
               {'t': 'Productos', 'v': len(datos), 'tipo': 'num'}]
    t = _salida(f"Stock inmovilizado (sin movimiento hace {f['dias']} días o más)", cols, datos, ['cantidad', 'valor'], resumen)
    t['nota'] = 'Mayor valor arriba: es la plata que está quieta en el almacén. Piensa en promociones o en devolver al proveedor.'
    return t


# ============================================================
# TABLERO (pantalla de inicio)
# ============================================================
def _inicio_mes(d, atras=0):
    y, m = d.year, d.month - atras
    while m <= 0:
        m += 12
        y -= 1
    return date(y, m, 1)


def tablero(f):
    hoy = date.today()
    ini = _inicio_mes(hoy)
    ini_ant = _inicio_mes(hoy, 1)
    fin_ant = ini - timedelta(days=1)
    ini_12 = _inicio_mes(hoy, 11)
    base = dict(f, usuario_id=None, cliente='', producto='')
    docs = _ventas(base, desde=ini_12, hasta=hoy)
    mes = [d for d in docs if d['fecha'].date() >= ini]
    ant = [d for d in docs if ini_ant <= d['fecha'].date() <= fin_ant]
    # Mismo avance: del 1 al día de hoy en el mes anterior
    dia_lim = min(hoy.day, fin_ant.day)
    ant_avance = [d for d in ant if d['fecha'].day <= dia_lim]

    def tot(lista, k='total'):
        return sum((d[k] for d in lista), Decimal('0'))
    venta_mes, venta_ant, venta_avance = tot(mes), tot(ant), tot(ant_avance)
    costo_mes = sum((l['costo'] for d in mes for l in d['lineas']), Decimal('0'))
    sub_mes = tot(mes, 'subtotal')
    por_mes = []
    for i in range(11, -1, -1):
        m0 = _inicio_mes(hoy, i)
        lista = [d for d in docs if d['fecha'].year == m0.year and d['fecha'].month == m0.month]
        por_mes.append({'mes': f'{MESES[m0.month - 1]} {str(m0.year)[2:]}', 'total': _n(tot(lista)),
                        'subtotal': _n(tot(lista, 'subtotal')), 'documentos': len(lista)})
    g = _agrupar_docs(docs, lambda d: (d['ruc'] or d['cliente']).upper(), lambda d: d['cliente'])
    top = sorted(g.values(), key=lambda x: -x['total'])[:5]
    cxc = _cuentas('cxc', dict(base, estado=''))
    cxp = _cuentas('cxp', dict(base, estado=''))
    perdidos = rep_clientes_perdidos(dict(base, dias=60))
    # Cotizaciones del mes
    emp_sql, emp_par = _sql_empresa('empresa_id', f) if _tiene('cotizaciones', 'empresa_id') else ('', ())
    cot = db_query(f"""SELECT COUNT(*) AS n, COALESCE(SUM(total), 0) AS monto,
                              COUNT(*) FILTER (WHERE UPPER(TRIM(estado)) = ANY(%s)) AS conv
                       FROM cotizaciones WHERE COALESCE(UPPER(TRIM(estado)), '') NOT IN ('ANULADA', 'ANULADO')
                         AND fecha_creacion >= %s {emp_sql}""", (list(COT_CONVERTIDA), ini) + emp_par)[0]
    vencidas = [x for x in cxc if x['estado'] == 'Vencido']
    return {
        'hoy': hoy.isoformat(), 'mes_nombre': f'{MESES[hoy.month - 1]} {hoy.year}',
        'mes_anterior_nombre': f'{MESES[ini_ant.month - 1]} {ini_ant.year}',
        'venta_mes': _n(venta_mes), 'venta_mes_sub': _n(sub_mes), 'venta_mes_anterior': _n(venta_ant),
        'venta_mismo_avance': _n(venta_avance), 'dia_avance': dia_lim,
        'variacion': _n((venta_mes - venta_ant) * 100 / venta_ant, 1) if venta_ant else None,
        'variacion_avance': _n((venta_mes - venta_avance) * 100 / venta_avance, 1) if venta_avance else None,
        'documentos_mes': len(mes), 'ticket_mes': _n(venta_mes / len(mes)) if mes else 0,
        'clientes_mes': len({d['cliente'] for d in mes}),
        'utilidad_mes': _n(sub_mes - costo_mes), 'margen_mes': _n((sub_mes - costo_mes) * 100 / sub_mes, 1) if sub_mes else None,
        'notas_credito_mes': _n(_notas_credito(base, ini, hoy)),
        'por_mes': por_mes,
        'top_clientes': [{'cliente': x['grupo'], 'total': _n(x['total']), 'documentos': x['documentos'],
                          'ultima': x['ultima'].strftime('%d/%m/%Y')} for x in top],
        'cxc_vencido': _n(sum(x['saldo_soles'] for x in vencidas)), 'cxc_vencido_docs': len(vencidas),
        'cxc_total': _n(sum(x['saldo_soles'] for x in cxc)),
        'cxc_por_vencer': _n(sum(x['saldo_soles'] for x in cxc if x['estado'] == 'Por vencer')),
        'cxc_peores': [{'cliente': x['tercero'], 'documento': x['documento'], 'saldo': x['saldo'],
                        'dias': x['dias_atraso']} for x in vencidas[:5]],
        'cxp_vencido': _n(sum(x['saldo_soles'] for x in cxp if x['estado'] == 'Vencido')),
        'cxp_por_vencer': _n(sum(x['saldo_soles'] for x in cxp if x['estado'] == 'Por vencer')),
        'perdidos': [{'cliente': r['cliente'], 'dias': r['dias'], 'total': r['total']} for r in perdidos['filas'][:5]],
        'perdidos_total': len(perdidos['filas']),
        'cotizaciones_mes': int(cot['n']), 'cotizado_mes': _n(cot['monto']), 'cotizaciones_conv': int(cot['conv']),
    }


# ============================================================
# CATÁLOGO DE REPORTES
# ============================================================
AG_VENTAS = [('documento', 'Por comprobante'), ('dia', 'Por día'), ('mes', 'Por mes'), ('empresa', 'Por empresa'),
             ('vendedor', 'Por vendedor'), ('cliente', 'Por cliente'), ('producto', 'Por producto')]
REPORTES = {
    'ventas': {'titulo': 'Ventas', 'tab': 'ventas', 'fn': rep_ventas, 'agrupar': AG_VENTAS,
               'filtros': ['fechas', 'empresa', 'usuario', 'cliente', 'producto'],
               'ayuda': 'Comprobantes emitidos (sin anulados ni borradores).'},
    'utilidad': {'titulo': 'Utilidad y margen', 'tab': 'ventas', 'fn': rep_utilidad,
                 'agrupar': [('producto', 'Por producto'), ('cliente', 'Por cliente'), ('vendedor', 'Por vendedor')],
                 'filtros': ['fechas', 'empresa', 'usuario', 'cliente', 'producto'], 'ayuda': 'Venta sin IGV menos costo.'},
    'top_productos': {'titulo': 'Más vendidos', 'tab': 'ventas', 'fn': rep_top_productos,
                      'filtros': ['fechas', 'empresa', 'usuario', 'cliente', 'limite', 'orden'], 'ayuda': 'Top 10, 20 o 50.'},
    'top_clientes': {'titulo': 'Mejores clientes', 'tab': 'ventas', 'fn': rep_top_clientes,
                     'filtros': ['fechas', 'empresa', 'usuario', 'producto', 'limite'], 'ayuda': 'Clientes con más compras.'},
    'clientes_perdidos': {'titulo': 'Dejaron de comprar', 'tab': 'ventas', 'fn': rep_clientes_perdidos,
                          'filtros': ['empresa', 'dias'], 'dias': 60, 'ayuda': 'Compraban y ya no.'},
    'compras': {'titulo': 'Compras', 'tab': 'compras', 'fn': rep_compras,
                'agrupar': [('documento', 'Por orden de compra'), ('mes', 'Por mes'), ('empresa', 'Por empresa'),
                            ('proveedor', 'Por proveedor'), ('producto', 'Por producto')],
                'filtros': ['fechas', 'empresa', 'proveedor', 'producto'],
                'ayuda': 'Órdenes de compra confirmadas (sin anuladas, borradores ni por aprobar).'},
    'stock': {'titulo': 'Stock', 'tab': 'inventario', 'fn': rep_stock,
              'agrupar': [('detalle', 'Detalle'), ('producto', 'Por producto'), ('empresa', 'Por empresa'), ('sede', 'Por sede')],
              'filtros': ['empresa', 'sede', 'producto'], 'ayuda': 'Stock de hoy valorizado a costo promedio.'},
    'kardex': {'titulo': 'Kardex valorizado', 'tab': 'inventario', 'fn': rep_kardex,
               'filtros': ['fechas', 'empresa', 'sede', 'producto'], 'ayuda': 'Entradas y salidas con costo.'},
    'rotacion': {'titulo': 'Rotación', 'tab': 'inventario', 'fn': rep_rotacion,
                 'filtros': ['fechas', 'empresa', 'sede', 'producto'], 'ayuda': 'Qué se vende rápido y qué no.'},
    'inmovilizado': {'titulo': 'Stock inmovilizado', 'tab': 'inventario', 'fn': rep_inmovilizado,
                     'filtros': ['empresa', 'sede', 'producto', 'dias'], 'dias': 90, 'ayuda': 'Sin movimiento en X días.'},
    'cxc': {'titulo': 'Cuentas por cobrar', 'tab': 'cxc', 'fn': rep_cxc, 'filtros': ['empresa', 'cliente', 'estado'],
            'ayuda': 'Lo que nos deben.'},
    'cxp': {'titulo': 'Cuentas por pagar', 'tab': 'cxp', 'fn': rep_cxp, 'filtros': ['empresa', 'proveedor', 'estado'],
            'ayuda': 'Lo que debemos.'},
    'kpi': {'titulo': 'KPI por colaborador', 'tab': 'kpi', 'fn': rep_kpi, 'filtros': ['fechas', 'empresa', 'usuario'],
            'solo_admin': True, 'ayuda': '¿Quién vende más, cuánto y en cuánto tiempo?'},
}


def _salida(titulo, columnas, filas, sumar, resumen=None, grafico=None):
    totales = {k: _n(sum((r.get(k) or 0) for r in filas if isinstance(r.get(k), (int, float))))
               for k in sumar}
    return {'titulo': titulo, 'columnas': columnas, 'filas': filas, 'totales': totales, 'resumen': resumen or [],
            'grafico': {'etiqueta': grafico[0], 'valor': grafico[1]} if grafico else None}


def _json_val(v):
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d %H:%M')
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, Decimal):
        return float(v)
    return v


def _para_json(t):
    t = dict(t)
    t['filas'] = [{k: _json_val(v) for k, v in r.items()} for r in t['filas']]
    return t


def _texto_filtros(f, info):
    partes = []
    if 'fechas' in info['filtros']:
        partes.append(f"Del {f['desde'].strftime('%d/%m/%Y')} al {f['hasta'].strftime('%d/%m/%Y')}")
    partes.append('Empresa: ' + (_nombre_empresa(f['empresa_id']) if f['empresa_id'] else 'Todas'))
    if f['almacen_id']:
        s = db_query('SELECT nombre FROM almacenes WHERE id = %s', (f['almacen_id'],))
        partes.append('Sede: ' + (s[0]['nombre'] if s else str(f['almacen_id'])))
    if f['usuario_id']:
        partes.append('Usuario: ' + _usuarios().get(f['usuario_id'], {}).get('nombre', str(f['usuario_id'])))
    for k, t in (('cliente', 'Cliente'), ('proveedor', 'Proveedor'), ('producto', 'Producto'), ('estado', 'Estado')):
        if f[k] and k in info['filtros'] + ['producto']:
            partes.append(f'{t}: {f[k]}')
    if 'dias' in info['filtros']:
        partes.append(f"Días: {f['dias']}")
    return ' · '.join(partes)


def _generar(tipo, args=None):
    info = REPORTES[tipo]
    a = dict((args if args is not None else request.args).items())
    if 'dias' in info and not a.get('dias'):
        a['dias'] = str(info['dias'])
    f = leer_filtros(a)
    _asegurar()
    t = info['fn'](f)
    t['filtros_texto'] = _texto_filtros(f, info)
    return t


# ============================================================
# EXPORTAR (Excel, CSV, PDF)
# ============================================================
def _valor_export(v, tipo):
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%d/%m/%Y %H:%M') if tipo == 'fechahora' else v.strftime('%d/%m/%Y')
    if isinstance(v, date):
        return v.strftime('%d/%m/%Y')
    return v


def _nombre_archivo(tipo, ext):
    return f"reporte_{tipo}_{datetime.now().strftime('%Y%m%d_%H%M')}.{ext}"


def exportar_csv(tipo, t):
    sio = io.StringIO()
    w = csv.writer(sio, delimiter=';', lineterminator='\r\n')
    w.writerow([t['titulo']])
    w.writerow([t['filtros_texto']])
    w.writerow([c['t'] for c in t['columnas']])
    for r in t['filas']:
        fila = []
        for c in t['columnas']:
            v = _valor_export(r.get(c['k']), c['tipo'])
            fila.append(f'{v:.2f}' if isinstance(v, float) else v)
        w.writerow(fila)
    if t['totales']:
        w.writerow([('TOTAL' if i == 0 else (f"{t['totales'][c['k']]:.2f}" if c['k'] in t['totales'] else ''))
                    for i, c in enumerate(t['columnas'])])
    return Response('﻿' + sio.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition': f'attachment; filename={_nombre_archivo(tipo, "csv")}'})


def exportar_xlsx(tipo, t):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r'[\\/*?:\[\]]', '', t['titulo'])[:31] or 'Reporte'
    n = len(t['columnas'])
    ws.cell(row=1, column=1, value=t['titulo']).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=t['filtros_texto']).font = Font(italic=True, color='475569')
    ws.cell(row=3, column=1, value=f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')} por "
                                   f"{session.get('nombre_completo') or session.get('usuario') or ''}").font = Font(color='64748B', size=9)
    borde = Border(bottom=Side(style='thin', color='CBD5E1'))
    for i, c in enumerate(t['columnas'], 1):
        cel = ws.cell(row=5, column=i, value=c['t'])
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = PatternFill('solid', fgColor='1E3A8A')
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    fila = 6
    formatos = {'dinero': '#,##0.00', 'num': '#,##0.##', 'pct': '0.0"%"', 'margen': '0.0"%"', 'semaforo': '0.0"%"',
                'fecha': 'DD/MM/YYYY', 'fechahora': 'DD/MM/YYYY HH:MM', 'dias1': '0.0'}
    rojo = PatternFill('solid', fgColor='FEE2E2')
    ambar = PatternFill('solid', fgColor='FEF3C7')
    verde = PatternFill('solid', fgColor='DCFCE7')
    for r in t['filas']:
        for i, c in enumerate(t['columnas'], 1):
            v = r.get(c['k'])
            if isinstance(v, datetime) and c['tipo'] == 'fecha':
                v = v.date()
            cel = ws.cell(row=fila, column=i, value=v)
            cel.border = borde
            if c['tipo'] in formatos and v is not None and not isinstance(v, str):
                cel.number_format = formatos[c['tipo']]
            if c['tipo'] == 'estado':
                cel.fill = {'Vencido': rojo, 'Por vencer': ambar, 'Al día': verde}.get(v, PatternFill())
                cel.font = Font(bold=True)
        fila += 1
    if t['totales']:
        ws.cell(row=fila, column=1, value='TOTAL').font = Font(bold=True)
        for i, c in enumerate(t['columnas'], 1):
            if c['k'] in t['totales']:
                cel = ws.cell(row=fila, column=i, value=t['totales'][c['k']])
                cel.font = Font(bold=True)
                cel.number_format = formatos.get(c['tipo'], '#,##0.00')
                cel.fill = PatternFill('solid', fgColor='FEF3C7')
    for i, c in enumerate(t['columnas'], 1):
        largo = max([len(str(c['t']))] + [len(str(_valor_export(r.get(c['k']), c['tipo']))) for r in t['filas'][:300]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(largo + 2, 9), 55)
    ws.freeze_panes = 'A6'
    if t['filas']:
        ws.auto_filter.ref = f'A5:{get_column_letter(n)}{5 + len(t["filas"])}'
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': f'attachment; filename={_nombre_archivo(tipo, "xlsx")}'})


def _fmt_pdf(v, tipo):
    v = _valor_export(v, tipo)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if tipo in ('pct', 'margen', 'semaforo'):
            return f'{v:.1f}%'
        if tipo == 'dinero':
            return f'{v:,.2f}'
        return f'{v:,.2f}'.rstrip('0').rstrip('.') if isinstance(v, float) else f'{v:,}'
    return str(v)


def exportar_pdf(tipo, t):
    from html import escape
    from weasyprint import HTML
    num = {'dinero', 'num', 'pct', 'margen', 'semaforo', 'dias1', 'atraso'}
    cab = ''.join(f'<th>{escape(c["t"])}</th>' for c in t['columnas'])
    cuerpo = []
    for r in t['filas']:
        tds = []
        for c in t['columnas']:
            v = r.get(c['k'])
            cls = 'n' if c['tipo'] in num else ''
            if c['tipo'] == 'estado':
                cls = {'Vencido': 'rojo', 'Por vencer': 'ambar', 'Al día': 'verde'}.get(v, '')
            tds.append(f'<td class="{cls}">{escape(_fmt_pdf(v, c["tipo"]))}</td>')
        cuerpo.append('<tr>' + ''.join(tds) + '</tr>')
    if t['totales']:
        cuerpo.append('<tr class="tot">' + ''.join(
            f'<td class="n">{escape(_fmt_pdf(t["totales"][c["k"]], c["tipo"]))}</td>' if c['k'] in t['totales']
            else ('<td>TOTAL</td>' if i == 0 else '<td></td>') for i, c in enumerate(t['columnas'])) + '</tr>')
    resumen = ''.join(f'<div class="k"><small>{escape(x["t"])}</small><b>{escape(_fmt_pdf(x["v"], x["tipo"]))}</b></div>'
                      for x in t['resumen'])
    color = (empresa_actual() or {}).get('color_primario') or '#1E3A8A'
    html = f"""<html><head><meta charset="utf-8"><style>
        @page {{ size: A4 landscape; margin: 12mm; @bottom-right {{ content: "Página " counter(page) " de " counter(pages); font-size: 8pt; }} }}
        body {{ font-family: DejaVu Sans, Arial, sans-serif; font-size: 8pt; color: #0F172A; }}
        h1 {{ font-size: 15pt; margin: 0; color: {color}; }}
        .f {{ color: #475569; margin: 2px 0 8px; }}
        .ks {{ display: flex; gap: 8px; margin-bottom: 8px; }}
        .k {{ border: 1px solid #CBD5E1; border-left: 4px solid {color}; border-radius: 6px; padding: 4px 8px; }}
        .k small {{ display: block; color: #64748B; font-size: 7pt; }} .k b {{ font-size: 11pt; }}
        table {{ width: 100%; border-collapse: collapse; }}
        th {{ background: #1E3A8A; color: #fff; padding: 4px; font-size: 7.5pt; }}
        td {{ border-bottom: 1px solid #E2E8F0; padding: 3px 4px; }}
        td.n {{ text-align: right; white-space: nowrap; }}
        td.rojo {{ background: #FEE2E2; color: #B91C1C; font-weight: bold; }}
        td.ambar {{ background: #FEF3C7; color: #92400E; font-weight: bold; }}
        td.verde {{ background: #DCFCE7; color: #166534; font-weight: bold; }}
        tr.tot td {{ background: #FEF3C7; font-weight: bold; }}
        </style></head><body>
        <h1>{escape(t['titulo'])}</h1>
        <div class="f">{escape(t['filtros_texto'])} · Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}</div>
        <div class="ks">{resumen}</div>
        <table><thead><tr>{cab}</tr></thead><tbody>{''.join(cuerpo) or '<tr><td>Sin datos</td></tr>'}</tbody></table>
        {f'<p class="f">{escape(t["nota"])}</p>' if t.get('nota') else ''}
        </body></html>"""
    pdf = HTML(string=html).write_pdf()
    return Response(pdf, mimetype='application/pdf',
                    headers={'Content-Disposition': f'attachment; filename={_nombre_archivo(tipo, "pdf")}'})


# ============================================================
# RUTAS
# ============================================================
def _permiso_reporte(tipo):
    if tipo not in REPORTES:
        return jsonify({'success': False, 'error': 'Ese reporte no existe.'}), 404
    if not puede_ver_montos():
        return _prohibido()
    if REPORTES[tipo].get('solo_admin') and not puede_ver_kpi():
        return _prohibido('El KPI por colaborador solo lo ve la gerencia / administración.')
    return None


@reportes_bp.before_request
def _solo_erika():
    # REGLA DE ORO (Erika 25/10/2026): Reportes solo para Erika: pantalla, consultas y exportaciones
    from utils import es_duena_reportes
    if 'usuario_id' not in session:
        return None  # cada ruta ya pide iniciar sesión
    if not es_duena_reportes():
        if '/api/' in request.path:
            return jsonify({'success': False, 'error': 'Reportes es solo para Erika.'}), 403
        from flask import flash, redirect
        flash('Reportes es solo para Erika.', 'warning')
        return redirect('/index')
    return None


@reportes_bp.route('/reportes')
def reportes():
    if not puede_ver_montos():
        return render_template('reportes.html', permitido=False, es_admin=False, catalogo={}), 403
    catalogo = {k: {'titulo': v['titulo'], 'tab': v['tab'], 'filtros': v['filtros'], 'agrupar': v.get('agrupar') or [],
                    'ayuda': v['ayuda'], 'dias': v.get('dias'), 'orden': i}
                for i, (k, v) in enumerate(REPORTES.items()) if not v.get('solo_admin') or puede_ver_kpi()}
    return render_template('reportes.html', permitido=True, es_admin=puede_ver_kpi(), catalogo=catalogo,
                           empresa_fija=not es_todas() and empresa_filtro() is not None)


@reportes_bp.route('/reportes/api/catalogos')
def api_catalogos():
    if not puede_ver_montos():
        return _prohibido()
    try:
        _asegurar()
        usuarios = _usuarios()
        clientes = [r['n'] for r in db_query(
            "SELECT DISTINCT razon_social AS n FROM clientes WHERE COALESCE(razon_social, '') <> '' ORDER BY 1 LIMIT 3000")]
        proveedores = [r['n'] for r in db_query(
            "SELECT DISTINCT razon_social AS n FROM proveedores WHERE COALESCE(razon_social, '') <> '' ORDER BY 1 LIMIT 3000")]
        productos = [f"{r['codigo']} · {r['descripcion']}" for r in db_query(
            "SELECT codigo, descripcion FROM productos WHERE COALESCE(activo, true) ORDER BY descripcion LIMIT 5000")]
        sedes = [{'id': r['id'], 'nombre': r['nombre']} for r in db_query(
            'SELECT id, nombre FROM almacenes WHERE COALESCE(activo, true) = true ORDER BY id')] if _existe('almacenes') else []
        actual = empresa_actual()
        return jsonify({'success': True,
                        'empresas': [{'id': e['id'], 'codigo': e['codigo']} for e in lista_empresas()],
                        'empresa_fija': ({'id': actual['id'], 'codigo': actual['codigo']} if actual and not es_todas() else None),
                        'sedes': sedes,
                        'usuarios': sorted(({'id': k, 'nombre': v['nombre']} for k, v in usuarios.items()
                                            if v['estado'] not in ('inactivo', 'inactiva')), key=lambda u: u['nombre'].lower()),
                        'clientes': clientes, 'proveedores': proveedores, 'productos': productos,
                        'hoy': date.today().isoformat(), 'inicio_mes': date.today().replace(day=1).isoformat()})
    except Exception as e:
        print(f'❌ reportes catálogos: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@reportes_bp.route('/reportes/api/tablero')
def api_tablero():
    if not puede_ver_montos():
        return _prohibido()
    try:
        _asegurar()
        return jsonify({'success': True, 'data': tablero(leer_filtros())})
    except Exception as e:
        print(f'❌ reportes tablero: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@reportes_bp.route('/reportes/api/reporte/<tipo>')
def api_reporte(tipo):
    bloqueo = _permiso_reporte(tipo)
    if bloqueo:
        return bloqueo
    try:
        return jsonify({'success': True, 'data': _para_json(_generar(tipo))})
    except ErrorReporte as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        print(f'❌ reporte {tipo}: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500


@reportes_bp.route('/reportes/api/reporte/<tipo>/exportar')
def api_exportar(tipo):
    bloqueo = _permiso_reporte(tipo)
    if bloqueo:
        return bloqueo
    formato = (request.args.get('formato') or 'xlsx').lower()
    try:
        t = _generar(tipo)
        if formato == 'csv':
            return exportar_csv(tipo, t)
        if formato == 'pdf':
            return exportar_pdf(tipo, t)
        if formato == 'xlsx':
            return exportar_xlsx(tipo, t)
        return jsonify({'success': False, 'error': 'Formato no válido (xlsx, csv o pdf).'}), 400
    except ErrorReporte as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        print(f'❌ exportar {tipo}: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500
