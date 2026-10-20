# routes/exportar.py - Botón "Exportar" de todas las pantallas (pedido de Erika 19/10/2026)
#
# El navegador manda la tabla TAL COMO SE VE (con los filtros y la búsqueda aplicados: solo las filas visibles)
# y aquí se arma un Excel .xlsx con:
#   - título del módulo, fecha y hora de descarga (hora de Lima), quién lo descargó y la empresa;
#   - los filtros que estaban puestos;
#   - los mismos nombres de columna que se ven en pantalla.
# El nombre del archivo lleva la fecha y hora de Lima: ventas_cotizaciones_2026-10-19_1530.xlsx
import io
import re
import unicodedata
from datetime import datetime
from zoneinfo import ZoneInfo

from flask import Blueprint, Response, jsonify, request, session

from routes.auditoria import registrar_auditoria

exportar_bp = Blueprint('exportar', __name__)
LIMA = ZoneInfo('America/Lima')
MAX_FILAS = 20000
MAX_COLUMNAS = 80
RE_NUMERO = re.compile(r'^(?:S/|US\$|\$)?\s*-?[\d,]*\.?\d+$')


def ahora_lima():
    return datetime.now(LIMA).replace(tzinfo=None)


def _slug(texto):
    t = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', '_', t).strip('_')[:60] or 'datos'


# Columnas que son códigos o documentos: se quedan como texto (RUC, DNI, Nr°, teléfono...), nunca como número
RE_COLUMNA_TEXTO = re.compile(r'c[óo]d|nr°|n°|nro|ruc|dni|documento|tel[ée]fono|celular|serie|lote|cuenta|operaci', re.I)


def _valor(v, como_texto=False):
    """Texto de la celda -> número cuando es un número o un monto ("S/ 1,250.00"); si no, texto tal cual."""
    if v is None:
        return ''
    s = str(v).strip()
    if not s:
        return ''
    if not como_texto and RE_NUMERO.match(s) and not re.match(r'^0\d', s.replace('S/', '').strip()):
        limpio = re.sub(r'[^\d.\-]', '', s)
        try:
            n = float(limpio)
            return int(n) if n.is_integer() and '.' not in limpio else n
        except ValueError:
            pass
    return s[:32000]


def armar_excel(titulo, columnas, filas, filtros, usuario, empresa):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r'[\[\]\*\?/\\:]', ' ', titulo)[:31] or 'Datos'
    ahora = ahora_lima()
    ws.append([titulo])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([f'Descargado el {ahora:%d/%m/%Y} a las {ahora:%H:%M} (hora de Lima)'
               + (f' por {usuario}' if usuario else '') + (f' · Empresa: {empresa}' if empresa else '')])
    ws.append(['Filtros: ' + ('; '.join(f'{k}: {v}' for k, v in filtros) if filtros else 'ninguno (se exportó todo lo que se ve)')])
    ws.append([f'{len(filas)} filas'])
    ws.append([])
    ws.append(columnas)
    fila_enc = ws.max_row
    relleno = PatternFill('solid', fgColor='1F2937')
    for c in ws[fila_enc]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = relleno
        c.alignment = Alignment(vertical='center', wrap_text=True)
    anchos = [len(str(c)) for c in columnas]
    texto = [bool(RE_COLUMNA_TEXTO.search(c)) for c in columnas]
    for f in filas:
        valores = [_valor(x, texto[i]) for i, x in enumerate(f[:len(columnas)])]
        ws.append(valores)
        for celda in ws[ws.max_row]:
            if isinstance(celda.value, str) and celda.value.startswith('='):
                celda.data_type = 's'   # se ve como texto, nunca como fórmula
        for i, v in enumerate(valores):
            anchos[i] = max(anchos[i], min(len(str(v)), 60))
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(8, min(a + 2, 62))
    ws.freeze_panes = ws.cell(row=fila_enc + 1, column=1)
    if filas:
        ws.auto_filter.ref = f'A{fila_enc}:{get_column_letter(len(columnas))}{ws.max_row}'
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue(), ahora


@exportar_bp.route('/api/exportar/tabla', methods=['POST'])
def api_exportar_tabla():
    if 'usuario_id' not in session:
        return jsonify({'success': False, 'error': 'Sesión expirada. Inicia sesión.'}), 401
    d = request.get_json(silent=True) or {}
    titulo = str(d.get('titulo') or 'Datos').strip()[:80]
    columnas = [str(c or '').strip()[:120] or f'Columna {i + 1}' for i, c in enumerate((d.get('columnas') or [])[:MAX_COLUMNAS])]
    filas = [[('' if x is None else str(x)) for x in (f or [])] for f in (d.get('filas') or [])[:MAX_FILAS]
             if isinstance(f, list)]
    filtros = [(str(k)[:80], str(v)[:200]) for k, v in (d.get('filtros') or [])[:30]
               if isinstance(k, str) and str(v or '').strip()]
    if not columnas:
        return jsonify({'success': False, 'error': 'No hay ninguna tabla para exportar en esta pantalla.'}), 400
    usuario = session.get('nombre_completo') or session.get('usuario') or ''
    empresa = (session.get('empresa') or '').upper()
    contenido, ahora = armar_excel(titulo, columnas, filas, filtros, usuario, empresa)
    nombre = f'{_slug(titulo)}_{ahora:%Y-%m-%d_%H%M}.xlsx'
    registrar_auditoria('exportar', titulo, tabla=None,
                        detalle=f'Exportó "{titulo}" a Excel ({len(filas)} filas'
                                + (', filtros: ' + '; '.join(f'{k}: {v}' for k, v in filtros) if filtros else '') + ')')
    return Response(contenido, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': f'attachment; filename="{nombre}"',
                             'X-Nombre-Archivo': nombre, 'Cache-Control': 'no-store'})
