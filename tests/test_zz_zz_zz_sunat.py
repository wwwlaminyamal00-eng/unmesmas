"""
Módulo SUNAT e impuestos (pedido de Erika del 21/10/2026). Datos inventados, solo en la base local.
  - Solo gerencia o quien tenga el permiso SUNAT (Reportes ya no da acceso).
  - IGV con el orden del PDT 621: ventas - NC - compras - saldo a favor - percepciones - retenciones.
  - Lo importado del SIRE/Excel se une al ERP sin contar dos veces el mismo comprobante.
  - La hoja del contador se lee y se compara; avisa si su total no suma la ONP.
  - Pagos del mes con vencimiento SUNAT por último dígito del RUC, check de pagado y candado para desmarcar.
"""
import io
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from conftest import _login
from tests.test_flujos import sql
from tests.test_permisos_usuarios import base, _usuario, _guardar   # noqa: F401  (base es fixture)
from werkzeug.security import generate_password_hash

PERIODO = '2026-03'


@pytest.fixture(scope='module')
def kcf(base):
    sql("UPDATE erp_empresas SET ruc = '20602095704' WHERE codigo = 'KCF'")
    from routes import multiempresa
    multiempresa.empresas(refrescar=True)
    eid = base['emp']['KCF']
    sql("""INSERT INTO comprobantes (tipo_comprobante, serie, numero, fecha_emision, cliente_tipo_doc, cliente_numero_doc,
               cliente_nombre, subtotal, igv, total, items_json, estado_sunat, empresa_id)
           VALUES ('Factura', 'F901', 1, '2026-03-05', 'RUC', '20100070970', 'CLIENTE SUNAT S.A.C.', 1000, 180, 1180,
                   '[]', 'Emitido', %s),
                  ('Factura', 'F901', 2, '2026-03-06', 'RUC', '20100070970', 'CLIENTE SUNAT S.A.C.', 500, 90, 590,
                   '[]', 'Borrador', %s)""", (eid, eid))
    sql("""INSERT INTO notas_credito (serie, numero, fecha_emision, cliente_numero_doc, cliente_nombre, monto, estado,
               comprobante_asociado, empresa_id)
           VALUES ('FC01', '901', '2026-03-10', '20100070970', 'CLIENTE SUNAT S.A.C.', 118, 'Emitida', 'F901-1', %s)""", (eid,))
    sql("""INSERT INTO comprobantes_proveedor (tipo, numero, fecha, monto, ruc_proveedor, proveedor_nombre, estado, empresa_id)
           VALUES ('Factura', 'F001-555', '2026-03-08', 590, '20600011111', 'PROVEEDOR ACEROS S.A.C.', 'Pagado', %s),
                  ('Boleta', 'B001-9', '2026-03-09', 100, '20600011111', 'PROVEEDOR ACEROS S.A.C.', 'Pagado', %s)""", (eid, eid))
    return eid


def _resumen(cliente, eid, periodo=PERIODO):
    r = cliente.get(f'/sunat/api/resumen?empresa_id={eid}&periodo={periodo}')
    assert r.status_code == 200, r.json
    return r.json['data']


def test_solo_gerencia_o_permiso_sunat(app, admin, vendedor, base, kcf):
    assert admin.get('/sunat').status_code == 200
    html = admin.get('/sunat').get_data(as_text=True)
    assert 'SUNAT e impuestos' in html and 'Cálculo preliminar' in html and 'data-modulo="sunat"' in html
    r = vendedor.get('/sunat/api/resumen')
    assert r.status_code == 403 and r.json['codigo'] == 'sin_permiso'
    assert vendedor.get('/sunat').status_code == 302
    # Usuario configurado: Reportes ya no abre SUNAT; el permiso SUNAT sí
    sql("""INSERT INTO erp_modulos (orden, codigo, nombre, estado) SELECT 8, 'REPORTES', 'Reportes', 'activo'
           WHERE NOT EXISTS (SELECT 1 FROM erp_modulos WHERE codigo = 'REPORTES')""")
    for i, cod in enumerate(('REPORTE_VENTAS', 'SUNAT'), 1):
        sql("""INSERT INTO erp_submodulos (modulo_id, orden, codigo, nombre, estado)
               SELECT m.id, %s, %s, %s, 'activo' FROM erp_modulos m
               WHERE m.codigo = 'REPORTES' AND NOT EXISTS (SELECT 1 FROM erp_submodulos WHERE codigo = %s)""", (i, cod, cod, cod))
    uid, _ = _usuario(base, 'sunat_helen', empresas=('KCF',))
    sql("UPDATE usuarios SET password = %s WHERE id = %s", (generate_password_hash('Sunat#2026!'), uid))
    _guardar(admin, uid, base['emp']['KCF'], [{'submodulo_codigo': 'REPORTE_VENTAS', 'puede_ver': True}])
    from routes.acceso import limpiar_cache, MENU_CODIGOS
    assert 'SUNAT' not in MENU_CODIGOS['reportes'] and MENU_CODIGOS['sunat'] == ['SUNAT']
    limpiar_cache()
    h = _login(app, 'sunat_helen', 'Sunat#2026!')
    html = h.get('/index').get_data(as_text=True)
    assert ':not([data-modulo="reportes"])' in html and ':not([data-modulo="sunat"])' not in html
    assert h.get('/sunat/api/resumen').status_code == 403
    _guardar(admin, uid, base['emp']['KCF'], [{'submodulo_codigo': 'SUNAT', 'puede_ver': True}])
    limpiar_cache()
    assert h.get('/sunat').status_code == 200
    assert h.get(f'/sunat/api/resumen?empresa_id={kcf}&periodo={PERIODO}').status_code == 200


def test_cronograma_por_ruc_y_dias_habiles():
    from routes.sunat import vencimiento_cronograma, _habil_n, _ultimo_habil
    assert vencimiento_cronograma('20602095704', 2026, 9) == date(2026, 10, 21)   # KCF: termina en 4
    assert vencimiento_cronograma('20610910671', 2026, 9) == date(2026, 10, 19)   # AGD: termina en 1
    assert vencimiento_cronograma('20602095704', 2026, 10) == date(2026, 11, 19)
    assert vencimiento_cronograma('20602095700', 2026, 12) == date(2027, 1, 18)
    assert _habil_n(2026, 10, 5) == date(2026, 10, 7)      # AFP de setiembre
    assert _habil_n(2026, 4, 5) == date(2026, 4, 9)        # Semana Santa no cuenta
    assert _ultimo_habil(2026, 10) == date(2026, 10, 30)


def test_liquidar_igual_que_el_contador():
    from routes.sunat import liquidar
    # Hoja SET 2026 del contador de KCF
    r = liquidar(Decimal('9822'), Decimal('6501.49'), {'retenciones_mes': '2732.14', 'saldo_retenciones_anterior': '14369'})
    assert r['tributo'] == Decimal('3320.51') and r['igv_pagar'] == 0
    assert r['importe_contador'] == Decimal('-13780.63') and r['saldo_retenciones_siguiente'] == Decimal('13780.63')
    # Saldo a favor: las retenciones pasan completas al mes siguiente
    r = liquidar(Decimal('100'), Decimal('500'), {'retenciones_mes': '30'})
    assert r['igv_pagar'] == 0 and r['saldo_favor_siguiente'] == 400 and r['saldo_retenciones_siguiente'] == 30
    r = liquidar(Decimal('1000'), Decimal('400'), {'retenciones_mes': '100', 'percepciones_mes': '50'})
    assert r['igv_pagar'] == 450


def test_resumen_igv_con_notas_y_retenciones(admin, kcf):
    d = _resumen(admin, kcf)
    assert d['sv']['n_facturas'] == 1 and d['sv']['n_nc'] == 1          # el borrador no cuenta
    assert d['sv']['base'] == 900 and d['sv']['igv'] == 162             # 1000 - 100 de la NC
    assert d['sc']['n_facturas'] == 1 and d['sc']['base'] == 500 and d['sc']['igv'] == 90
    assert d['sc']['no_gravado'] == 100                                  # la boleta no da crédito fiscal
    assert d['liq']['igv_pagar'] == 72 and d['renta']['monto'] == 9      # 1% de 900 (Régimen MYPE)
    r = admin.post('/sunat/api/periodo', json={'empresa_id': kcf, 'periodo': PERIODO,
                                                'manual': {'retenciones_mes': '10', 'saldo_retenciones_anterior': '50'},
                                                'planilla': {'base_essalud': '1130', 'base_onp': '1000'}})
    assert r.status_code == 200, r.json
    d = _resumen(admin, kcf)
    assert d['liq']['igv_pagar'] == 12 and d['plame']['essalud'] == 101.7 and d['plame']['onp'] == 130
    assert d['total_sunat'] == 12 + 9 + 101.7 + 130
    assert admin.post('/sunat/api/periodo', json={'empresa_id': kcf, 'periodo': PERIODO,
                                                  'manual': {'retenciones_mes': '-5'}}).status_code == 400
    assert any(p['grupo'] == 'Reglas de oro' for p in d['plan'])
    assert any('Completa las facturas' in p['titulo'] for p in d['plan'])   # todavía no se importó SIRE


def _txt_sire(filas):
    cab = ('RUC|Apellidos y Nombres o Razón social|Periodo|CAR SUNAT|Fecha de emisión|Fecha Vcto/Pago|Tipo CP/Doc.|'
           'Serie del CDP|Nro CP o Doc. Nro Inicial (Rango)|Nro Final (Rango)|Tipo Doc Identidad|Nro Doc Identidad|'
           'Apellidos Nombres/ Razón Social|Valor Facturado Exportación|BI Gravada|Dscto BI|IGV / IPM|Dscto IGV / IPM|'
           'Mto Exonerado|Mto Inafecto|ISC|BI Grav IVAP|IVAP|ICBPER|Otros Tributos|Total CP|Moneda|Tipo Cambio')
    lineas = [cab]
    for f in filas:
        lineas.append('|'.join(['20602095704', 'KCF', '202603', '', f[0], '', f[1], f[2], f[3], '', '6', f[4], f[5], '0',
                                f[6], '0', f[7], '0', '0', '0', '0', '0', '0', '0', '0', f[8], 'PEN', '']))
    return '\n'.join(lineas).encode('utf-8')


def test_importar_sire_sin_duplicar(admin, kcf):
    antes = sql("SELECT COUNT(*) FROM comprobantes")[0][0]
    contenido = _txt_sire([
        ('05/03/2026', '01', 'F901', '00000001', '20100070970', 'CLIENTE SUNAT S.A.C.', '1000.00', '180.00', '1180.00'),  # ya está en el ERP
        ('15/03/2026', '01', 'F901', '00000003', '20555555555', 'OTRO CLIENTE S.A.C.', '2000.00', '360.00', '2360.00'),
        ('20/04/2026', '01', 'F901', '00000004', '20555555555', 'OTRO MES S.A.C.', '50.00', '9.00', '59.00'),
    ])
    r = admin.post('/sunat/api/importar/vista', data={'empresa_id': kcf, 'periodo': PERIODO, 'libro': 'VENTAS',
                                                      'archivo': (io.BytesIO(contenido), 'LE2060209570420260300140400.txt')},
                   content_type='multipart/form-data')
    assert r.status_code == 200, r.json
    v = r.json['data']
    assert v['resumen']['n'] == 2 and v['resumen']['en_erp'] == 1 and any('otro mes' in a for a in v['avisos'])
    r = admin.post('/sunat/api/importar/guardar', json={'empresa_id': kcf, 'periodo': PERIODO, 'libro': 'VENTAS',
                                                         'origen': 'SIRE', 'archivo': v['archivo'], 'filas': v['filas']})
    assert r.status_code == 200, r.json
    d = _resumen(admin, kcf)
    assert d['sv']['n_facturas'] == 2 and d['sv']['base'] == 2900 and d['sv']['igv'] == 522
    assert d['sv']['n_ambos'] == 1 and d['sv']['n_archivo'] == 1
    docs = admin.get(f'/sunat/api/documentos?empresa_id={kcf}&periodo={PERIODO}&libro=VENTAS').json['data']
    assert {x['fuente'] for x in docs} == {'ERP y archivo', 'Solo archivo', 'Solo ERP'}   # la NC sigue solo en el ERP
    # Quitar la importación deja todo como antes y no toca el ERP
    lote = admin.get(f'/sunat/api/importados?empresa_id={kcf}&periodo={PERIODO}').json['data'][0]['lote']
    assert admin.post('/sunat/api/importados/quitar', json={'empresa_id': kcf, 'lote': lote}).status_code == 200
    assert _resumen(admin, kcf)['sv']['base'] == 900
    assert sql("SELECT COUNT(*) FROM comprobantes")[0][0] == antes


def test_importar_excel_de_las_chicas(admin, kcf):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Fecha de emisión', 'Tipo', 'Serie', 'Número', 'Proveedor', 'RUC / DNI', 'Base imponible', 'IGV',
               'No gravado', 'Total', 'Moneda', 'Tipo de cambio'])
    ws.append([date(2026, 3, 8), 'Factura', 'F001', '555', 'PROVEEDOR ACEROS S.A.C.', '20600011111', 500, 90, 0, 590, 'PEN', None])
    ws.append([date(2026, 2, 20), 'Factura', 'E001', '77', 'SERVICIOS SAC', '20777777777', 1000, 180, 0, 1180, 'PEN', None])
    out = io.BytesIO()
    wb.save(out)
    r = admin.post('/sunat/api/importar/vista', data={'empresa_id': kcf, 'periodo': PERIODO, 'libro': 'COMPRAS',
                                                      'archivo': (io.BytesIO(out.getvalue()), 'compras_marzo.xlsx')},
                   content_type='multipart/form-data')
    assert r.status_code == 200, r.json
    v = r.json['data']
    assert v['resumen']['n'] == 2 and v['resumen']['en_erp'] == 1   # la de febrero se puede anotar en marzo
    assert admin.post('/sunat/api/importar/guardar', json={'empresa_id': kcf, 'periodo': PERIODO, 'libro': 'COMPRAS',
                                                           'origen': 'EXCEL', 'archivo': v['archivo'],
                                                           'filas': v['filas']}).status_code == 200
    d = _resumen(admin, kcf)
    assert d['sc']['n_facturas'] == 2 and d['sc']['igv'] == 270
    r = admin.post('/sunat/api/importar/vista', data={'empresa_id': kcf, 'periodo': PERIODO, 'libro': 'COMPRAS',
                                                      'archivo': (io.BytesIO(b'hola,mundo\n1,2'), 'x.csv')},
                   content_type='multipart/form-data')
    assert r.status_code == 400 and 'títulos' in r.json['error']
    plantilla = admin.get('/sunat/api/plantilla.xlsx')
    assert plantilla.status_code == 200 and plantilla.data[:2] == b'PK'


def _hoja_contador():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'MAR 2026'
    ws['A2'] = 'KCF CORPORACION E.I.R.L.'
    ws['A5'] = 'LIQUIDACION DE IMPUESTOS  Y CONTRIBUCIONES MARZO 2026'
    ws['D7'], ws['E7'] = 'VENCIMIENTO', '21 ABRIL 2026'
    filas = {12: ('B', 'VENTAS NETAS LOCALES', 2900, 522), 19: ('C', 'LOCALES DESTINADAS A VENTAS GRAVADAS', 1500, -270),
             27: ('C', 'IMPUESTO RESULTANTE O SALDO A FAVOR', None, 252), 28: ('C', 'SALDO FAVOR DEL PERIODO ANTERIOR', None, 0),
             29: ('C', 'TRIBUTO A PAGAR O  SALDO A FAVOR', None, 252), 34: ('C', 'RETENCIONES DECLARADAS EN EL PERIODO', None, -10),
             35: ('C', 'SALDO DE RETENCIONES DE PERIODOS ANTERIORES', None, -50),
             38: ('C', 'IMPORTE A PAGAR O SALDO  A FAVOR', None, 192), 48: ('C', 'IMPUESTO DEL MES', None, 29),
             70: ('B', 'ESSALUD', 1130, 101.7), 71: ('B', 'SNP', 1000, 130),
             89: ('B', 'TOTAL TRIBUTOS Y CONTRIBUCIONES POR PAGAR', None, 322.7)}
    for fila, (col, texto, base, imp) in filas.items():
        ws[f'{col}{fila}'] = texto
        if base is not None:
            ws[f'F{fila}'] = base
        ws[f'H{fila}'] = imp
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_hoja_del_contador_se_compara(admin, kcf):
    # Vuelve a cargar las ventas del SIRE para que el preliminar tenga 2,900 de base
    sql("""INSERT INTO sunat_registros (empresa_id, periodo, libro, lote, origen, tipo_cp, serie, numero, fecha, ruc, nombre,
               base, igv, total) VALUES (%s, %s, 'VENTAS', 'L1', 'SIRE', '01', 'F901', '3', '2026-03-15', '20555555555',
               'OTRO CLIENTE S.A.C.', 2000, 360, 2360)""", (kcf, PERIODO))
    r = admin.post('/sunat/api/contador', data={'empresa_id': kcf, 'periodo': PERIODO,
                                                'archivo': (io.BytesIO(_hoja_contador()), 'KCF_IMPUESTOS.xlsx')},
                   content_type='multipart/form-data')
    assert r.status_code == 200, r.json
    assert r.json['data']['ventas_base'] == '2900.00' and r.json['data']['onp'] == '130.00'
    d = _resumen(admin, kcf)
    filas = {f['clave']: f for f in d['preliminar']}
    assert filas['ventas_igv']['contador'] == 522 and not filas['ventas_igv']['alerta']
    assert filas['renta']['contador'] == 29 and not filas['renta']['alerta']   # 1% de 2,900
    assert filas['importe_igv']['erp'] == 192 and not filas['importe_igv']['alerta']
    assert any('ONP' in o for o in d['observaciones_contador'])
    xlsx = admin.get(f'/sunat/api/preliminar.xlsx?empresa_id={kcf}&periodo={PERIODO}')
    assert xlsx.status_code == 200 and xlsx.data[:2] == b'PK'
    wb = openpyxl.load_workbook(io.BytesIO(xlsx.data))
    ws = wb.worksheets[0]
    assert ws['B12'].value == 'VENTAS NETAS LOCALES' and ws['F12'].value == 2900 and 'PRELIMINAR' in ws['A1'].value
    r = admin.post('/sunat/api/contador', data={'empresa_id': kcf, 'periodo': '2026-05',
                                                'archivo': (io.BytesIO(_hoja_contador()), 'KCF.xlsx')},
                   content_type='multipart/form-data')
    assert r.status_code == 400 and 'No encontré la hoja' in r.json['error']


def test_pagos_del_mes_con_candado(admin, vendedor, kcf):
    r = admin.get(f'/sunat/api/pagos?empresa_id={kcf}&mes=2026-04')
    assert r.status_code == 200, r.json
    filas = {f['codigo']: f for f in r.json['data']['filas']}
    assert filas['IGV']['periodo'] == PERIODO and filas['IGV']['vencimiento'] == '2026-04-22'
    assert filas['AFP']['vencimiento'] == '2026-04-09' and filas['REMUNERACIONES']['periodo'] == '2026-04'
    assert filas['ESSALUD']['monto'] == 101.7 and filas['IGV']['estado'] == 'Pendiente'
    assert 'CTS' not in filas and 'GRATIFICACION' not in filas
    assert filas['IGV']['semaforo'] == 'vencido'          # abril ya pasó
    # Sin el N° de operación del voucher no se puede marcar pagado
    r = admin.post('/sunat/api/pagos/marcar', json={'empresa_id': kcf, 'periodo': PERIODO, 'codigo': 'IGV', 'estado': 'Pagado',
                                                    'monto': '192', 'fecha_pago': '2026-04-20'})
    assert r.status_code == 400 and 'operación' in r.json['error']
    r = admin.post('/sunat/api/pagos/marcar', json={'empresa_id': kcf, 'periodo': PERIODO, 'codigo': 'IGV', 'estado': 'Pagado',
                                                    'monto': '192', 'fecha_pago': '2026-04-20', 'numero_operacion': '123456'})
    assert r.status_code == 200, r.json
    filas = {f['codigo']: f for f in admin.get(f'/sunat/api/pagos?empresa_id={kcf}&mes=2026-04').json['data']['filas']}
    assert filas['IGV']['estado'] == 'Pagado' and filas['IGV']['monto'] == 192 and filas['IGV']['marcado_por']
    # Desmarcar exige motivo y queda en el historial
    r = admin.post('/sunat/api/pagos/marcar', json={'empresa_id': kcf, 'periodo': PERIODO, 'codigo': 'IGV', 'estado': 'Pendiente'})
    assert r.status_code == 400 and 'motivo' in r.json['error']
    assert admin.post('/sunat/api/pagos/marcar', json={'empresa_id': kcf, 'periodo': PERIODO, 'codigo': 'IGV',
                                                       'estado': 'Pendiente', 'observacion': 'Me equivoqué de mes'}).status_code == 200
    hist = sql("SELECT historial FROM sunat_obligaciones WHERE codigo = 'IGV' AND periodo = %s", (PERIODO,))[0][0]
    assert [h['estado'] for h in hist] == ['Pagado', 'Pendiente']
    assert admin.post('/sunat/api/pagos/marcar', json={'empresa_id': kcf, 'periodo': PERIODO, 'codigo': 'ONP',
                                                       'estado': 'No aplica'}).status_code == 200
    noviembre = {f['codigo'] for f in admin.get(f'/sunat/api/pagos?empresa_id={kcf}&mes=2026-11').json['data']['filas']}
    assert 'CTS' in noviembre
    assert vendedor.post('/sunat/api/pagos/marcar', json={'empresa_id': kcf, 'periodo': PERIODO, 'codigo': 'IGV',
                                                          'estado': 'Pagado'}).status_code == 403


def test_regimen_solo_gerencia(admin, kcf):
    r = admin.post('/sunat/api/config', json={'empresa_id': kcf, 'regimen': 'GENERAL', 'tasa_renta': '1.5'})
    assert r.status_code == 200, r.json
    assert _resumen(admin, kcf)['renta']['tasa'] == 1.5
    assert admin.post('/sunat/api/config', json={'empresa_id': kcf, 'regimen': 'XYZ'}).status_code == 400
    admin.post('/sunat/api/config', json={'empresa_id': kcf, 'regimen': 'RMT', 'tasa_renta': '1'})


def test_plan_de_accion_mes_actual_no_falla(admin, kcf):
    hoy = date.today()
    d = _resumen(admin, kcf, f'{hoy.year}-{hoy.month:02d}')
    assert d['es_mes_actual'] and {p['clave'] for p in d['proyeccion']['partes']} >= {'guias', 'oc', 'pc', 'compras_extra'}
    assert d['plan'][-1]['grupo'] == 'Reglas de oro'


def test_plan_no_mezcla_empresas(admin, kcf):
    """Una guía sin factura de KCF no debe salir en el plan de AGD (error visto con la copia de datos reales)."""
    agd = sql("SELECT id FROM erp_empresas WHERE codigo = 'AGD'")[0][0]
    sql("""INSERT INTO guias_remision (serie, numero, fecha_emision, destinatario_nombre, ruc_destinatario, motivo_traslado,
               estado_sunat, empresa_id) VALUES ('T991', '777', NOW(), 'CLIENTE SOLO KCF', '20100070970', 'VENTA', 'ACEPTADO', %s)""", (kcf,))
    from routes import sunat
    assert any(g['guia'] == 'T991-777' for g in sunat.guias_sin_factura(kcf))
    assert not any(g['guia'] == 'T991-777' for g in sunat.guias_sin_factura(str(agd)))


def test_sin_hoja_del_contador_el_igv_no_se_marca_no_aplica(admin, kcf):
    """Si el ERP no tiene datos del mes, un IGV en cero puede ser falta de datos: queda Pendiente, sin monto."""
    agd = str(sql("SELECT id FROM erp_empresas WHERE codigo = 'AGD'")[0][0])
    filas = {f['codigo']: f for f in admin.get(f'/sunat/api/pagos?empresa_id={agd}&mes=2026-02').json['data']['filas']}
    assert filas['IGV']['estado'] == 'Pendiente' and filas['IGV']['monto'] is None
    assert filas['RENTA_3RA']['estado'] == 'Pendiente'


def test_resumen_trae_el_voucher_pagado(admin, kcf):
    """El cuadro del resumen muestra lo pagado de verdad (voucher) de IGV y renta."""
    d = _resumen(admin, kcf, PERIODO)
    assert 'voucher' in d
    admin.post('/sunat/api/pagos/marcar', json={'empresa_id': kcf, 'periodo': PERIODO, 'codigo': 'RENTA_3RA', 'estado': 'Pagado',
                                                'monto': '55', 'fecha_pago': '2026-04-20', 'numero_operacion': 'NPS-777'})
    v = _resumen(admin, kcf, PERIODO)['voucher']['RENTA_3RA']
    assert v['estado'] == 'Pagado' and v['monto'] == 55 and v['numero_operacion'] == 'NPS-777'
