"""
Pruebas del módulo 8 - Reportes (pedido de Erika, 16/10/2026).
Crea sus propios datos (clientes "ZZREP ...", productos ZZR-..., usuarios zzrep_*) y verifica los totales de:
ventas, cuentas por cobrar/pagar, top productos, utilidad, compras, inmovilizado, KPI por colaborador,
exportación CSV / Excel / PDF y la protección por rol en el servidor.
"""
import io
import json
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from werkzeug.security import generate_password_hash

from tests.test_inventario import _sql
from tests.conftest import _login

# El ERP trabaja con la fecha de Lima (main.py fija TZ=America/Lima): las pruebas usan esa misma fecha
HOY = datetime.now(ZoneInfo('America/Lima')).date()
RANGO = {'desde': (HOY - timedelta(days=30)).isoformat(), 'hasta': HOY.isoformat()}


def _q(**kw):
    p = dict(RANGO)
    p.update(kw)
    return '&'.join(f'{k}={v}' for k, v in p.items())


@pytest.fixture(scope='module')
def rep(app, admin):
    # Ver "Todas las empresas" para que el resultado no dependa de la empresa que dejaron otras pruebas
    assert admin.post('/api/empresa_activa', json={'codigo': 'TODAS'}).status_code == 200
    usuarios = {}
    for u, rol in (('zzrep_vend1', 'vendedor'), ('zzrep_vend2', 'vendedor'), ('zzrep_fin', 'finanzas'), ('zzrep_ven3', 'vendedor')):
        usuarios[u] = _sql("""INSERT INTO usuarios (usuario_sistema, usuario, nombres_apellidos, nombre_completo, correo, password, rol, estado)
                              VALUES (%s, %s, %s, %s, %s, %s, %s, 'activo') RETURNING id""",
                           (u, u, u.upper(), u.upper(), f'{u}@test.pe', generate_password_hash('Clave#2026!'), rol))[0][0]
    _sql("""
        INSERT INTO productos (familia, codigo, descripcion, marca, unidad, costo_unitario, precio_unitario, stock, activo) VALUES
            ('GENERAL', 'ZZR-001', 'ZZREP BOTA PUNTA ACERO', 'X', 'PAR', 10, 30, 0, true),
            ('GENERAL', 'ZZR-002', 'ZZREP ARNES CUERPO ENTERO', 'X', 'UND', 20, 50, 0, true);
        INSERT INTO erp_empresas (codigo, nombre_corto, nombre_comercial, razon_social, ruc, estado, color_primario, color_secundario, color_pastel)
        SELECT 'KCF', 'KCF', 'KCF', 'KCF S.A.C.', '20111111111', 'activo', '#D6243A', '#1D1A1F', '#FDE8EB' WHERE NOT EXISTS (SELECT 1 FROM erp_empresas WHERE codigo='KCF');
        INSERT INTO almacenes (codigo, nombre, direccion, activo)
        SELECT 'SMP', 'San Martín de Porres', 'Av. Perú 123', true WHERE NOT EXISTS (SELECT 1 FROM almacenes WHERE codigo='SMP');
        INSERT INTO proveedores (razon_social, ruc, direccion, codigo_proveedor, activo)
        VALUES ('ZZREP PROVEEDOR SAC', '20999111111', 'Av. 1', 'PROV-ZZR', true);
    """, fetch=False)
    prod = dict(_sql("SELECT codigo, id FROM productos WHERE codigo LIKE 'ZZR-%'"))
    cli = _sql("SELECT id FROM clientes ORDER BY id LIMIT 1")[0][0]
    v1, v2 = usuarios['zzrep_vend1'], usuarios['zzrep_vend2']

    def cot(numero, uid, estado, total, dias_atras):
        return _sql("""INSERT INTO cotizaciones (numero_cotizacion, cliente_id, estado, subtotal, igv, total, usuario_id, fecha_creacion)
                       VALUES (%s, %s, %s, %s, %s, %s, %s, NOW() - %s * INTERVAL '1 day') RETURNING id""",
                    (numero, cli, estado, total / 1.18, total - total / 1.18, total, uid, dias_atras))[0][0]
    c1 = cot('COT-ZZR-1', v1, 'Aceptada por Cliente', 590, 6)
    cot('COT-ZZR-2', v1, 'Generada', 300, 5)              # no convertida
    cot('COT-ZZR-3', v2, 'Generada', 1180, 4)             # convertida por la factura
    cot('COT-ZZR-4', v1, 'Anulada', 9999, 3)              # anulada: no cuenta
    _sql("""INSERT INTO pedido_compra_pc (numero, cliente, ruc, monto, cotizacion_id, cotizacion_numero, estado, fecha)
            VALUES ('PC-ZZR-1', 'ZZREP CLIENTE UNO', '20999000001', 590, %s, 'COT-ZZR-1', 'PC conforme', NOW() - INTERVAL '4 days')""",
         (c1,), fetch=False)

    def comp(numero, cliente, ruc, sub, estado, items, dias_atras, venc_dias, asociado='', pc=''):
        igv = round(sub * 0.18, 2)
        return _sql("""INSERT INTO comprobantes (tipo_comprobante, serie, numero, fecha_emision, cliente_tipo_doc, cliente_numero_doc,
                           cliente_nombre, subtotal, igv, total, monto_a_pagar, items_json, estado_sunat, documento_asociado,
                           pc_vinculado, es_credito, fecha_vencimiento, creado_por)
                       VALUES ('FACTURA', 'F001', %s, NOW() - %s * INTERVAL '1 day', 'RUC', %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, true,
                               %s::date + %s, %s) RETURNING id""",
                    (numero, dias_atras, ruc, cliente, sub, igv, sub + igv, sub + igv, json.dumps(items), estado, asociado, pc,
                     HOY, venc_dias, str(usuarios['zzrep_ven3'])))[0][0]
    f1 = comp('90000001', 'ZZREP CLIENTE UNO', '20999000001', 500, 'Emitido',
              [{'codigo': 'ZZR-001', 'producto': 'ZZREP BOTA PUNTA ACERO', 'cantidad': 10, 'valorVenta': 30},
               {'codigo': 'ZZR-002', 'producto': 'ZZREP ARNES CUERPO ENTERO', 'cantidad': 5, 'valorVenta': 40}],
              2, 3, pc='PC-ZZR-1')
    comp('90000002', 'ZZREP CLIENTE DOS', '20999000002', 1000, 'Emitido',
         [{'producto_id': prod['ZZR-002'], 'codigo': 'ZZR-002', 'producto': 'ZZREP ARNES CUERPO ENTERO', 'cantidad': 20, 'valorVenta': 50}],
         1, -10, asociado='COT-ZZR-3')
    comp('90000003', 'ZZREP CLIENTE UNO', '20999000001', 8000, 'Anulado', [], 1, 0)     # anulado: no cuenta
    comp('90000004', 'ZZREP CLIENTE UNO', '20999000001', 7000, 'BORRADOR', [], 1, 0)    # borrador: no cuenta
    # Cobro parcial de la primera factura (90) y un cobro anulado (no cuenta)
    _sql("""INSERT INTO finanzas_pagos (tipo, documento_tabla, documento_id, documento_numero, tercero_nombre, fecha, monto, estado)
            VALUES ('COBRO', 'comprobantes', %s, 'F001-90000001', 'ZZREP CLIENTE UNO', CURRENT_DATE, 90, 'Registrado'),
                   ('COBRO', 'comprobantes', %s, 'F001-90000001', 'ZZREP CLIENTE UNO', CURRENT_DATE, 100, 'Anulado')""",
         (f1, f1), fetch=False)
    # Compras: una OC aprobada (2 líneas) y una anulada; factura del proveedor vencida hace 5 días
    prov = _sql("SELECT id FROM proveedores WHERE ruc = '20999111111' ORDER BY id LIMIT 1")[0][0]
    oc = _sql("""INSERT INTO ordenes_compra (numero_orden, proveedor_id, usuario_id, estado, subtotal, igv, total, fecha_creacion)
                 VALUES ('OC-ZZR-1', %s, %s, 'Aprobada', 400, 72, 472, NOW() - INTERVAL '2 days') RETURNING id""", (prov, v1))[0][0]
    _sql("""INSERT INTO ordenes_compra_items (orden_compra_id, producto, cantidad, precio_unitario, total)
            VALUES (%s, 'ZZREP BOTA PUNTA ACERO', 10, 10, 100), (%s, 'ZZREP ARNES CUERPO ENTERO', 15, 20, 300)""", (oc, oc), fetch=False)
    _sql("""INSERT INTO ordenes_compra (numero_orden, proveedor_id, usuario_id, estado, subtotal, igv, total, fecha_creacion)
            VALUES ('OC-ZZR-2', %s, %s, 'Anulada', 5000, 900, 5900, NOW())""", (prov, v1), fetch=False)
    _sql("""INSERT INTO comprobantes_proveedor (tipo, numero, fecha, monto, ruc_proveedor, proveedor_nombre, orden_compra_id, estado)
            VALUES ('Factura', 'E001-ZZR1', NOW() - INTERVAL '35 days', 472, '20999111111', 'ZZREP PROVEEDOR SAC', %s, 'Pendiente')""",
         (oc,), fetch=False)
    # Fechas desde el "hoy" de Python (el mismo que usa el reporte), no CURRENT_DATE: si la base está en otra zona
    # horaria, entre las 19:00 y 24:00 de Lima ambos "hoy" no coinciden y los días de atraso salían con uno de diferencia.
    _sql("UPDATE comprobantes_proveedor SET fecha_vencimiento = %s::date - 5 WHERE numero = 'E001-ZZR1'", (HOY,), fetch=False)
    # Stock parado: 7 botas sin movimiento hace 200 días
    emp = _sql("SELECT id FROM erp_empresas WHERE codigo = 'KCF'")[0][0]
    sede = _sql("SELECT id FROM almacenes WHERE codigo = 'SMP'")[0][0]
    _sql("""INSERT INTO inv_stock (empresa_id, almacen_id, producto_id, cantidad, costo_promedio, valor, ultimo_movimiento)
            VALUES (%s, %s, %s, 7, 10, 70, NOW() - INTERVAL '200 days')""", (emp, sede, prod['ZZR-001']), fetch=False)
    return {'usuarios': usuarios, 'prod': prod}


def _reporte(c, tipo, **kw):
    r = c.get(f'/reportes/api/reporte/{tipo}?' + _q(**kw))
    assert r.status_code == 200, r.get_data(as_text=True)
    return r.get_json()['data']


def test_pagina_y_tablero(rep, admin):
    assert admin.get('/reportes').status_code == 200
    d = admin.get('/reportes/api/tablero').get_json()
    assert d['success'] and len(d['data']['por_mes']) == 12
    assert d['data']['venta_mes'] >= 1770 or HOY.day <= 2   # las facturas de prueba son de hace 1 y 2 días


def test_ventas_totales_excluyen_anulados_y_borradores(rep, admin):
    t = _reporte(admin, 'ventas', cliente='ZZREP', agrupar='documento')
    assert len(t['filas']) == 2
    assert t['totales']['total'] == 1770.0 and t['totales']['subtotal'] == 1500.0
    por_cliente = _reporte(admin, 'ventas', cliente='ZZREP', agrupar='cliente')
    assert {f['grupo'].split(' (')[0]: f['total'] for f in por_cliente['filas']} == {'ZZREP CLIENTE DOS': 1180.0,
                                                                                      'ZZREP CLIENTE UNO': 590.0}


def test_ventas_por_vendedor_sale_de_la_cotizacion(rep, admin):
    t = _reporte(admin, 'ventas', cliente='ZZREP', agrupar='vendedor')
    assert {f['grupo']: f['total'] for f in t['filas']} == {'ZZREP_VEND1': 590.0, 'ZZREP_VEND2': 1180.0}
    solo = _reporte(admin, 'ventas', usuario_id=rep['usuarios']['zzrep_vend2'], agrupar='documento')
    assert [f['cliente'] for f in solo['filas']] == ['ZZREP CLIENTE DOS']


def test_top_productos_y_utilidad(rep, admin):
    t = _reporte(admin, 'top_productos', cliente='ZZREP', limite=10)
    assert [f['codigo'] for f in t['filas']] == ['ZZR-002', 'ZZR-001']
    assert t['filas'][0]['venta'] == 1200.0 and t['filas'][0]['cantidad'] == 25
    assert t['filas'][1]['venta'] == 300.0
    por_cant = _reporte(admin, 'top_productos', cliente='ZZREP', limite=1, orden='cantidad')
    assert len(por_cant['filas']) == 1
    u = _reporte(admin, 'utilidad', cliente='ZZREP', agrupar='producto')
    util = {f['grupo'].split(' · ')[0]: f['utilidad'] for f in u['filas']}
    assert util == {'ZZR-002': 700.0, 'ZZR-001': 200.0}   # 1200 - 25x20 y 300 - 10x10
    assert u['totales']['utilidad'] == 900.0


def test_cuentas_por_cobrar_vencido_primero(rep, admin):
    t = _reporte(admin, 'cxc', cliente='ZZREP')
    assert [(f['estado'], f['saldo']) for f in t['filas']] == [('Vencido', 1180.0), ('Por vencer', 500.0)]
    assert t['filas'][0]['dias_atraso'] == 10
    assert t['totales']['saldo'] == 1680.0
    solo_vencido = _reporte(admin, 'cxc', cliente='ZZREP', estado='Vencido')
    assert len(solo_vencido['filas']) == 1


def test_cuentas_por_pagar_y_compras(rep, admin):
    t = _reporte(admin, 'cxp', proveedor='ZZREP')
    assert len(t['filas']) == 1 and t['filas'][0]['estado'] == 'Vencido' and t['filas'][0]['dias_atraso'] == 5
    c = _reporte(admin, 'compras', proveedor='ZZREP', agrupar='documento')
    assert [f['documento'] for f in c['filas']] == ['OC-ZZR-1'] and c['totales']['total'] == 472.0
    p = _reporte(admin, 'compras', proveedor='ZZREP', agrupar='producto')
    assert {f['producto']: f['subtotal'] for f in p['filas']} == {'ZZREP ARNES CUERPO ENTERO': 300.0, 'ZZREP BOTA PUNTA ACERO': 100.0}


def test_stock_inmovilizado(rep, admin):
    t = _reporte(admin, 'inmovilizado', producto='ZZR-001', dias=90)
    assert len(t['filas']) == 1 and t['filas'][0]['valor'] == 70.0 and t['filas'][0]['dias'] >= 199
    assert _reporte(admin, 'inmovilizado', producto='ZZR-001', dias=300)['filas'] == []
    s = _reporte(admin, 'stock', producto='ZZR-001')
    assert s['totales']['valor'] == 70.0


def test_kpi_por_colaborador(rep, admin):
    t = _reporte(admin, 'kpi')
    fila = {f['colaborador']: f for f in t['filas']}
    v1 = fila['ZZREP_VEND1']
    assert v1['cotizaciones'] == 2 and v1['convertidas'] == 1 and v1['conversion'] == 50.0   # la anulada no cuenta
    assert v1['venta'] == 590.0 and v1['ventas'] == 1 and v1['ticket'] == 590.0
    assert v1['dias_cierre'] == 2.0                      # cotización hace 6 días, PC hace 4
    v2 = fila['ZZREP_VEND2']
    assert v2['cotizaciones'] == 1 and v2['convertidas'] == 1 and v2['venta'] == 1180.0
    # Solo personas de la tabla usuarios
    nombres = {r[0] for r in _sql("SELECT COALESCE(nombres_apellidos, nombre_completo) FROM usuarios")}
    assert set(fila) <= nombres


def test_exportar_csv_xlsx_pdf(rep, admin):
    r = admin.get('/reportes/api/reporte/ventas/exportar?' + _q(cliente='ZZREP', agrupar='documento', formato='csv'))
    assert r.status_code == 200 and 'text/csv' in r.headers['Content-Type']
    texto = r.get_data(as_text=True)
    assert 'Comprobante' in texto and 'F001-90000001' in texto and '1770.00' in texto and '90000003' not in texto
    r = admin.get('/reportes/api/reporte/ventas/exportar?' + _q(cliente='ZZREP', agrupar='documento', formato='xlsx'))
    assert r.status_code == 200
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.data)).active
    filas = list(ws.iter_rows(values_only=True))
    total = next(f for f in filas if f[0] == 'TOTAL')
    assert 1770.0 in total and 1500.0 in total
    r = admin.get('/reportes/api/reporte/cxc/exportar?' + _q(cliente='ZZREP', formato='pdf'))
    assert r.status_code == 200 and r.data[:4] == b'%PDF'


def test_proteccion_por_rol(rep, app, vendedor, aprobador):
    # REGLA DE ORO (Erika 25/10/2026): Reportes solo para Erika. Vendedor, validadora y Finanzas ya no entran.
    fin = _login(app, 'zzrep_fin', 'Clave#2026!')
    for c in (vendedor, aprobador, fin):
        assert c.get('/reportes').status_code == 302
        assert c.get('/reportes/api/reporte/ventas').status_code == 403
        assert c.get('/reportes/api/reporte/ventas/exportar?formato=csv').status_code == 403
        assert c.get('/reportes/api/reporte/cxc').status_code == 403
        assert c.get('/reportes/api/tablero').status_code == 403
        assert c.get('/reportes/api/reporte/kpi').status_code == 403
    # Sin sesión
    assert app.test_client().get('/reportes/api/reporte/ventas').status_code == 401


def test_reporte_inexistente_y_agrupacion_invalida(rep, admin):
    assert admin.get('/reportes/api/reporte/nada').status_code == 404
    assert admin.get('/reportes/api/reporte/ventas?agrupar=xyz').status_code == 400
    admin.post('/api/empresa_activa', json={'codigo': 'KCF'})
