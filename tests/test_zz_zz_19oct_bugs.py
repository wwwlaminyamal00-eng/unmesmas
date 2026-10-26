"""
Pedido de Erika del 19/10/2026 (tarde): bugs de Compras, Proveedores, PV y Exportar + revisión integral.
  BUG 2: no se puede borrar un proveedor con productos (mensaje claro) y se puede desactivar.
  BUG 3: Exportar arma un Excel de la tabla que se ve, con filtros y fecha/hora de Lima en el nombre.
  BUG 4: el resumen de la PV conoce el descuento de la cotización.
  Audio: vencimientos con semáforo (6 y 3 meses), fecha de vencimiento del producto, fecha de creación/modificación.
  Revisión: dos órdenes de compra creadas a la vez no repiten número.
Datos inventados, solo en la base local de pruebas.
"""
import io
import threading
from datetime import date, timedelta

import pytest

from tests.conftest import _login
from tests.test_inventario import _sql


@pytest.fixture(scope='module')
def base19(admin):
    _sql("""
        INSERT INTO erp_empresas (codigo, nombre_corto, nombre_comercial, razon_social, ruc, estado, color_primario, color_secundario, color_pastel)
        SELECT 'KCF', 'KCF', 'KCF', 'KCF S.A.C.', '20111111111', 'activo', '#D6243A', '#1D1A1F', '#FDE8EB'
        WHERE NOT EXISTS (SELECT 1 FROM erp_empresas WHERE codigo = 'KCF');
        INSERT INTO almacenes (codigo, nombre, direccion, activo)
        SELECT 'SMP', 'San Martín de Porres', 'Av. Perú 123, SMP', true WHERE NOT EXISTS (SELECT 1 FROM almacenes WHERE codigo = 'SMP');
    """, fetch=False)
    emp = _sql("SELECT id FROM erp_empresas WHERE codigo = 'KCF'")[0][0]
    sede = _sql("SELECT id FROM almacenes WHERE codigo = 'SMP'")[0][0]
    return {'emp': emp, 'sede': sede}


def _proveedor(nombre, ruc):
    return _sql("""INSERT INTO proveedores (razon_social, ruc, direccion, condicion_pago, activo)
                   VALUES (%s, %s, 'Av. Prueba 1', 'Contado', true) RETURNING id""", (nombre, ruc))[0][0]


# ---------------------------------------------------------------- BUG 2
def test_proveedor_con_productos_no_se_borra_y_se_explica(admin):
    pid = _proveedor('PROVEEDOR CON PRODUCTOS ZZ', '20555555501')
    _sql("""INSERT INTO productos (familia, codigo, descripcion, marca, unidad, costo_unitario, precio_unitario, stock, activo, proveedor_id)
            VALUES ('GENERAL', 'ZZB2-001', 'PRODUCTO DEL PROVEEDOR ZZ', 'ZZ', 'UND', 1, 2, 0, true, %s)""", (pid,), fetch=False)
    r = admin.delete(f'/maestros/api/proveedores/{pid}')
    assert r.status_code == 409
    assert r.json['error'] == ('No se puede eliminar el proveedor porque tiene productos asociados. '
                               'Primero reasigne los productos a otro proveedor.')
    assert r.json['puede_desactivar'] is True and r.json['usos']['productos'] == 1
    # no se tocó nada
    assert _sql("SELECT activo FROM proveedores WHERE id = %s", (pid,))[0][0] is True
    assert _sql("SELECT COUNT(*) FROM productos WHERE proveedor_id = %s", (pid,))[0][0] == 1
    # desactivar = baja lógica, se guarda el historial y el producto sigue con su proveedor
    r = admin.put(f'/maestros/api/proveedores/{pid}/desactivar')
    assert r.status_code == 200 and r.json['success'], r.json
    assert _sql("SELECT activo FROM proveedores WHERE id = %s", (pid,))[0][0] is False
    assert _sql("SELECT COUNT(*) FROM productos WHERE proveedor_id = %s", (pid,))[0][0] == 1
    lista = admin.get('/maestros/api/proveedores/listar').json['data']
    assert pid not in [p['id'] for p in lista]          # ya no aparece para compras nuevas


def test_proveedor_con_ordenes_tampoco_se_borra(admin, base19):
    pid = _proveedor('PROVEEDOR CON OC ZZ', '20555555502')
    _sql("INSERT INTO ordenes_compra (numero_orden, proveedor_id, estado, total) VALUES ('OC-ZZB2-1', %s, 'Borrador', 10)",
         (pid,), fetch=False)
    r = admin.delete(f'/maestros/api/proveedores/{pid}')
    assert r.status_code == 409 and 'historial' in r.json['error']
    assert r.json['usos'] == {'ordenes_compra': 1}
    assert _sql("SELECT COUNT(*) FROM proveedores WHERE id = %s", (pid,))[0][0] == 1


def test_proveedor_sin_uso_si_se_borra(admin):
    pid = _proveedor('PROVEEDOR SIN USO ZZ', '20555555503')
    r = admin.delete(f'/maestros/api/proveedores/{pid}')
    assert r.status_code == 200 and r.json['success']
    assert _sql("SELECT COUNT(*) FROM proveedores WHERE id = %s", (pid,))[0][0] == 0


# ---------------------------------------------------------------- BUG 3
def test_exportar_tabla_a_excel(admin, anonimo):
    from openpyxl import load_workbook
    r = admin.post('/api/exportar/tabla', json={
        'titulo': 'Ventas - Cotizaciones', 'columnas': ['N° cotización', 'Cliente', 'Monto'],
        'filas': [['COT-1', 'CLIENTE A', 'S/ 1,250.50'], ['COT-2', '=HACK()', '0012']],
        'filtros': [['Estado', 'Aceptada'], ['Buscar', 'cliente']]})
    assert r.status_code == 200
    assert r.mimetype.endswith('spreadsheetml.sheet')
    nombre = r.headers['X-Nombre-Archivo']
    assert nombre.startswith('ventas_cotizaciones_') and nombre.endswith('.xlsx')
    assert f'filename="{nombre}"' in r.headers['Content-Disposition']
    ws = load_workbook(io.BytesIO(r.data)).active
    textos = [c.value for c in ws['A'] if c.value]
    assert textos[0] == 'Ventas - Cotizaciones'
    assert 'hora de Lima' in textos[1]
    assert 'Estado: Aceptada' in textos[2]
    enc = [c.value for c in ws[6]]
    assert enc == ['N° cotización', 'Cliente', 'Monto']
    assert [c.value for c in ws[7]] == ['COT-1', 'CLIENTE A', 1250.5]     # el monto queda como número
    assert ws['B8'].value == '=HACK()' and ws['B8'].data_type == 's'      # se guarda como texto, nunca como fórmula
    assert ws['C8'].value == '0012'                                         # los códigos con 0 adelante se respetan
    r2 = admin.post('/api/exportar/tabla', json={'titulo': 'Proveedores', 'columnas': ['RUC', 'Código de proveedor', 'Contacto', 'Total'],
                                                 'filas': [['20600011111', '15', '-', '30']]})
    fila = [c.value for c in load_workbook(io.BytesIO(r2.data)).active[7]]
    assert fila == ['20600011111', '15', '-', 30]                           # RUC y códigos como texto; el guion tal cual
    assert anonimo.post('/api/exportar/tabla', json={'columnas': ['a']}).status_code == 401
    assert admin.post('/api/exportar/tabla', json={'columnas': []}).status_code == 400


def test_botones_exportar_ya_no_son_de_mentira(admin):
    js = admin.get('/static/js/exportar_tabla.js').get_data(as_text=True)
    assert 'window.exportData' in js and 'window.bdExportarExcel' in js and '/api/exportar/tabla' in js
    for pagina in ('/ventas', '/compras', '/inventario', '/maestros', '/finanzas'):
        html = admin.get(pagina, follow_redirects=True).get_data(as_text=True)
        assert 'js/exportar_tabla.js' in html, pagina
    ventas = admin.get('/static/js/ventas.js').get_data(as_text=True)
    assert "showToast(`Exportando datos de ${module}...`, 'info');\r\n}" not in ventas


# ---------------------------------------------------------------- BUG 4
def test_descuento_de_la_cotizacion_para_el_resumen_de_la_pc(vendedor):
    cid = _sql("""INSERT INTO cotizaciones (numero_cotizacion, cliente_id, usuario_id, estado, subtotal,
                                            descuento_tipo, descuento_porcentaje, descuento_monto)
                  VALUES ('COT-ZZB4-1', 1, 1, 'Aceptada por Cliente', 1000, '%', 10, 100) RETURNING id""")[0][0]
    p = _sql("SELECT id, codigo FROM productos WHERE codigo = 'PRD-0001'")[0]
    _sql("INSERT INTO cotizacion_detalle (cotizacion_id, producto_id, cantidad, precio_venta_unitario) VALUES (%s, %s, 10, 100)",
         (cid, p[0]), fetch=False)
    r = vendedor.get('/ventas/api/cotizaciones/descuento', query_string={'numero': 'COT-ZZB4-1'})
    assert r.status_code == 200, r.json
    d = r.json['data']
    assert d['descuento_tipo'] == '%' and d['descuento_porcentaje'] == 10
    assert d['subtotal_lista'] == 1000 and d['subtotal_neto'] == 900 and d['descuento_total'] == 100
    assert d['productos'][p[1]] == {'lista': 100.0, 'neto': 90.0}
    assert vendedor.get('/ventas/api/cotizaciones/descuento?numero=NO-EXISTE').json['data'] is None
    html = vendedor.get('/ventas').get_data(as_text=True)
    assert 'id="pcResumenDescuento"' in html and 'Subtotal − Descuento = Valor venta' in html
    # Orden pedido por Erika (19/10 noche): Subtotal, Valor venta, Descuento, IGV, Total
    pos = [html.index(x) for x in ('id="pcResumenSubtotal"', 'id="pcResumenValorVenta"', 'id="pcResumenDescuento"', 'id="pcResumenIgv"', 'id="pcResumenTotal"')]
    assert pos == sorted(pos)


# ---------------------------------------------------------------- Vencimientos
def test_semaforo_de_vencimiento():
    from routes.inventario import nivel_vencimiento
    hoy = date(2026, 10, 19)
    assert nivel_vencimiento(date(2026, 10, 1), hoy)[0] == 'vencido'
    assert nivel_vencimiento(date(2026, 12, 1), hoy)[0] == 'rojo'          # menos de 3 meses
    assert nivel_vencimiento(date(2027, 3, 1), hoy)[0] == 'amarillo'       # entre 3 y 6 meses
    assert nivel_vencimiento(date(2027, 8, 1), hoy)[0] == 'ok'


def test_vencimientos_por_producto_y_por_lote(admin, base19):
    from routes.productos import asegurar_columnas_productos
    asegurar_columnas_productos()
    hoy = date.today()
    _sql("""INSERT INTO productos (familia, codigo, descripcion, marca, unidad, costo_unitario, precio_unitario, stock, activo, fecha_vencimiento)
            VALUES ('QUIMICOS', 'ZZV-ROJO', 'ALCOHOL ZZV', 'ZZ', 'UND', 1, 2, 0, true, %s),
                   ('QUIMICOS', 'ZZV-AMAR', 'DESINFECTANTE ZZV', 'ZZ', 'UND', 1, 2, 0, true, %s),
                   ('QUIMICOS', 'ZZV-LEJOS', 'JABON ZZV', 'ZZ', 'UND', 1, 2, 0, true, %s)""",
         (hoy + timedelta(days=60), hoy + timedelta(days=150), hoy + timedelta(days=400)), fetch=False)
    prod = dict(_sql("SELECT codigo, id FROM productos WHERE codigo LIKE 'ZZV-%'"))
    for codigo in prod:
        _sql("INSERT INTO inv_stock (empresa_id, almacen_id, producto_id, cantidad) VALUES (%s, %s, %s, 5)",
             (base19['emp'], base19['sede'], prod[codigo]), fetch=False)
    r = admin.get('/inventario/api/vencimientos')
    assert r.status_code == 200, r.json
    por_codigo = {v['codigo']: v for v in r.json['data']}
    assert por_codigo['ZZV-ROJO']['nivel'] == 'rojo' and por_codigo['ZZV-ROJO']['origen'] == 'Producto'
    assert por_codigo['ZZV-AMAR']['nivel'] == 'amarillo'
    assert 'ZZV-LEJOS' not in por_codigo                                    # vence en más de 6 meses
    assert r.json['conteo']['rojo'] >= 1 and r.json['conteo']['amarillo'] >= 1
    todo = {v['codigo'] for v in admin.get('/inventario/api/vencimientos?dias=36500').json['data']}
    assert 'ZZV-LEJOS' in todo
    solo_rojo = admin.get('/inventario/api/vencimientos?nivel=rojo').json['data']
    assert all(v['nivel'] == 'rojo' for v in solo_rojo)
    res = admin.get('/inventario/api/resumen').json['data']
    assert res['vencimientos_semaforo']['rojo'] >= 1


def test_producto_guarda_fecha_de_vencimiento(admin):
    r = admin.post('/productos/api/productos', json={
        'descripcion': 'LEJIA ZZV', 'modelo': 'X', 'marca': 'ZZ', 'familia': 'QUIMICOS', 'unidad': 'UND',
        'transporte': 'auto', 'fecha_vencimiento': '2027-01-31'})
    assert r.status_code == 200 and r.json['success'], r.json
    pid = r.json['data']['id']
    assert str(_sql("SELECT fecha_vencimiento FROM productos WHERE id = %s", (pid,))[0][0]) == '2027-01-31'
    assert admin.get(f'/productos/api/productos/{pid}').json['data']['fecha_vencimiento'].startswith('2027-01-31')
    # otra pantalla que edita el producto sin mandar la fecha no la borra
    r = admin.put(f'/productos/api/productos/{pid}', json={'descripcion': 'LEJIA ZZV 2', 'modelo': 'X', 'marca': 'ZZ',
                                                          'familia': 'QUIMICOS', 'unidad': 'UND', 'transporte': 'auto'})
    assert r.status_code == 200, r.json
    assert str(_sql("SELECT fecha_vencimiento FROM productos WHERE id = %s", (pid,))[0][0]) == '2027-01-31'
    r = admin.put(f'/productos/api/productos/{pid}', json={'descripcion': 'LEJIA ZZV 2', 'modelo': 'X', 'marca': 'ZZ',
                                                          'familia': 'QUIMICOS', 'unidad': 'UND', 'transporte': 'auto',
                                                          'fecha_vencimiento': ''})
    assert _sql("SELECT fecha_vencimiento FROM productos WHERE id = %s", (pid,))[0][0] is None
    r = admin.put(f'/productos/api/productos/{pid}', json={'descripcion': 'LEJIA ZZV 2', 'modelo': 'X', 'marca': 'ZZ',
                                                          'familia': 'QUIMICOS', 'unidad': 'UND', 'transporte': 'auto',
                                                          'fecha_vencimiento': 'mañana'})
    assert r.status_code == 400


# ---------------------------------------------------------------- Fecha y hora, nombres claros
def test_listas_traen_fecha_de_creacion_y_modificacion(admin):
    _sql("""INSERT INTO guias_remision (serie, numero, fecha_emision, documento_asociado, destinatario_nombre, ruc_destinatario, estado_sunat)
            VALUES ('T0ZZ', '00000001', NOW(), 'COT-ZZB4-1', 'CLIENTE ZZ', '20888888899', 'Emitida');
            INSERT INTO comprobantes (tipo_comprobante, serie, numero, fecha_emision, cliente_tipo_doc, cliente_nombre, cliente_numero_doc,
                                      subtotal, igv, total, documento_asociado, estado_sunat, items_json)
            VALUES ('Factura', 'F0ZZ', '1', NOW(), '6', 'CLIENTE ZZ', '20888888899', 100, 18, 118, 'COT-ZZB4-1', 'Emitida', '[]')""",
         fetch=False)
    for url in ('/ventas/api/cotizaciones/listar', '/ventas/api/guias/listar', '/ventas/api/comprobantes/listar'):
        data = admin.get(url).json['data']
        assert data, url
        assert 'creado_en' in data[0] and 'modificado_en' in data[0], url
    html = admin.get('/ventas').get_data(as_text=True)
    assert '>N° cotización</th>' in html and '>N° despacho</th>' in html and '<th>Comprobante</th>' not in html
    base = admin.get('/inventario', follow_redirects=True).get_data(as_text=True)
    assert 'window.tiemposHTML' in base and 'window.condPagoTexto' in base


# ---------------------------------------------------------------- BUG 1
def test_condicion_de_pago_credito_se_lee_claro(admin):
    html = admin.get('/compras').get_data(as_text=True)
    assert 'Número de días de crédito' in html and 'id="ocCondTexto"' in html
    base = admin.get('/compras').get_data(as_text=True)
    assert "'Crédito ' + n + ' días'" in base
    modal = admin.get('/maestros', follow_redirects=True).get_data(as_text=True)
    assert 'id="prov_diasCredito"' in modal and "personalizarDiasCredito('prov_condicion'" not in modal


# ---------------------------------------------------------------- Estrés: correlativos
def test_dos_ordenes_a_la_vez_no_repiten_numero(app, base19):
    prov = _proveedor('PROVEEDOR ESTRES ZZ', '20555555504')
    prod = _sql("SELECT id FROM productos WHERE codigo = 'PRD-0001'")[0][0]
    clientes = [_login(app, 'admin', 'Admin#2026!') for _ in range(4)]
    numeros, errores = [], []
    barrera = threading.Barrier(len(clientes))

    def crear(c, n):
        try:
            barrera.wait()
            r = c.post('/compras/api/flujo/ordenes', json={
                'empresa_id': base19['emp'], 'proveedor_id': prov, 'almacen_entrega_id': base19['sede'],
                'condicion_pago': 'Crédito', 'dias_credito': 30, 'modalidad_entrega': 'Proveedor entrega',
                'fecha_compromiso': '2026-10-30', 'motivo_compra': 'Stock estratégico',
                'items': [{'producto_id': prod, 'cantidad': n, 'precio_unitario': 5}]})   # cantidades distintas: no son duplicadas
            if r.status_code == 200:
                numeros.append(r.json['numero'] if 'numero' in r.json else r.json['id'])
            else:
                errores.append(r.json)
        except Exception as e:      # pragma: no cover
            errores.append(str(e))

    hilos = [threading.Thread(target=crear, args=(c, i + 1)) for i, c in enumerate(clientes)]
    [h.start() for h in hilos]
    [h.join() for h in hilos]
    assert not errores, errores
    filas = _sql("SELECT numero_orden FROM ordenes_compra WHERE proveedor_id = %s", (prov,))
    assert len(filas) == 4 and len({f[0] for f in filas}) == 4
    assert _sql("SELECT condicion_pago FROM ordenes_compra WHERE proveedor_id = %s LIMIT 1", (prov,))[0][0] == 'Crédito 30 días'
